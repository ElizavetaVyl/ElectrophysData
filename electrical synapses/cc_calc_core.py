CC_MODE_BLOCK_KEYS = ("cc", "cc_last_pos", "cc_most_neg", "cc_neg_pos")
CC_MODE_SELECTION_LABELS = {
    "cc": "all_pre_spike",
    "cc_last_pos": "last_positive_pre_spike",
    "cc_most_neg": "most_negative_pre_spike",
    "cc_neg_pos": "most_negative_and_most_positive_pre_spike",
}
CC_MODE_SELECTION_NOTES = {
    "cc": "all sweeps before the first sweep with at least one spike",
    "cc_last_pos": "last positive sweep before the first sweep with at least one spike",
    "cc_most_neg": "most negative sweep before the first sweep with at least one spike",
    "cc_neg_pos": "most negative sweep and most positive sweep before the first sweep with at least one spike",
}
CC_MODE_RUN_DIRS = {
    "cc": "CC_multi_sweeps",
    "cc_last_pos": "CC_last_positive_pre_spike",
    "cc_most_neg": "CC_most_negative_pre_spike",
    "cc_neg_pos": "CC_most_negative_most_positive_pre_spike",
}

"""Core logic for coupling coefficient batch analysis (imported by notebook).

Notebook: ``CC calculation.ipynb`` (cells 0-3).
Batch loop and Excel writers live in the notebook, not here.
"""

import os
import re
import statistics
import traceback
from datetime import datetime

import numpy as np
import pyabf
import scipy.signal
from scipy.ndimage import gaussian_filter1d
from scipy.signal import find_peaks

# =============================================================================
# CONFIG
# =============================================================================

EXPECTED_CHANNELS = 4  # double_cciv: 2x Vm + 2x I

CC_SPIKE_HEIGHT = -10  # mV; AP if find_peaks sees Vm above this
CC_SPIKE_DISTANCE = 10  # samples; min distance between peaks
CC_MIN_DELTA_I_PA = 10  # pA; skip CC if |stim current step| smaller
CC_MIN_DELTA_V_MV = 10  # mV; skip CC if |delta_V_active| (pre-post) is smaller

RIN_VMIN = -80  # mV; primary I–V window (mean Vm in stim epoch)
RIN_VMAX = -50  # mV
RIN_VMIN_FALLBACK = -95  # mV; wider window if too few points in primary
RIN_VMAX_FALLBACK = -50  # mV
RIN_MIN_POINTS = 2  # minimum sweeps for V(I) linear fit
# Unit conversion: R [MΩ] = (ΔV [mV] / ΔI [pA]) × 1000  OR  slope(V vs I) [mV/pA] × 1000

# Cell properties (Cell Properties for one folder.ipynb); separate from CC Rin
CP_SPIKE_HEIGHT = -10  # mV
CP_SPIKE_DISTANCE = 10  # samples
CP_MIN_APS = 4  # sweep must have >= this many APs in stim window
CP_RIN_VMIN = -85  # mV; V_rest / Rin_abs stage 1
CP_RIN_VMAX = -50  # mV
CP_RIN_VMIN_FALLBACK = -95  # mV; stage 2 (same idea as CC Rin abs fallback)
CP_RIN_VMAX_FALLBACK = -50  # mV
CP_RIN_VMIN_WIDE = -95  # mV; stage 3 if still <2 points (later depolarized files)
CP_RIN_VMAX_WIDE = -40  # mV
CP_RIN_MIN_POINTS = 2  # linear V_rest / Rin_abs needs >=2 I–V points
CP_RIN_REL_MIN_DI_PA = 0.15  # pA; skip per-sweep Rin_rel if |delta I| smaller
# V_rest = intercept at I=0, so only linear stages (not Rin relative dV/dI)
CP_VREST_STAGES = (
    (1, CP_RIN_VMIN, CP_RIN_VMAX, "1_primary"),
    (2, CP_RIN_VMIN_FALLBACK, CP_RIN_VMAX_FALLBACK, "2_extended"),
    (3, CP_RIN_VMIN_WIDE, CP_RIN_VMAX_WIDE, "3_wide"),
)

SAVE_QC_PLOTS = True  # save AP + I–V + tau/Cm PNGs (cell-properties style QC)
CELL_PROPS_PLOTS_SUBDIR = "Cell_properties_plots"  # AP / Rin / tau figures
QC_PLOTS_SUBDIR = CELL_PROPS_PLOTS_SUBDIR  # backward-compatible alias
SAVE_CC_PLOTS = True  # save CC QC PNGs when cc_plots_dir is passed
SAVE_CC_TRACE_PLOTS = True  # *_CC_traces.png (sweep QC for CC/Gj)
SAVE_CC_VPOST_PLOTS = True  # *_CC_vs_Vpost.png per file
CC_PLOTS_SUBDIR = "CC_plots"  # subfolder for coupling-coefficient figures
PLOT_DPI = 100  # PNG resolution (lower = faster writes; was 150)
CC_VM_FIT_LINEAR = True  # solid line on folder CC_norm vs Vm
CC_VM_LINEAR_MIN_POINTS = 3  # no line if the file has fewer CC_norm points
CC_VM_CMAP = "viridis"  # file color = first→last by file number in the name

# Which analysis to run (notebook dialog can change this per session)
RUN_CC = True  # CC, Gj, Rin-for-CC, CC QC + folder CC plots
RUN_CELL_PROPS = True  # V_rest, firing metrics, AP / I–V QC plots
RUN_TAU_CM = True  # tau/Cm (uses Rin; enable CC or Rin will still be computed)
RUN_SPIKELETS = True  # spike/spikelet; AP start from cp_* inflections (not the full cell-props block)
ANALYSIS_BLOCKS = {
    "cc": RUN_CC,
    "cell_props": RUN_CELL_PROPS,
    "tau_cm": RUN_TAU_CM,
    "spikelets": RUN_SPIKELETS,
}
_SESSION_BLOCKS = None  # filled after the once-per-run chooser window

# Spikelet coupling (AP2+ on first >=4 AP sweep, else 3, else 2; else AP1 if only 1-spike sweeps)
SPIKELET_BASELINE_MS = 1.0  # passive mean Vm in [t_start-1ms, t_start); not used as t=0
SPIKELET_PEAK_MS = 15.0  # search passive peak in [t_start, t_start+15ms]
SPIKELET_PEAK_SMOOTH_MS = 0.3  # Gaussian σ for local-max search (kills 1-sample jitter)
SPIKELET_MIN_PROMINENCE_MV = 0.15  # original per-AP scheme: drop after the top
SPIKELET_MEANTRACE_MIN_PROMINENCE_MV = 0.05  # meantrace: looser drop than 0.15 mV
SPIKELET_USE_NOISE_GATE = False  # amplitude vs noise; peak shape is separate (prominence)
SPIKELET_NOISE_K = 1.0  # unused while SPIKELET_USE_NOISE_GATE is False
SPIKELET_NOISE_K_REF = 3.0  # old prestim 3× bar, unused on QC
SPIKELET_NOISE_LOCAL_MS = 10.0  # MAD on passive in [t0-10ms, t0), pooled per sweep
SPIKELET_MIN_AMP_MV = 0.15  # unused while SPIKELET_USE_NOISE_GATE is False
SPIKELET_MIN_APS = 2  # prefer AP2+; fall back to 1-AP sweeps if none exist
SPIKELET_DELAY_FRAC = 0.10  # delay_10: 10% of AP amp and 10% of spikelet amp
SPIKELET_MEANTRACE_PEAK_MS = 15.0  # parallel scheme: average passive trace in [t0, t0+15ms]
SAVE_SPIKELET_PLOTS = True
SPIKELET_PLOTS_SUBDIR = "Spikelet_plots"
SPIKELET_DIR_TAG = {"ch0->ch2": "12", "ch2->ch0": "21"}

# Tau / Cm (Tau Cm calculations.ipynb); Rin = same CC/Gj rin_for_channel
TCM_VMIN_LIMIT = -90  # mV; skip sweep if min Vm in post epoch is below this
TCM_DV_SEARCH_MIN = -35  # mV; delta_V = mean(post) - mean(pre) search range
TCM_DV_SEARCH_MAX = -10  # mV
TCM_PRE_MS = 0.125  # s; prestim mean window before stim (125 ms)

CELL_PROPS_FIELD_SUFFIXES = (
    "AP21_ratio",
    "init_freq_Hz",
    "late_freq_Hz",
    "mean_freq_Hz",
    "FWHM_ms",
    "delay_AP1_ms",
    "Rin_abs_MOhm",
    "Rin_rel_MOhm",
    "V_rest_mV",
    "V_rest_skip_reason",
    "V_rest_note",
    "V_rest_stage",
    "hold_V_mV",
    "inj_current_pA",
    "R2_abs_Rin",
    "props_sweep",
    "props_skip_reason",
)

SPIKELET_SUMMARY_SUFFIXES = (
    "sweep",
    "n_AP_active",
    "sweep_tier",
    "n_AP_used",
    "n_spikelet_detected",
    "n_with_amp",
    "mean_amp_active_mV",
    "mean_amp_spikelet_mV",
    "mean_amp_ratio",
    "mean_delay_ms",
    "mean_delay_10_ms",
    "n_delay_negative",
    "n_delay_10_negative",
    "n_ratio_gt_1",
    "avg_amp_active_mV",
    "avg_amp_spikelet_mV",
    "avg_amp_ratio",
    "avg_delay_ms",
    "avg_delay_10_ms",
    "meantrace10_amp_active_mV",
    "meantrace10_amp_spikelet_mV",
    "meantrace10_amp_ratio",
    "meantrace10_delay_ms",
    "meantrace10_delay_10_ms",
    "meantrace10_detected",
    "meantrace10_skip_reason",
    "metric_source",
    "ap1_fallback",
    "rms_noise_mV",
    "noise_thr_mV",
    "skip_reason",
    "n_sweeps",
    "mean_vm_begin_mV",
    "mean_baseline_passive_mV",
)

SPIKELET_AP_KEYS = (
    "file",
    "recording_datetime",
    "direction",
    "sweep",
    "ap_index",
    "t_start_ms",
    "t_peak_active_ms",
    "t_peak_passive_ms",
    "t_10_active_ms",
    "t_10_spikelet_ms",
    "amp_active_mV",
    "vm_begin_active_mV",
    "baseline_passive_mV",
    "amp_spikelet_mV",
    "amp_ratio",
    "delay_ms",
    "delay_peak_ms",
    "delay_10_ms",
    "delay_negative",
    "delay_10_negative",
    "ratio_gt_1",
    "rms_local_mV",
    "noise_thr_mV",
    "detected",
    "used_in_average",
    "skip_reason",
    "metric_source",
)


SPIKELET_SWEEP_KEYS = (
    "file",
    "recording_datetime",
    "direction",
    "sweep",
    "n_AP_active",
    "n_AP_used",
    "n_spikelet_detected",
    "n_with_amp",
    "mean_amp_ratio",
    "mean_amp_active_mV",
    "mean_amp_spikelet_mV",
    "mean_delay_ms",
    "mean_delay_10_ms",
    "n_delay_negative",
    "n_delay_10_negative",
    "n_ratio_gt_1",
    "mean_vm_begin_mV",
    "mean_baseline_passive_mV",
    "noise_thr_mV",
    "avg_amp_ratio",
    "avg_amp_active_mV",
    "avg_amp_spikelet_mV",
    "avg_delay_ms",
    "avg_delay_10_ms",
    "meantrace10_amp_active_mV",
    "meantrace10_amp_spikelet_mV",
    "meantrace10_amp_ratio",
    "meantrace10_delay_ms",
    "meantrace10_delay_10_ms",
    "meantrace10_detected",
    "meantrace10_skip_reason",
    "metric_source",
    "ap1_fallback",
    "skip_reason",
    "is_primary",
)


def empty_spikelet_summary_fields():
    out = {}
    for tag in ("12", "21"):
        for key in SPIKELET_SUMMARY_SUFFIXES:
            out[f"spikelet_{key}_{tag}"] = None
    return out


TAU_CM_FIELD_SUFFIXES = (
    "tau_ms",
    "Cm_pF",
    "V_pre_mV",
    "V_post_mV",
    "V_post_min_mV",
    "delta_V_mV",
    "tau_sweep",
    "tau_selection_note",
    "tau_skip_reason",
    "Cm_skip_reason",
)


def empty_cell_props_fields():
    """Default None values for cell-property columns on both channels."""
    out = {}
    for ch in ("ch0", "ch2"):
        for key in CELL_PROPS_FIELD_SUFFIXES:
            out[f"{key}_{ch}"] = None
    return out


def empty_tau_cm_fields():
    out = {}
    for ch in ("ch0", "ch2"):
        for key in TAU_CM_FIELD_SUFFIXES:
            out[f"{key}_{ch}"] = None
    return out


# Per-sweep columns only (file-level metrics -> File_summary sheet)
ALL_DATA_SWEEP_KEYS = (
    "file",
    "recording_datetime",
    "sweep",
    "direction",
    "CC_selection_mode",
    "cur_step_pA",
    "delta_V_active_mV",
    "delta_V_passive_mV",
    "Vm_active_stim_mV",
    "CC",
    "CC_norm",
    "CC_skip_reason",
    "Gj_sweep_nS",
    "Gj_sweep_skip_reason",
)


def sweep_only_row(row_dict):
    return {k: row_dict.get(k) for k in ALL_DATA_SWEEP_KEYS}


def time_period_borders(data_rate):
    sr = int(data_rate)
    if sr == 20000:
        return dict(
            start10=0, end10=2800, start11=6000, end11=12800,
            start20=28000, end20=30800, start21=34000, end21=40800,
            Rtime_ch0=8000, Rtime_ch2=36000,
        )
    if sr == 10000:
        return dict(
            start10=0, end10=1400, start11=3000, end11=6400,
            start20=14000, end20=15400, start21=16600, end21=20000,
            Rtime_ch0=4000, Rtime_ch2=18000,
        )
    raise ValueError(f"Unsupported sampling rate: {sr} Hz")


def recording_datetime_str(abf):
    dt = getattr(abf, "abfDateTime", None)
    if dt is None:
        return getattr(abf, "abfDateTimeString", None)
    if isinstance(dt, datetime):
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    return str(dt)


def resolve_analysis_blocks(blocks=None):
    """Normalized analysis flags. At least one stays on."""
    out = {
        "cc": True,
        "cc_last_pos": False,
        "cc_most_neg": False,
        "cc_neg_pos": False,
        "cell_props": True,
        "tau_cm": True,
        "spikelets": True,
    }
    src = ANALYSIS_BLOCKS if blocks is None else blocks
    if src:
        for key in out:
            if key in src:
                out[key] = bool(src[key])
    chosen_cc = [key for key in CC_MODE_BLOCK_KEYS if out.get(key)]
    if len(chosen_cc) > 1:
        keep = chosen_cc[0]
        for key in CC_MODE_BLOCK_KEYS:
            out[key] = (key == keep)
    if not any(out.values()):
        out["spikelets"] = True
    return out


def selected_cc_mode(blocks=None):
    """Return the selected CC block key, or None if no CC mode is enabled."""
    chosen = resolve_analysis_blocks(blocks)
    for key in CC_MODE_BLOCK_KEYS:
        if chosen.get(key):
            return key
    return None


def cc_mode_run_dir_name(mode_key):
    """Folder name for one CC mode run."""
    return CC_MODE_RUN_DIRS.get(mode_key)


def ask_analysis_blocks(initial=None):
    """Tk window: choose which analysis blocks to run this session."""
    from tkinter import BooleanVar, Button, Checkbutton, Frame, Label, Tk

    chosen = resolve_analysis_blocks(initial)
    print(">>> A window should open: Analysis blocks.")
    print(">>> If you do not see it, check behind Jupyter or the taskbar.")
    print(">>> Uncheck blocks you do not need, then OK.")

    root = Tk()
    root.title("Analysis blocks")
    root.geometry("620x340+120+80")
    try:
        root.attributes("-topmost", True)
        root.lift()
        root.focus_force()
    except Exception:
        pass
    root.update_idletasks()

    Label(
        root,
        text="Select analysis blocks for this run",
        font=("Segoe UI", 14, "bold"),
        pady=8,
        padx=12,
    ).pack(anchor="w")
    Label(
        root,
        text=(
            "Uncheck a block to skip it (faster).\n"
            "Choose one CC block: multi-sweeps, last positive pre-spike, most negative pre-spike, or most negative + most positive.\n"
            "Spike/spikelet AP start uses inflections (peaks + d²V) inside that block.\n"
            "The Cell properties block is not required for spikelets."
        ),
        justify="left",
        padx=12,
    ).pack(anchor="w")

    vars_ = {}
    labels = (
        ("cc", "CC / Gj / Rin   (all sweeps before first spike)"),
        ("cc_last_pos", "CC / Gj / Rin   (last positive sweep before first spike)"),
        ("cc_most_neg", "CC / Gj / Rin   (most negative sweep before first spike)"),
        ("cc_neg_pos", "CC / Gj / Rin   (most negative + most positive before first spike)"),
        ("cell_props", "Cell properties   (V_rest, firing, AP / I–V QC plots)"),
        ("tau_cm", "Tau / Cm"),
        ("spikelets", "Spike / spikelet   (QC + amplitude/delays vs time and vs Vm)"),
    )
    for key, text in labels:
        var = BooleanVar(value=chosen[key])
        vars_[key] = var
        Checkbutton(root, text=text, variable=var, anchor="w", font=("Segoe UI", 11)).pack(
            fill="x", padx=16, pady=2
        )

    def _read():
        return {key: bool(var.get()) for key, var in vars_.items()}

    def _finish():
        chosen.clear()
        chosen.update(resolve_analysis_blocks(_read()))
        try:
            root.grab_release()
        except Exception:
            pass
        root.destroy()

    def _all():
        for var in vars_.values():
            var.set(True)

    btns = Frame(root)
    btns.pack(pady=14)
    Button(btns, text="Select all", command=_all, width=14, height=2).pack(side="left", padx=6)
    Button(btns, text="OK", command=_finish, width=10, height=2).pack(side="left", padx=6)
    root.protocol("WM_DELETE_WINDOW", _finish)
    try:
        root.grab_set()
    except Exception:
        pass
    root.resizable(False, False)
    root.mainloop()
    print(">>> Selected blocks:", chosen)
    return chosen


def _ask_analysis_blocks_console():
    """Fallback if the Tk window cannot open."""
    print("Tk window did not open. Press Enter for all blocks,")
    print("or type a subset: cc, cc_last_pos, cc_most_neg, cc_neg_pos, cell_props, tau_cm, spikelets")
    try:
        raw = input("Blocks: ").strip().lower()
    except Exception:
        raw = ""
    if not raw:
        return resolve_analysis_blocks({
            "cc": True, "cc_last_pos": False, "cc_most_neg": False, "cc_neg_pos": False,
            "cell_props": True, "tau_cm": True, "spikelets": True,
        })
    wanted = {p.strip().replace("-", "_") for p in raw.replace(";", ",").split(",") if p.strip()}
    aliases = {
        "spikelet": "spikelets", "spikelets": "spikelets",
        "cell": "cell_props",
        "cc_multi": "cc",
        "cclastpos": "cc_last_pos",
        "ccmostneg": "cc_most_neg",
        "ccnegpos": "cc_neg_pos",
    }
    wanted = {aliases.get(x, x) for x in wanted}
    return resolve_analysis_blocks({
        "cc": "cc" in wanted,
        "cc_last_pos": "cc_last_pos" in wanted,
        "cc_most_neg": "cc_most_neg" in wanted,
        "cc_neg_pos": "cc_neg_pos" in wanted,
        "cell_props": "cell_props" in wanted,
        "tau_cm": "tau_cm" in wanted,
        "spikelets": "spikelets" in wanted,
    })


def ensure_analysis_blocks(blocks=None, force_ask=False):
    """Ask once per kernel session; reuse after that."""
    global ANALYSIS_BLOCKS, _SESSION_BLOCKS
    if blocks is not None:
        _SESSION_BLOCKS = resolve_analysis_blocks(blocks)
        ANALYSIS_BLOCKS = _SESSION_BLOCKS
        return _SESSION_BLOCKS
    if _SESSION_BLOCKS is not None and not force_ask:
        return _SESSION_BLOCKS
    try:
        chosen = ask_analysis_blocks()
    except Exception as exc:
        print("Analysis-block window failed:", exc)
        chosen = _ask_analysis_blocks_console()
    _SESSION_BLOCKS = resolve_analysis_blocks(chosen)
    ANALYSIS_BLOCKS = _SESSION_BLOCKS
    on = [k for k, v in _SESSION_BLOCKS.items() if v]
    print(">>> This run will compute:", ", ".join(on) if on else "(none)")
    return _SESSION_BLOCKS


def mean_delta_voltage(sweep_y, pre_start, pre_end, post_start, post_end):
    pre_seg = sweep_y[pre_start:pre_end]
    post_seg = sweep_y[post_start:post_end]
    if len(pre_seg) == 0 or len(post_seg) == 0:
        raise ValueError(
            f"empty epoch window (pre {pre_start}:{pre_end}, post {post_start}:{post_end})"
        )
    pre = statistics.mean(pre_seg)
    post = statistics.mean(post_seg)
    return pre - post


def n_spikes_in_window(sweep_y, win_start, win_end, height=CC_SPIKE_HEIGHT, distance=CC_SPIKE_DISTANCE):
    peaks, _ = find_peaks(sweep_y[win_start:win_end], height=height, distance=distance)
    return len(peaks)


def cc_sweep_indices_direction(
    abf,
    active_ch,
    passive_ch,
    pre_start,
    pre_end,
    post_start,
    post_end,
    height=CC_SPIKE_HEIGHT,
    distance=CC_SPIKE_DISTANCE,
):
    """Subthreshold sweeps: no AP on active or passive in stimulus period (pre_start:post_end)."""
    indices = []
    for sn in abf.sweepList:
        abf.setSweep(sweepNumber=sn, channel=active_ch)
        if n_spikes_in_window(abf.sweepY, pre_start, post_end, height, distance) > 0:
            break
        abf.setSweep(sweepNumber=sn, channel=passive_ch)
        if n_spikes_in_window(abf.sweepY, pre_start, post_end, height, distance) > 0:
            break
        indices.append(sn)
    return indices


def _cc_direction_step_map(
    abf,
    sweeps,
    current_ch,
    pre_start,
    pre_end,
    post_start,
    post_end,
):
    """Map sweep -> delta_I [pA] for one CC direction."""
    out = {}
    for sn in sweeps:
        abf.setSweep(sweepNumber=sn, channel=current_ch)
        pre = float(statistics.mean(abf.sweepY[pre_start:pre_end]))
        post = float(statistics.mean(abf.sweepY[post_start:post_end]))
        out[sn] = post - pre
    return out


def cc_select_sweeps_direction(
    abf,
    active_ch,
    passive_ch,
    current_ch,
    pre_start,
    pre_end,
    post_start,
    post_end,
    mode_key="cc",
    height=CC_SPIKE_HEIGHT,
    distance=CC_SPIKE_DISTANCE,
):
    """Pick sweeps for one CC mode and direction."""
    indices = cc_sweep_indices_direction(
        abf, active_ch, passive_ch, pre_start, pre_end, post_start, post_end,
        height=height, distance=distance,
    )
    if not indices or mode_key == "cc":
        return indices

    step_map = _cc_direction_step_map(
        abf, indices, current_ch, pre_start, pre_end, post_start, post_end,
    )
    negative = [sn for sn in indices if step_map.get(sn, 0.0) < 0]
    positive = [sn for sn in indices if step_map.get(sn, 0.0) > 0]

    if mode_key == "cc_last_pos":
        if positive:
            return [positive[-1]]
        return [max(indices, key=lambda sn: step_map[sn])] if indices else []
    if mode_key == "cc_most_neg":
        if negative:
            return [min(negative, key=lambda sn: step_map[sn])]
        return [min(indices, key=lambda sn: step_map[sn])] if indices else []
    if mode_key == "cc_neg_pos":
        picks = []
        if negative:
            picks.append(min(negative, key=lambda sn: step_map[sn]))
        elif indices:
            picks.append(min(indices, key=lambda sn: step_map[sn]))
        if positive:
            picks.append(max(positive, key=lambda sn: step_map[sn]))
        elif indices:
            picks.append(max(indices, key=lambda sn: step_map[sn]))
        return sorted(set(picks))

    return indices


# backward-compatible name (single channel, given window)
def cc_sweep_indices(abf, channel, win_start, win_end, height=CC_SPIKE_HEIGHT, distance=CC_SPIKE_DISTANCE):
    indices = []
    for sn in abf.sweepList:
        abf.setSweep(sweepNumber=sn, channel=channel)
        if n_spikes_in_window(abf.sweepY, win_start, win_end, height, distance) > 0:
            break
        indices.append(sn)
    return indices


def cc_skip_reason(delta_i, delta_v_active, passive_sweep_y, pre_start, pre_end, post_start, post_end):
    if n_spikes_in_window(passive_sweep_y, pre_start, post_end) > 0:
        return (
            f"spike on passive cell in stimulus period "
            f"(samples {pre_start}-{post_end})"
        )
    if abs(delta_i) < CC_MIN_DELTA_I_PA:
        return (
            f"|delta_I| < {CC_MIN_DELTA_I_PA} pA in stim current channel "
            f"(pre {pre_start}-{pre_end}, post {post_start}-{post_end})"
        )
    if abs(delta_v_active) < CC_MIN_DELTA_V_MV:
        return (
            f"|delta_V_active| < {CC_MIN_DELTA_V_MV} mV "
            f"(pre {pre_start}-{pre_end}, post {post_start}-{post_end})"
        )
    return None


def coupling_block(abf, sweeps, active_ch, passive_ch, cur_ch, windows):
    s10, e10, s11, e11 = windows
    rows = []
    for sn in sweeps:
        abf.setSweep(sweepNumber=sn, channel=active_ch)
        dva = mean_delta_voltage(abf.sweepY, s10, e10, s11, e11)
        vm_stim = float(statistics.mean(abf.sweepY[s11:e11]))
        abf.setSweep(sweepNumber=sn, channel=cur_ch)
        di = mean_delta_voltage(abf.sweepY, s10, e10, s11, e11)
        abf.setSweep(sweepNumber=sn, channel=passive_ch)
        dvp = mean_delta_voltage(abf.sweepY, s10, e10, s11, e11)
        reason = cc_skip_reason(di, dva, abf.sweepY, s10, e10, s11, e11)
        cc = None if reason else round(dvp / dva, 4)
        rows.append({
            "sweep": sn,
            "CC_selection_mode": None,
            "cur_step_pA": round(di, 1),
            "delta_V_active_mV": round(dva, 4),
            "delta_V_passive_mV": round(dvp, 4),
            "Vm_active_stim_mV": round(vm_stim, 4),
            "CC": cc,
            "CC_norm": None,
            "CC_skip_reason": reason,
        })
    return rows


def cc_normalize_block(block, cc_mean):
    """Set CC_norm = CC / file mean CC for this direction (valid sweeps only)."""
    if cc_mean in (None, 0):
        return block
    for r in block:
        cc = r.get("CC")
        r["CC_norm"] = round(cc / cc_mean, 4) if cc is not None else None
    return block


def last_subthreshold_sweep(
    abf, voltage_ch, spike_start, spike_end,
    height=CC_SPIKE_HEIGHT, distance=CC_SPIKE_DISTANCE,
):
    """
    Last sweep index before the first AP on voltage_ch in [spike_start:spike_end].

    If no AP in any sweep -> last sweep in the file.
    If sweep 0 already has an AP -> None (no usable subthreshold sweeps).
    """
    last_ok = None
    for sn in abf.sweepList:
        abf.setSweep(sweepNumber=sn, channel=voltage_ch)
        if n_spikes_in_window(abf.sweepY, spike_start, spike_end, height, distance) > 0:
            return last_ok
        last_ok = sn
    return last_ok


def collect_iv_points(
    abf, voltage_ch, current_ch, v_start, v_end, i_index, vmin, vmax, sweep_max=None,
):
    """Absolute (I at Rtime, mean Vm in stim epoch) for linear Rin fit."""
    currents, voltages = [], []
    for sn in abf.sweepList:
        if sweep_max is not None and sn > sweep_max:
            break
        abf.setSweep(sweepNumber=sn, channel=voltage_ch)
        vm = statistics.mean(abf.sweepY[v_start:v_end])
        if vmin <= vm <= vmax:
            voltages.append(vm)
            abf.setSweep(sweepNumber=sn, channel=current_ch)
            currents.append(abf.sweepY[i_index])
    return currents, voltages


def collect_rin_delta_points(
    abf, voltage_ch, current_ch, pre_start, pre_end, post_start, post_end, vmin, vmax,
    sweep_max=None,
):
    """Per sweep: ΔV/ΔI if mean post-stim Vm in [vmin, vmax] and |ΔI| large enough."""
    delta_is, delta_vs = [], []
    for sn in abf.sweepList:
        if sweep_max is not None and sn > sweep_max:
            break
        abf.setSweep(sweepNumber=sn, channel=voltage_ch)
        v_post = statistics.mean(abf.sweepY[post_start:post_end])
        if not (vmin <= v_post <= vmax):
            continue
        v_pre = statistics.mean(abf.sweepY[pre_start:pre_end])
        abf.setSweep(sweepNumber=sn, channel=current_ch)
        i_pre = statistics.mean(abf.sweepY[pre_start:pre_end])
        i_post = statistics.mean(abf.sweepY[post_start:post_end])
        di = i_post - i_pre
        if abs(di) < CC_MIN_DELTA_I_PA:
            continue
        delta_is.append(di)
        delta_vs.append(v_post - v_pre)
    return delta_is, delta_vs


def rin_r2(currents, voltages):
    """Linear fit V(I); inputs pA and mV -> Rin in MΩ (= slope [mV/pA] × 1000)."""
    if len(currents) < RIN_MIN_POINTS:
        return None, None
    try:
        fit = np.polyfit(currents, voltages, 1, full=True)
        rin = np.round(fit[0][0] * 1000.0, 3)
        sse = fit[1][0]
        sst = ((np.array(voltages) - statistics.mean(voltages)) ** 2).sum()
        r2 = None if sst == 0 else np.round(1 - sse / sst, 3)
        return rin, r2
    except (IndexError, ZeroDivisionError, np.linalg.LinAlgError, TypeError, ValueError):
        return None, None


def rin_from_delta(di, dv):
    """Rin [MΩ] from one step: (ΔV [mV] / ΔI [pA]) × 1000."""
    if abs(di) < CC_MIN_DELTA_I_PA:
        return None
    return round(dv / di * 1000.0, 3)


def collect_cc_iv_series(
    abf, voltage_ch, current_ch, v_start, v_end, i_index, spike_start=None, spike_end=None,
):
    """All-sweep absolute (I, V) for Rin QC plot; flags spike-free and Vm windows."""
    if spike_start is None:
        spike_start = v_start
    if spike_end is None:
        spike_end = v_end
    sweep_max = last_subthreshold_sweep(abf, voltage_ch, spike_start, spike_end)

    currents, voltages = [], []
    flags_primary, flags_fallback, flags_subth = [], [], []
    for sn in abf.sweepList:
        abf.setSweep(sweepNumber=sn, channel=voltage_ch)
        vm = float(statistics.mean(abf.sweepY[v_start:v_end]))
        abf.setSweep(sweepNumber=sn, channel=current_ch)
        i_val = float(abf.sweepY[i_index])
        voltages.append(vm)
        currents.append(i_val)
        is_sub = sweep_max is not None and sn <= sweep_max
        flags_subth.append(is_sub)
        flags_primary.append(is_sub and RIN_VMIN <= vm <= RIN_VMAX)
        flags_fallback.append(is_sub and RIN_VMIN_FALLBACK <= vm <= RIN_VMAX_FALLBACK)
    return {
        "currents": currents,
        "voltages": voltages,
        "in_primary": flags_primary,
        "in_fallback": flags_fallback,
        "subthreshold": flags_subth,
        "sweep_max_subth": sweep_max,
    }


def rin_for_channel(
    abf, voltage_ch, current_ch, pre_start, pre_end, post_start, post_end, i_index
):
    """
    Rin for CC / Gj / Cm (priority), using only sweeps before the first AP
    on voltage_ch in the post epoch [post_start:post_end]:

    1) Absolute: linear fit V vs I, mean post Vm in -80...-50 mV (>=2 points)
    2) Absolute: same fit, mean post Vm in -95...-50 mV (>=2 points)
    3) Relative: dV/dI (mean if several) for sweeps with mean Vm in -95...-50
       (>=1 point; no linear fit)
    """
    stim_label = f"stim Vm samples {post_start}-{post_end}, I at sample {i_index}"
    plot = {"mode": None, "rin_mohm": None, "r2": None, "vm_range": None}

    sweep_max = last_subthreshold_sweep(abf, voltage_ch, post_start, post_end)
    if sweep_max is None:
        return (
            None, None, 0,
            (
                f"no subthreshold sweeps (AP already on sweep 0 in "
                f"[{post_start}:{post_end}]) ({stim_label})"
            ),
            None, None, plot,
        )
    sub_note = f"sweeps 0...{sweep_max} (before first AP)"

    # --- 1) Absolute, primary Vm window (>=2 points) ---
    currents, voltages = collect_iv_points(
        abf, voltage_ch, current_ch, post_start, post_end, i_index,
        RIN_VMIN, RIN_VMAX, sweep_max=sweep_max,
    )
    n_primary = len(voltages)
    if n_primary >= RIN_MIN_POINTS:
        rin, r2 = rin_r2(currents, voltages)
        if rin is not None and rin > 0:
            plot.update(
                mode="abs_linear", rin_mohm=rin, r2=r2,
                vm_range=f"{RIN_VMIN}:{RIN_VMAX} mV",
            )
            return (
                rin, r2, n_primary, None,
                f"{RIN_VMIN}:{RIN_VMAX} mV",
                f"Rin_abs primary ({sub_note})",
                plot,
            )

    # --- 2) Absolute, extended Vm window (>=2 points) ---
    currents_fb, voltages_fb = collect_iv_points(
        abf, voltage_ch, current_ch, post_start, post_end, i_index,
        RIN_VMIN_FALLBACK, RIN_VMAX_FALLBACK, sweep_max=sweep_max,
    )
    n_fb = len(voltages_fb)
    if n_fb >= RIN_MIN_POINTS:
        rin, r2 = rin_r2(currents_fb, voltages_fb)
        if rin is not None and rin > 0:
            note = (
                f"Rin_abs extended ({RIN_VMIN_FALLBACK}...{RIN_VMAX_FALLBACK} mV, "
                f"{n_fb} points; {sub_note})"
            )
            plot.update(
                mode="abs_linear", rin_mohm=rin, r2=r2,
                vm_range=f"{RIN_VMIN_FALLBACK}:{RIN_VMAX_FALLBACK} mV",
            )
            return (
                rin, r2, n_fb, None,
                f"{RIN_VMIN_FALLBACK}:{RIN_VMAX_FALLBACK} mV", note, plot,
            )

    # --- 3) Relative: only -95...-50, >=1 point, no linear fit ---
    dis, dvs = collect_rin_delta_points(
        abf, voltage_ch, current_ch, pre_start, pre_end, post_start, post_end,
        RIN_VMIN_FALLBACK, RIN_VMAX_FALLBACK, sweep_max=sweep_max,
    )
    if len(dis) >= 1:
        ratios = [dv / di for di, dv in zip(dis, dvs)]
        rin = round(statistics.mean(ratios) * 1000.0, 3)
        if rin > 0:
            note = (
                f"Rin_rel dV/dI ({len(dis)} sweep(s), "
                f"Vm in [{RIN_VMIN_FALLBACK}, {RIN_VMAX_FALLBACK}] mV; {sub_note})"
            )
            plot.update(mode="rel_delta", rin_mohm=rin, vm_range=note)
            return (
                rin, None, len(dis), None,
                f"{RIN_VMIN_FALLBACK}:{RIN_VMAX_FALLBACK} mV", note, plot,
            )

    return (
        None, None, max(n_primary, n_fb),
        (
            f"no positive Rin_abs (>= {RIN_MIN_POINTS} points in "
            f"[{RIN_VMIN},{RIN_VMAX}] or [{RIN_VMIN_FALLBACK},{RIN_VMAX_FALLBACK}] mV) "
            f"and no positive Rin_rel (>=1 sweep in [{RIN_VMIN_FALLBACK},{RIN_VMAX_FALLBACK}] mV) "
            f"among {sub_note} ({stim_label})"
        ),
        None, None, plot,
    )


# =============================================================================
# Cell properties (from Cell Properties for one folder.ipynb)
# =============================================================================


def _stim_channel_for_voltage(voltage_ch):
    """Channel index whose sweepC carries the stimulus command for voltage_ch."""
    if voltage_ch == 0:
        return 0
    if voltage_ch == 2:
        return 1
    raise ValueError(f"unsupported voltage channel {voltage_ch}")


def channel_epochs(abf, voltage_ch):
    """Stim start/stop and analysis windows from sweepC (125 ms pre, 250 ms late stim, 500 ms post)."""
    sr = int(abf.dataRate)
    points_pre = int(0.125 * sr)
    points_in = int(0.25 * sr)
    points_after = int(0.5 * sr)

    stim_ch = _stim_channel_for_voltage(voltage_ch)
    abf.setSweep(sweepNumber=0, channel=stim_ch)
    stim_idx = np.array(np.nonzero(abf.sweepC))[0]
    if len(stim_idx) == 0:
        return None

    start = int(stim_idx[0])
    stop = int(stim_idx[-1])
    return {
        "start": start,
        "stop": stop,
        "pre_start": start - points_pre,
        "post_stop": stop + points_after,
        "within_point": stop - points_in,
    }


def cp_sweep_4aps(abf, voltage_ch, start, stop, v_level=CP_SPIKE_HEIGHT):
    for sweep_num in abf.sweepList:
        abf.setSweep(sweepNumber=sweep_num, channel=voltage_ch)
        peaks, _ = find_peaks(
            abf.sweepY[start:stop], height=v_level, distance=CP_SPIKE_DISTANCE
        )
        if len(peaks) >= CP_MIN_APS:
            return sweep_num
    return None


def cp_sweep_1ap(abf, voltage_ch, start, stop, v_level=CP_SPIKE_HEIGHT):
    """Last sweep before the first AP. None if sweep 0 already has an AP."""
    last_no_ap = None
    for sweep_num in abf.sweepList:
        abf.setSweep(sweepNumber=sweep_num, channel=voltage_ch)
        peaks, _ = find_peaks(
            abf.sweepY[start:stop], height=v_level, distance=CP_SPIKE_DISTANCE
        )
        if len(peaks) >= 1:
            return last_no_ap
        last_no_ap = sweep_num
    return last_no_ap


def cp_infl_points(abf, sweep, voltage_ch, start, stop):
    if sweep is None:
        return np.array([], dtype=int)
    abf.setSweep(sweepNumber=sweep, channel=voltage_ch)
    smooth = gaussian_filter1d(abf.sweepY[start:stop], 3)
    smooth_d2 = np.gradient(np.gradient(smooth))
    infls = np.where(np.diff(np.sign(smooth_d2)))[0]
    return infls + start


def cp_peak_indices(abf, sweep, voltage_ch, start, stop, v_level=CP_SPIKE_HEIGHT):
    if sweep is None:
        return np.array([], dtype=int)
    abf.setSweep(sweepNumber=sweep, channel=voltage_ch)
    peaks, _ = find_peaks(
        abf.sweepY[start:stop], height=v_level, distance=CP_SPIKE_DISTANCE
    )
    return peaks + start


def cp_spike_begin(ind_peaks, ind_infls):
    start_ap_ind, end_ap_ind = [], []
    if len(ind_peaks) == 0 or len(ind_infls) == 0:
        return start_ap_ind, end_ap_ind

    for peak in ind_peaks:
        infl_before = [i for i in ind_infls if i < peak]
        if len(infl_before) >= 2:
            start_ap_ind.append(infl_before[-2])
        else:
            start_ap_ind.append(np.nan)

        infl_after = [i for i in ind_infls if i > peak]
        if len(infl_after) >= 2:
            end_ap_ind.append(infl_after[1])
        else:
            end_ap_ind.append(np.nan)

    return start_ap_ind, end_ap_ind


def cp_fwhm_details(abf, sweep, voltage_ch, ap_start, ap_end, ap_peak, sampling_rate, height=0.5):
    """
    FWHM of 1st AP relative to AP start baseline (not prominence).

    half_level = V(ap_start) + height * (V(peak) - V(ap_start))
    Width = time between rising and falling crossings of half_level.
    """
    if sweep is None or not np.isfinite(ap_start) or not np.isfinite(ap_end):
        return None, None, None
    if not np.isfinite(ap_peak):
        return None, None, None

    abf.setSweep(sweepNumber=sweep, channel=voltage_ch)
    i0 = int(ap_start)
    i1 = int(ap_end)
    ip = int(ap_peak)
    y = abf.sweepY
    if i0 < 0 or i1 <= i0 + 1 or ip <= i0 or ip >= i1:
        return None, None, None

    v_base = float(y[i0])
    v_peak = float(y[ip])
    amp = v_peak - v_base
    if amp <= 0:
        return None, None, None

    half_level = v_base + height * amp

    def _cross_up(j0, j1, level):
        for i in range(j0, j1):
            y0, y1 = float(y[i]), float(y[i + 1])
            if y0 < level <= y1:
                if y1 == y0:
                    return float(i)
                return i + (level - y0) / (y1 - y0)
        return None

    def _cross_down(j0, j1, level):
        for i in range(j0, j1):
            y0, y1 = float(y[i]), float(y[i + 1])
            if y0 >= level > y1:
                if y0 == y1:
                    return float(i)
                return i + (y0 - level) / (y0 - y1)
        return None

    left = _cross_up(i0, ip, half_level)
    right = _cross_down(ip, i1 - 1, half_level)
    if left is None or right is None or right <= left:
        return None, None, None

    fwhm_ms = round(((right - left) / sampling_rate) * 1000, 3)
    left_i = int(round(left))
    right_i = int(round(right))
    return fwhm_ms, (left_i, right_i), half_level


def cp_fwhm_1st_spike(abf, sweep, voltage_ch, ap_start, ap_end, ap_peak, sampling_rate, height=0.5):
    fwhm_ms, _, _ = cp_fwhm_details(
        abf, sweep, voltage_ch, ap_start, ap_end, ap_peak, sampling_rate, height
    )
    return fwhm_ms


def cp_ap21_ratio(abf, sweep, voltage_ch, ap_start, ap_peaks):
    if sweep is None or len(ap_peaks) < 2:
        return None
    abf.setSweep(sweepNumber=sweep, channel=voltage_ch)
    amplitudes = [
        abf.sweepY[int(peak)] - abf.sweepY[int(start)]
        for peak, start in zip(ap_peaks, ap_start)
        if np.isfinite(start)
    ]
    if len(amplitudes) < 2 or amplitudes[0] == 0:
        return None
    return round(amplitudes[1] / amplitudes[0], 2)


def cp_frequencies(ap_peaks, sampling_rate):
    if len(ap_peaks) < 2:
        return None, None, None
    ap_peaks = np.asarray(ap_peaks, dtype=int)
    freq_12 = round(sampling_rate / (ap_peaks[1] - ap_peaks[0]), 2)
    freq_late = round(sampling_rate / (ap_peaks[-1] - ap_peaks[-2]), 2)
    if len(ap_peaks) >= 3:
        freq_list = [
            sampling_rate / (ap_peaks[i + 1] - ap_peaks[i])
            for i in range(1, len(ap_peaks) - 1)
        ]
        mean_freq = round(statistics.mean(freq_list), 2) if freq_list else None
    else:
        mean_freq = None
    return freq_12, freq_late, mean_freq


def cp_delay_1st_ap(stim_start, ap_peaks, sampling_rate):
    if len(ap_peaks) == 0:
        return None
    return round(((ap_peaks[0] - stim_start) / sampling_rate) * 1000, 2)


def cp_collect_iv_rin_data(abf, voltage_ch, current_ch, pre_start, pre_end, post_start, post_end):
    """I–V points and fit for cell-properties Rin_abs / Rin_rel / V_rest.

    V_rest is the linear V vs I intercept at I=0. Stages (need >=2 points):
    1) post Vm in -85...-50 mV
    2) post Vm in -95...-50 mV
    3) post Vm in -95...-40 mV
    Relative dV/dI (CC Rin stage 3) cannot give an intercept, so it is not used.
    """
    empty = {
        "currents": [],
        "voltages": [],
        "range_indices": [],
        "rin_rel": None,
        "rin_abs": None,
        "v_rest": None,
        "r2_abs": None,
        "slope": None,
        "intercept": None,
        "n_subthreshold_sweeps": 0,
        "v_rest_skip": None,
        "v_rest_note": None,
        "v_rest_stage": None,
    }
    last_sub_sweep = cp_sweep_1ap(abf, voltage_ch, pre_end, post_end)
    if last_sub_sweep is None:
        empty["v_rest_skip"] = (
            "AP already on sweep 0 in stim window; no subthreshold I–V for V_rest"
        )
        return empty

    voltages_list, currents_list = [], []
    voltages_delta_list, currents_delta_list = [], []

    for sn in range(last_sub_sweep + 1):
        abf.setSweep(sweepNumber=sn, channel=voltage_ch)
        voltage_pre = float(np.mean(abf.sweepY[pre_start:pre_end]))
        voltage_post = float(np.mean(abf.sweepY[post_start:post_end]))
        voltages_delta_list.append(voltage_post - voltage_pre)
        voltages_list.append(voltage_post)
        abf.setSweep(sweepNumber=sn, channel=current_ch)
        current_pre = float(np.mean(abf.sweepY[pre_start:pre_end]))
        current_post = float(np.mean(abf.sweepY[post_start:post_end]))
        currents_delta_list.append(current_post - current_pre)
        currents_list.append(current_post)

    rin_relative_list = [
        x / y if abs(y) > CP_RIN_REL_MIN_DI_PA else np.nan
        for x, y in zip(voltages_delta_list, currents_delta_list)
    ]

    def _in_window(vmin, vmax):
        return [
            i for i, vm in enumerate(voltages_list)
            if vmin <= vm <= vmax
        ]

    vm_min = min(voltages_list) if voltages_list else None
    vm_max = max(voltages_list) if voltages_list else None
    vm_span = (
        f"subth post Vm {vm_min:.1f}…{vm_max:.1f} mV, {len(voltages_list)} sweep(s)"
        if vm_min is not None else "no subthreshold sweeps"
    )

    # Same idea as rin_for_channel: try narrower I–V windows first.
    # V_rest needs a linear V(I) fit (intercept at I=0), so relative dV/dI
    # (Rin stage 3) cannot provide V_rest.
    stages = CP_VREST_STAGES
    stage_counts = []
    fit_idx = []
    v_rest_note = None
    v_rest_stage = None
    rin_abs = None
    v_rest = None
    r2_abs = None
    slope = None
    intercept = None
    v_rest_skip = None

    for n_stage, vmin, vmax, tag in stages:
        idx = _in_window(vmin, vmax)
        stage_counts.append(f"stage{n_stage} {vmin}…{vmax} mV: {len(idx)}")
        if len(idx) < CP_RIN_MIN_POINTS:
            continue
        currents_ar = np.array(currents_list)[idx]
        voltages_ar = np.array(voltages_list)[idx]
        try:
            fit = np.polyfit(currents_ar, voltages_ar, 1, full=True)
            slope, intercept = fit[0]
            rin_abs = round(float(slope) * 1000, 3)
            v_rest = round(float(intercept), 3)
            if not np.isfinite(rin_abs) or not np.isfinite(v_rest):
                v_rest_skip = f"I–V intercept not finite at stage {n_stage}; {vm_span}"
                rin_abs = v_rest = None
                continue
            sse = fit[1][0] if len(fit[1]) else 0.0
            sst = np.sum((voltages_ar - np.mean(voltages_ar)) ** 2)
            r2_abs = round(1 - sse / sst, 3) if sst != 0 else None
            fit_idx = idx
            v_rest_stage = tag
            v_rest_note = (
                f"V_rest stage {n_stage}/{len(stages)} ({tag}): "
                f"linear I–V in {vmin}…{vmax} mV, {len(idx)} points; {vm_span}"
            )
            v_rest_skip = None
            break
        except (IndexError, ZeroDivisionError, np.linalg.LinAlgError, TypeError, ValueError) as exc:
            v_rest_skip = f"I–V polyfit failed at stage {n_stage} ({exc}); {vm_span}"
            continue

    if v_rest is None and v_rest_skip is None:
        v_rest_skip = (
            f"need >={CP_RIN_MIN_POINTS} subthreshold I–V points for V_rest "
            f"(Rin relative dV/dI cannot give intercept); "
            f"{'; '.join(stage_counts)}; {vm_span}"
        )

    rin_idx = _in_window(CP_RIN_VMIN, CP_RIN_VMAX) or _in_window(
        CP_RIN_VMIN_FALLBACK, CP_RIN_VMAX_FALLBACK
    ) or _in_window(CP_RIN_VMIN_WIDE, CP_RIN_VMAX_WIDE)
    if rin_idx:
        rin_from_range = [
            x for x in (rin_relative_list[i] for i in rin_idx)
            if x is not None and np.isfinite(x)
        ]
        rin_rel = round(float(np.mean(rin_from_range)) * 1000, 3) if rin_from_range else None
    else:
        rin_rel = None

    return {
        "currents": currents_list,
        "voltages": voltages_list,
        "range_indices": fit_idx,
        "rin_rel": rin_rel,
        "rin_abs": rin_abs,
        "v_rest": v_rest,
        "r2_abs": r2_abs,
        "slope": slope,
        "intercept": intercept,
        "n_subthreshold_sweeps": last_sub_sweep + 1,
        "v_rest_skip": v_rest_skip,
        "v_rest_note": v_rest_note,
        "v_rest_stage": v_rest_stage,
    }


def cp_iv_rin(abf, voltage_ch, current_ch, pre_start, pre_end, post_start, post_end):
    """Rin_abs (V–I slope), Rin_rel (mean deltaV/deltaI), V_rest (intercept at I=0)."""
    data = cp_collect_iv_rin_data(
        abf, voltage_ch, current_ch, pre_start, pre_end, post_start, post_end
    )
    return data["rin_rel"], data["rin_abs"], data["v_rest"], data["r2_abs"]


def cp_iv_values(abf, sweep, voltage_ch, current_ch, pre_start, pre_end, in_start, in_end):
    """Hold potential (pre-stim Vm) and injection current on the >=4 AP sweep."""
    if sweep is None:
        return None, None
    abf.setSweep(sweepNumber=sweep, channel=voltage_ch)
    hold_v = round(float(np.mean(abf.sweepY[pre_start:pre_end])), 3)
    abf.setSweep(sweepNumber=sweep, channel=current_ch)
    current_pre = float(np.mean(abf.sweepY[pre_start:pre_end]))
    current_in = float(np.mean(abf.sweepY[in_start:in_end]))
    inj = round(current_in - current_pre, 3)
    return hold_v, inj


def compute_cell_properties(abf, voltage_ch, current_ch):
    """Per-cell firing and I–V metrics (Cell Properties notebook logic)."""
    label = f"ch{voltage_ch}"
    ch_tag = f"ch{voltage_ch}"
    epochs = channel_epochs(abf, voltage_ch)
    if epochs is None:
        return {
            f"props_skip_reason_{label}": "no stimulus detected in sweepC",
            f"V_rest_skip_reason_{ch_tag}": "no stimulus detected in sweepC",
        }

    start = epochs["start"]
    stop = epochs["stop"]
    pre_start = epochs["pre_start"]
    within = epochs["within_point"]
    sr = int(abf.dataRate)

    iv = cp_collect_iv_rin_data(
        abf, voltage_ch, current_ch, pre_start, start, within, stop
    )
    iv_fields = {
        f"Rin_abs_MOhm_{ch_tag}": iv["rin_abs"],
        f"Rin_rel_MOhm_{ch_tag}": iv["rin_rel"],
        f"V_rest_mV_{ch_tag}": iv["v_rest"],
        f"R2_abs_Rin_{ch_tag}": iv["r2_abs"],
        f"V_rest_skip_reason_{ch_tag}": iv.get("v_rest_skip"),
        f"V_rest_note_{ch_tag}": iv.get("v_rest_note"),
        f"V_rest_stage_{ch_tag}": iv.get("v_rest_stage"),
    }

    sweep_4 = cp_sweep_4aps(abf, voltage_ch, start, stop)
    if sweep_4 is None:
        return {
            **iv_fields,
            f"props_skip_reason_{label}": f"no sweep with >={CP_MIN_APS} APs in stim window",
        }

    ind_infls = cp_infl_points(abf, sweep_4, voltage_ch, start, stop)
    ind_peaks = cp_peak_indices(abf, sweep_4, voltage_ch, start, stop)
    if len(ind_peaks) < CP_MIN_APS:
        return {
            **iv_fields,
            f"props_skip_reason_{label}": (
                f"sweep {sweep_4} has {len(ind_peaks)} peaks (< {CP_MIN_APS})"
            ),
        }

    ap_start, ap_end = cp_spike_begin(ind_peaks, ind_infls)
    if not ap_start or not np.isfinite(ap_start[0]) or not np.isfinite(ap_end[0]):
        return {
            **iv_fields,
            f"props_skip_reason_{label}": "could not define 1st AP borders (inflection points)",
        }

    hold_v, inj = cp_iv_values(
        abf, sweep_4, voltage_ch, current_ch, pre_start, start, within, stop
    )
    freq_12, freq_late, mean_freq = cp_frequencies(ind_peaks, sr)

    return {
        **iv_fields,
        f"AP21_ratio_{ch_tag}": cp_ap21_ratio(abf, sweep_4, voltage_ch, ap_start, ind_peaks),
        f"init_freq_Hz_{ch_tag}": freq_12,
        f"late_freq_Hz_{ch_tag}": freq_late,
        f"mean_freq_Hz_{ch_tag}": mean_freq,
        f"FWHM_ms_{ch_tag}": cp_fwhm_1st_spike(
            abf, sweep_4, voltage_ch, ap_start[0], ap_end[0], ind_peaks[0], sr
        ),
        f"delay_AP1_ms_{ch_tag}": cp_delay_1st_ap(start, ind_peaks, sr),
        f"hold_V_mV_{ch_tag}": hold_v,
        f"inj_current_pA_{ch_tag}": inj,
        f"props_sweep_{ch_tag}": sweep_4,
        f"props_skip_reason_{ch_tag}": None,
    }


def cell_properties_for_file(abf):
    """Cell properties for both paired cells; keys suffixed _ch0 and _ch2."""
    out = empty_cell_props_fields()
    for voltage_ch, current_ch in ((0, 1), (2, 3)):
        props = compute_cell_properties(abf, voltage_ch, current_ch)
        out.update(props)
    return out


# =============================================================================
# Spikelet coupling (AP2+; first >=4, else 3, else 2; else AP1 on 1-spike sweeps)
# =============================================================================


def _ms_to_samples(ms, sr):
    return max(1, int(round(float(ms) * float(sr) / 1000.0)))


def _samples_to_ms(n_samples, sr):
    return (float(n_samples) / float(sr)) * 1000.0


def _frac_crossing_time_ms(y, i0, i1, level, sr):
    """First upward crossing of `level` in [i0, i1], linearly interpolated, ms from sample 0."""
    if y is None or sr in (None, 0):
        return None
    i0 = int(i0)
    i1 = int(i1)
    n = len(y)
    if i0 < 0 or i1 <= i0 + 1 or i0 >= n - 1:
        return None
    i1 = min(i1, n - 1)
    y = np.asarray(y, dtype=float)
    for i in range(i0, i1):
        y0, y1v = float(y[i]), float(y[i + 1])
        if y0 < level <= y1v:
            if y1v == y0:
                idx = float(i)
            else:
                idx = i + (level - y0) / (y1v - y0)
            return _samples_to_ms(idx, sr)
    return None


def _frac_rise_time_ms(y, i0, i_peak, v_base, amp, sr, frac=None):
    """Time of frac * amp above v_base on the rising phase [i0, i_peak]."""
    if frac is None:
        frac = SPIKELET_DELAY_FRAC
    if amp is None or amp <= 0 or v_base is None:
        return None
    level = float(v_base) + float(frac) * float(amp)
    return _frac_crossing_time_ms(y, i0, i_peak, level, sr)


def _spikelet_smooth_for_peak(y, sr, smooth_ms=None):
    """Light Gaussian on a short window; used only to find the peak index."""
    y = np.asarray(y, dtype=float).ravel()
    if smooth_ms is None:
        smooth_ms = SPIKELET_PEAK_SMOOTH_MS
    if y.size < 5 or sr is None or smooth_ms <= 0:
        return y
    sigma = float(smooth_ms) * float(sr) / 1000.0
    if sigma < 0.5:
        return y
    return gaussian_filter1d(y, sigma=sigma, mode="nearest")


def _spikelet_local_peak_index(seg, sr=None, prominence=None, smooth_ms=None):
    """Index of a peaked spikelet in the search window, or None.

    The full search window is scanned. Not ``argmax``.
    Digitizer jitter is first smoothed (default σ = SPIKELET_PEAK_SMOOTH_MS,
    or ``smooth_ms`` when given).
    A candidate must then be a real peak: the smoothed trace falls after
    the top by at least ``prominence`` (scipy; default
    ``SPIKELET_MIN_PROMINENCE_MV``). A monotonic rise, a slow coupling
    envelope with no peaked event, or a 1-sample wiggle therefore returns
    None.

    If several valid peaks exist, the highest smoothed one is taken.
    Amplitude is still measured on the raw trace at this index.
    """
    y_raw = np.asarray(seg, dtype=float).ravel()
    if y_raw.size < 5:
        return None
    if prominence is None:
        prominence = SPIKELET_MIN_PROMINENCE_MV
    y = _spikelet_smooth_for_peak(y_raw, sr, smooth_ms=smooth_ms)
    peaks, _props = find_peaks(y, prominence=float(prominence))
    if peaks.size == 0:
        return None
    best_i = int(peaks[int(np.argmax(y[peaks]))])
    return best_i


def _segment_rms_mad(y, i0, i1):
    """Robust noise (MAD→σ) on y[i0:i1]."""
    i0 = max(int(i0), 0)
    i1 = max(int(i1), i0 + 3)
    seg = np.asarray(y[i0:i1], dtype=float)
    if len(seg) < 3:
        return None
    med = float(np.median(seg))
    mad = float(np.median(np.abs(seg - med)))
    if mad > 0:
        return 1.4826 * mad
    return float(np.std(seg - np.mean(seg)))


def _pooled_rms_mad(chunks):
    """One MAD→σ from concatenated local-baseline windows."""
    parts = []
    for chunk in chunks or []:
        arr = np.asarray(chunk, dtype=float).ravel()
        if len(arr):
            parts.append(arr)
    if not parts:
        return None
    y = np.concatenate(parts)
    return _segment_rms_mad(y, 0, len(y))


def _passive_rms_prestim(abf, sweep, passive_ch, pre_start, stim_start):
    """Prestim noise on passive (fallback if no local 10 ms windows)."""
    abf.setSweep(sweepNumber=sweep, channel=passive_ch)
    return _segment_rms_mad(abf.sweepY, pre_start, stim_start)


def spikelet_amp_threshold(rms, k=None):
    """Amplitude bar: max(k × robust noise, SPIKELET_MIN_AMP_MV)."""
    if k is None:
        k = SPIKELET_NOISE_K
    if rms is None:
        return float(SPIKELET_MIN_AMP_MV)
    try:
        noise = float(rms)
    except (TypeError, ValueError):
        return float(SPIKELET_MIN_AMP_MV)
    if not np.isfinite(noise) or noise < 0:
        return float(SPIKELET_MIN_AMP_MV)
    return max(float(k) * noise, float(SPIKELET_MIN_AMP_MV))


def spikelet_amp_passes(amp, rms, k=None):
    """True if spikelet amplitude clears the noise gate."""
    if amp is None or amp <= 0:
        return False, "amp_spikelet_le_0"
    k = SPIKELET_NOISE_K if k is None else k
    need = spikelet_amp_threshold(rms, k=k)
    if amp < need:
        return False, (
            f"below_noise (amp={amp:.4f} < max({k}*noise, "
            f"{SPIKELET_MIN_AMP_MV})={need:.4f})"
        )
    return True, None


def select_spikelet_sweep(abf, voltage_ch, start, stop):
    """
    First sweep with >=4 APs; else 3, else 2, else 1.

    Returns (sweep, n_peaks, tier) or (None, 0, None).
    """
    for min_aps, tier in ((4, ">=4"), (3, "3"), (SPIKELET_MIN_APS, "2"), (1, "1")):
        for sweep_num in abf.sweepList:
            abf.setSweep(sweepNumber=sweep_num, channel=voltage_ch)
            peaks, _ = find_peaks(
                abf.sweepY[start:stop], height=CP_SPIKE_HEIGHT, distance=CP_SPIKE_DISTANCE
            )
            if len(peaks) >= min_aps:
                return sweep_num, int(len(peaks)), tier
    return None, 0, None


def list_spikelet_sweeps(abf, voltage_ch, start, stop):
    """
    Sweeps with >=2 APs from the first such sweep onward.

    If the file has no multi-spike sweep, use 1-AP sweeps from the first
    of those onward (AP1 is then measured). Later sweeps below the
    chosen minimum are skipped.
    """
    counted = []
    for sweep_num in abf.sweepList:
        abf.setSweep(sweepNumber=sweep_num, channel=voltage_ch)
        peaks, _ = find_peaks(
            abf.sweepY[start:stop], height=CP_SPIKE_HEIGHT, distance=CP_SPIKE_DISTANCE
        )
        counted.append((int(sweep_num), int(len(peaks))))
    min_aps = SPIKELET_MIN_APS if any(n >= SPIKELET_MIN_APS for _, n in counted) else 1
    started = False
    out = []
    for sweep_num, n in counted:
        if n >= min_aps:
            started = True
            out.append(sweep_num)
        elif started:
            continue
    return out


def _spikelet_ap_indices(n_ap):
    """AP indices to measure: AP2+ normally, AP1 only if the sweep has one spike."""
    n_ap = int(n_ap or 0)
    if n_ap < 1:
        return range(0)
    if n_ap == 1:
        return range(0, 1)
    return range(1, n_ap)


def _empty_spikelet_metrics(skip_reason):
    return {
        "sweep": None,
        "n_AP_active": 0,
        "sweep_tier": None,
        "n_AP_used": 0,
        "n_spikelet_detected": 0,
        "n_with_amp": 0,
        "mean_amp_active_mV": None,
        "mean_amp_spikelet_mV": None,
        "mean_amp_ratio": None,
        "mean_delay_ms": None,
        "mean_delay_10_ms": None,
        "n_delay_negative": 0,
        "n_delay_10_negative": 0,
        "n_ratio_gt_1": 0,
        "avg_amp_active_mV": None,
        "avg_amp_spikelet_mV": None,
        "avg_amp_ratio": None,
        "avg_delay_ms": None,
        "avg_delay_10_ms": None,
        "meantrace10_amp_active_mV": None,
        "meantrace10_amp_spikelet_mV": None,
        "meantrace10_amp_ratio": None,
        "meantrace10_delay_ms": None,
        "meantrace10_delay_10_ms": None,
        "meantrace10_detected": False,
        "meantrace10_skip_reason": None,
        "metric_source": None,
        "ap1_fallback": False,
        "rms_noise_mV": None,
        "noise_thr_mV": None,
        "skip_reason": skip_reason,
        "n_sweeps": 0,
        "mean_vm_begin_mV": None,
        "mean_baseline_passive_mV": None,
    }


def _round_or_none(val, nd=4):
    if val is None:
        return None
    try:
        if not np.isfinite(val):
            return None
    except TypeError:
        return None
    return round(float(val), nd)


def _ratio_gt_one(val):
    v = _finite_number(val)
    return bool(v is not None and v > 1.0)


def _plottable_amp_ratio(val):
    """Ratio for folder plots: drop unphysical spikelet/spike > 1."""
    v = _finite_number(val)
    if v is None or v > 1.0:
        return None
    return v


def _mark_spikelet_quality(row):
    """Flag unphysical delay (<0) and ratio (>1). Values stay in Excel."""
    d = _finite_number(row.get("delay_ms"))
    d10 = _finite_number(row.get("delay_10_ms"))
    ratio = _finite_number(row.get("amp_ratio"))
    row["delay_negative"] = bool(d is not None and d < 0)
    row["delay_10_negative"] = bool(d10 is not None and d10 < 0)
    row["ratio_gt_1"] = _ratio_gt_one(ratio)
    notes = []
    if row["delay_negative"]:
        notes.append(f"delay_peak<0 ({d} ms): spikelet peak before AP peak")
    if row["delay_10_negative"]:
        notes.append(f"delay_10<0 ({d10} ms): spikelet 10% before AP 10%")
    if row["ratio_gt_1"]:
        notes.append(f"ratio>1 ({ratio}): spikelet amp > AP amp")
    if notes:
        extra = "; ".join(notes)
        prev = row.get("skip_reason")
        row["skip_reason"] = f"{prev}; {extra}" if prev else extra
    return row


def _best_effort_spikelet_sweep(abf, voltage_ch, start, stop):
    """Sweep with the most APs in the stim window (may be 0 or 1)."""
    best_sn, best_n = None, -1
    for sn in abf.sweepList:
        abf.setSweep(sweepNumber=sn, channel=voltage_ch)
        peaks, _ = find_peaks(
            abf.sweepY[start:stop], height=CP_SPIKE_HEIGHT, distance=CP_SPIKE_DISTANCE
        )
        n = int(len(peaks))
        if n > best_n:
            best_n, best_sn = n, int(sn)
    if best_sn is None and len(abf.sweepList):
        best_sn, best_n = int(abf.sweepList[0]), 0
    return best_sn, best_n


def _fallback_spikelet_qc_meta(abf, active_ch, passive_ch, direction, reason, sweep=None):
    """QC overlay even when spikelet analysis cannot finish."""
    sr = int(abf.dataRate)
    n_pre = _ms_to_samples(SPIKELET_BASELINE_MS, sr)
    n_post = _ms_to_samples(SPIKELET_PEAK_MS, sr)
    epochs = channel_epochs(abf, active_ch)
    if epochs is None:
        stim_start, stim_stop, pre_start = 0, 1, 0
        n_ap = 0
        ind_peaks = np.array([], dtype=int)
        ap_starts = np.array([], dtype=float)
        if sweep is None:
            sweep = int(abf.sweepList[0]) if len(abf.sweepList) else 0
    else:
        stim_start, stim_stop = epochs["start"], epochs["stop"]
        pre_start = epochs["pre_start"]
        if sweep is None:
            sweep, n_ap = _best_effort_spikelet_sweep(
                abf, active_ch, stim_start, stim_stop
            )
        else:
            abf.setSweep(sweepNumber=sweep, channel=active_ch)
            peaks, _ = find_peaks(
                abf.sweepY[stim_start:stim_stop],
                height=CP_SPIKE_HEIGHT, distance=CP_SPIKE_DISTANCE,
            )
            n_ap = int(len(peaks))
        ind_infls = cp_infl_points(abf, sweep, active_ch, stim_start, stim_stop)
        ind_peaks = cp_peak_indices(abf, sweep, active_ch, stim_start, stim_stop)
        ap_starts, _ = cp_spike_begin(ind_peaks, ind_infls)
        n_ap = int(len(ind_peaks))
    metrics = _empty_spikelet_metrics(reason)
    metrics["sweep"] = sweep
    metrics["n_AP_active"] = n_ap
    return {
        "direction": direction,
        "sweep": sweep if sweep is not None else 0,
        "active_ch": active_ch,
        "passive_ch": passive_ch,
        "stim_start": stim_start,
        "stim_stop": stim_stop,
        "pre_start": pre_start,
        "sr": sr,
        "n_pre": n_pre,
        "n_post": n_post,
        "ind_peaks": ind_peaks,
        "ap_starts": ap_starts,
        "ap_rows": [],
        "mean_a": None,
        "mean_p": None,
        "snips_a": [],
        "snips_p": [],
        "snips_ok": [],
        "tier": None,
        "n_ap": n_ap,
        "rms": None,
        "noise_thr_mV": None,
        "metrics": metrics,
        "t10_a_rel_ms": None,
        "t10_p_rel_ms": None,
        "error": reason,
    }


def _finite_number(val):
    if val is None:
        return None
    try:
        fv = float(val)
    except (TypeError, ValueError):
        return None
    if not np.isfinite(fv):
        return None
    return fv


def _mean_row_field(rows, key, alt_key=None):
    """Mean of a numeric field over rows that have it (AP1 already excluded upstream)."""
    vals = []
    for row in rows or []:
        v = _finite_number(row.get(key))
        if v is None and alt_key is not None:
            v = _finite_number(row.get(alt_key))
        if v is not None:
            vals.append(v)
    return _round_or_none(statistics.mean(vals), 4) if vals else None


def _rows_with_spikelet_amp(ap_rows):
    """AP2+ rows that have a measured spikelet peak (noise gate not required)."""
    out = []
    for row in ap_rows or []:
        if _finite_number(row.get("amp_spikelet_mV")) is not None:
            out.append(row)
    return out


def _fill_sweep_means_from_ap_rows(metrics, ap_rows):
    """
    Sweep means for folder plots: AP1 excluded; average every AP2+ with a
    true local-max peak. Noise gate is off (``SPIKELET_USE_NOISE_GATE``).
    """
    measured = _rows_with_spikelet_amp(ap_rows)
    passed = [r for r in measured if r.get("detected")]
    metrics["n_AP_used"] = len(ap_rows or [])
    metrics["n_with_amp"] = len(measured)
    used = passed or measured
    if not used:
        return False
    ratio_ok = [r for r in used if not r.get("ratio_gt_1")]
    metrics["mean_amp_active_mV"] = _mean_row_field(used, "amp_active_mV")
    metrics["mean_amp_spikelet_mV"] = _mean_row_field(used, "amp_spikelet_mV")
    metrics["mean_amp_ratio"] = _mean_row_field(ratio_ok, "amp_ratio")
    metrics["mean_delay_ms"] = _mean_row_field(used, "delay_ms", "delay_peak_ms")
    metrics["mean_delay_10_ms"] = _mean_row_field(used, "delay_10_ms")
    metrics["mean_vm_begin_mV"] = _mean_row_field(used, "vm_begin_active_mV")
    metrics["mean_baseline_passive_mV"] = _mean_row_field(
        used, "baseline_passive_mV"
    )
    metrics["metric_source"] = "ap2_noise_gate" if passed else "ap2_with_amp"
    metrics["skip_reason"] = None
    metrics["n_delay_negative"] = sum(1 for r in used if r.get("delay_negative"))
    metrics["n_delay_10_negative"] = sum(1 for r in used if r.get("delay_10_negative"))
    metrics["n_ratio_gt_1"] = sum(1 for r in used if r.get("ratio_gt_1"))
    return True


def _compute_meantrace10_metrics(mean_a, mean_p, n_pre, sr, rms_unified=None):
    """Parallel scheme on mean traces: 15 ms passive window after t=0.

    Detection = local maximum with post-peak decline
    (``SPIKELET_MEANTRACE_MIN_PROMINENCE_MV`` on smoothed trace).
    No noise-amplitude gate.
    """
    out = {
        "meantrace10_amp_active_mV": None,
        "meantrace10_amp_spikelet_mV": None,
        "meantrace10_amp_ratio": None,
        "meantrace10_delay_ms": None,
        "meantrace10_delay_10_ms": None,
        "meantrace10_detected": False,
        "meantrace10_skip_reason": None,
    }
    if mean_a is None or mean_p is None or sr in (None, 0):
        out["meantrace10_skip_reason"] = "no mean traces"
        return out
    n_peak = _ms_to_samples(SPIKELET_MEANTRACE_PEAK_MS, sr)
    i1 = min(len(mean_p), n_pre + n_peak + 1)
    if i1 <= n_pre + 2 or len(mean_a) < i1:
        out["meantrace10_skip_reason"] = "meantrace10 window too short"
        return out
    mp = np.asarray(mean_p[:i1], dtype=float)
    ma = np.asarray(mean_a[:i1], dtype=float)
    mp_smooth = _spikelet_smooth_for_peak(mp, sr)
    seg_p = mp[n_pre:i1]
    seg_a = ma[n_pre:i1]
    base_p = float(np.mean(mp_smooth[:n_pre])) if n_pre > 0 else float(mp_smooth[0])
    base_a = float(ma[n_pre])
    i_rel_p = _spikelet_local_peak_index(
        seg_p, sr, prominence=SPIKELET_MEANTRACE_MIN_PROMINENCE_MV,
    )
    i_rel_a = _spikelet_local_peak_index(seg_a, sr)
    if i_rel_a is None:
        i_rel_a = int(np.argmax(seg_a)) if len(seg_a) else None
    if i_rel_p is None:
        out["meantrace10_skip_reason"] = "no local peak on meantrace10"
        return out
    if i_rel_a is None:
        out["meantrace10_skip_reason"] = "no AP peak on meantrace10"
        return out
    amp_p = float(mp_smooth[n_pre + int(i_rel_p)]) - base_p
    amp_a = float(seg_a[int(i_rel_a)]) - base_a
    out["meantrace10_amp_active_mV"] = _round_or_none(amp_a, 4)
    out["meantrace10_amp_spikelet_mV"] = _round_or_none(amp_p, 4)
    out["meantrace10_detected"] = True
    out["meantrace10_skip_reason"] = None
    if amp_a not in (None, 0):
        out["meantrace10_amp_ratio"] = _round_or_none(amp_p / amp_a, 4)
    out["meantrace10_delay_ms"] = _round_or_none(
        _samples_to_ms(int(i_rel_p) - int(i_rel_a), sr), 4
    )
    i_peak_a = n_pre + int(i_rel_a)
    i_peak_p = n_pre + int(i_rel_p)
    t10_a = _frac_rise_time_ms(ma, n_pre, i_peak_a, base_a, amp_a, sr)
    t10_p = _frac_rise_time_ms(mp_smooth, n_pre, i_peak_p, base_p, amp_p, sr)
    if t10_a is not None and t10_p is not None:
        out["meantrace10_delay_10_ms"] = _round_or_none(t10_p - t10_a, 4)
    return out


def analyze_spikelets_direction(abf, active_ch, passive_ch, direction):
    """
    Spikelet metrics on every sweep with >=2 APs (from the first such sweep on).

    Sweep means: AP1 excluded except on 1-AP fallback sweeps; mean over
    measured APs (gate is off). Noise is still measured for the QC title but does not
    reject peaks. Peak = interior local max with prominence in 15 ms
    (not window argmax: no peak → no spikelet).
    QC PNG is the primary sweep only: ``{stem}_12_spikelets.png`` and
    ``{stem}_21_spikelets.png`` in the Spikelet_plots folder.

    AP1 excluded. Baseline = mean passive in 1 ms before AP start.
    Peak search = Gaussian-smoothed 15 ms window
    (σ = SPIKELET_PEAK_SMOOTH_MS) then ``find_peaks`` with prominence
    ``SPIKELET_MIN_PROMINENCE_MV``; amplitude still from the raw trace.
    delay_ms / delay_peak_ms = t_peak_passive - t_peak_active.
    delay_10_ms = t_10_spikelet - t_10_active (10% of each event's own amplitude).
    Vm per sweep for vs-Vm plots = mean spikelet baseline (passive, 1 ms
    before t=0) over those same AP2+.
    """
    epochs = channel_epochs(abf, active_ch)
    sr = int(abf.dataRate)
    n_pre = _ms_to_samples(SPIKELET_BASELINE_MS, sr)
    n_post = _ms_to_samples(SPIKELET_PEAK_MS, sr)
    if epochs is None:
        reason = "no stimulus detected in sweepC"
        meta = _fallback_spikelet_qc_meta(abf, active_ch, passive_ch, direction, reason)
        return [], _empty_spikelet_metrics(reason), meta, []

    stim_start, stim_stop = epochs["start"], epochs["stop"]
    pre_start = epochs["pre_start"]

    sweep_nums = list_spikelet_sweeps(abf, active_ch, stim_start, stim_stop)
    if not sweep_nums:
        sn, n_found = _best_effort_spikelet_sweep(abf, active_ch, stim_start, stim_stop)
        reason = (
            f"no sweep with APs in stim window "
            f"(best sweep {sn} has {n_found} AP)"
        )
        meta = _fallback_spikelet_qc_meta(
            abf, active_ch, passive_ch, direction, reason, sweep=sn,
        )
        metrics = _empty_spikelet_metrics(reason)
        metrics["sweep"] = sn
        metrics["n_AP_active"] = n_found
        return [], metrics, meta, []

    primary_sn, _, _ = select_spikelet_sweep(abf, active_ch, stim_start, stim_stop)
    all_ap_rows = []
    sweep_metrics = []
    qc_meta = None
    for sn in sweep_nums:
        ap_rows, metrics, meta = _analyze_spikelets_sweep(
            abf, sn, active_ch, passive_ch, direction,
            stim_start, stim_stop, pre_start, sr, n_pre, n_post,
        )
        all_ap_rows.extend(ap_rows)
        sweep_metrics.append(metrics)
        if meta is not None and (qc_meta is None or sn == primary_sn):
            if sn == primary_sn:
                qc_meta = meta
            elif qc_meta is None:
                qc_meta = meta

    file_metrics = _aggregate_spikelet_sweep_metrics(sweep_metrics, primary_sn)
    return all_ap_rows, file_metrics, qc_meta, sweep_metrics


def _analyze_spikelets_sweep(
    abf, sweep, active_ch, passive_ch, direction,
    stim_start, stim_stop, pre_start, sr, n_pre, n_post,
):
    """Spikelet metrics for one sweep. AP1 excluded unless this sweep has only one AP."""
    ind_infls = cp_infl_points(abf, sweep, active_ch, stim_start, stim_stop)
    ind_peaks = cp_peak_indices(abf, sweep, active_ch, stim_start, stim_stop)
    ap_starts, _ap_ends = cp_spike_begin(ind_peaks, ind_infls)
    n_ap = int(len(ind_peaks))
    use_ap1 = n_ap == 1
    if n_ap < 1:
        reason = f"sweep {sweep} has 0 peaks"
        metrics = _empty_spikelet_metrics(reason)
        metrics["sweep"] = sweep
        metrics["n_AP_active"] = n_ap
        plot_meta = {
            "direction": direction,
            "sweep": sweep,
            "active_ch": active_ch,
            "passive_ch": passive_ch,
            "stim_start": stim_start,
            "stim_stop": stim_stop,
            "pre_start": pre_start,
            "sr": sr,
            "n_pre": n_pre,
            "n_post": n_post,
            "ind_peaks": ind_peaks,
            "ap_starts": ap_starts,
            "ap_rows": [],
            "mean_a": None,
            "mean_p": None,
            "snips_a": [],
            "snips_p": [],
            "snips_ok": [],
            "tier": None,
            "n_ap": n_ap,
            "rms": None,
            "noise_thr_mV": None,
            "metrics": metrics,
            "t10_a_rel_ms": None,
            "t10_p_rel_ms": None,
            "error": reason,
        }
        return [], metrics, plot_meta
    tier = "1" if use_ap1 else (">=4" if n_ap >= 4 else str(n_ap))

    rms_pre = _passive_rms_prestim(abf, sweep, passive_ch, pre_start, stim_start)
    n_local = _ms_to_samples(SPIKELET_NOISE_LOCAL_MS, sr)

    abf.setSweep(sweepNumber=sweep, channel=active_ch)
    y_a = np.asarray(abf.sweepY, dtype=float)
    abf.setSweep(sweepNumber=sweep, channel=passive_ch)
    y_p = np.asarray(abf.sweepY, dtype=float)
    n_y = len(y_a)

    ap_rows = []
    snips_a, snips_p, snips_ok = [], [], []
    amps_a_for_avg = []
    local_chunks = []

    for i in _spikelet_ap_indices(n_ap):
        i_peak_a = int(ind_peaks[i])
        i_start = ap_starts[i] if i < len(ap_starts) else np.nan
        row = {
            "ap_index": i + 1,
            "sweep": sweep,
            "t_start_ms": None,
            "t_peak_active_ms": _round_or_none(_samples_to_ms(i_peak_a, sr), 4),
            "t_peak_passive_ms": None,
            "t_10_active_ms": None,
            "t_10_spikelet_ms": None,
            "amp_active_mV": None,
            "vm_begin_active_mV": None,
            "baseline_passive_mV": None,
            "amp_spikelet_mV": None,
            "amp_ratio": None,
            "delay_ms": None,
            "delay_peak_ms": None,
            "delay_10_ms": None,
            "delay_negative": False,
            "delay_10_negative": False,
            "ratio_gt_1": False,
            "rms_local_mV": None,
            "noise_thr_mV": None,
            "i_peak_passive": None,
            "detected": False,
            "used_in_average": False,
            "skip_reason": None,
            "metric_source": "individual",
        }

        if not np.isfinite(i_start):
            row["skip_reason"] = "no_ap_start"
            ap_rows.append(row)
            continue
        i_start = int(i_start)
        row["t_start_ms"] = _round_or_none(_samples_to_ms(i_start, sr), 4)

        if i_start < n_pre or i_start + n_post >= n_y:
            row["skip_reason"] = "window_out_of_trace"
            ap_rows.append(row)
            continue

        isi_ok = True
        if i + 1 < n_ap:
            next_start = ap_starts[i + 1] if i + 1 < len(ap_starts) else np.nan
            if np.isfinite(next_start) and int(next_start) < i_start + n_post:
                isi_ok = False
        prev_start = ap_starts[i - 1] if i >= 1 and i - 1 < len(ap_starts) else np.nan
        if np.isfinite(prev_start) and (i_start - int(prev_start)) < (n_pre + n_post):
            isi_ok = False

        v_base_a = float(y_a[i_start])
        v_peak_a = float(y_a[i_peak_a])
        amp_a = v_peak_a - v_base_a
        row["amp_active_mV"] = _round_or_none(amp_a, 4)
        row["vm_begin_active_mV"] = _round_or_none(v_base_a, 4)
        if amp_a <= 0:
            row["skip_reason"] = "amp_active_le_0"
            ap_rows.append(row)
            continue
        row["t_10_active_ms"] = _round_or_none(
            _frac_rise_time_ms(y_a, i_start, i_peak_a, v_base_a, amp_a, sr), 4
        )

        baseline = float(np.mean(y_p[i_start - n_pre:i_start]))
        row["baseline_passive_mV"] = _round_or_none(baseline, 4)
        chunk = y_p[max(0, i_start - n_local):i_start]
        if len(chunk) >= 3:
            local_chunks.append(np.asarray(chunk, dtype=float))
        seg_p = y_p[i_start:i_start + n_post + 1]
        i_rel = _spikelet_local_peak_index(seg_p, sr)
        if i_rel is None:
            row["skip_reason"] = "no_local_peak"
        else:
            i_peak_p = i_start + i_rel
            amp_p = float(y_p[i_peak_p]) - baseline
            row["i_peak_passive"] = int(i_peak_p)
            row["t_peak_passive_ms"] = _round_or_none(_samples_to_ms(i_peak_p, sr), 4)
            row["amp_spikelet_mV"] = _round_or_none(amp_p, 4)
            d_peak = _round_or_none(_samples_to_ms(i_peak_p - i_peak_a, sr), 4)
            row["delay_ms"] = d_peak
            row["delay_peak_ms"] = d_peak
            row["t_10_spikelet_ms"] = _round_or_none(
                _frac_rise_time_ms(y_p, i_start, i_peak_p, baseline, amp_p, sr), 4
            )
            if row["t_10_active_ms"] is not None and row["t_10_spikelet_ms"] is not None:
                row["delay_10_ms"] = _round_or_none(
                    row["t_10_spikelet_ms"] - row["t_10_active_ms"], 4
                )
            if amp_a != 0:
                row["amp_ratio"] = _round_or_none(amp_p / amp_a, 4)
            _mark_spikelet_quality(row)

        if isi_ok:
            snips_a.append(y_a[i_start - n_pre:i_start + n_post])
            snips_p.append(y_p[i_start - n_pre:i_start + n_post])
            snips_ok.append(bool(row["detected"]))
            amps_a_for_avg.append(amp_a)
            row["used_in_average"] = True
        else:
            if row["skip_reason"] is None and not row["detected"]:
                row["skip_reason"] = "isi_too_short"
            elif not row["detected"] and row["skip_reason"]:
                pass
            elif row["detected"] and not isi_ok:
                row["used_in_average"] = False

        ap_rows.append(row)

    rms_unified = _pooled_rms_mad(local_chunks)
    if rms_unified is None:
        rms_unified = rms_pre
    thr_unified = spikelet_amp_threshold(rms_unified)
    rms_u = _round_or_none(rms_unified, 4)
    thr_u = _round_or_none(thr_unified, 4)
    snip_i = 0
    for row in ap_rows:
        row["rms_local_mV"] = rms_u
        row["noise_thr_mV"] = thr_u
        amp = row.get("amp_spikelet_mV")
        if amp is not None:
            if SPIKELET_USE_NOISE_GATE:
                ok, why = spikelet_amp_passes(amp, rms_unified)
            else:
                ok, why = True, None
            if ok:
                row["detected"] = True
                if row.get("skip_reason") in (None, "isi_too_short"):
                    row["skip_reason"] = None
            else:
                row["detected"] = False
                row["skip_reason"] = why
        if row.get("used_in_average"):
            if snip_i < len(snips_ok):
                snips_ok[snip_i] = bool(row["detected"])
            snip_i += 1

    detected = [r for r in ap_rows if r["detected"]]
    metrics = _empty_spikelet_metrics(None)
    metrics.update({
        "sweep": sweep,
        "n_AP_active": n_ap,
        "sweep_tier": tier,
        "n_AP_used": len(ap_rows),
        "n_spikelet_detected": len(detected),
        "n_with_amp": 0,
        "rms_noise_mV": rms_u,
        "noise_thr_mV": thr_u,
        "ap1_fallback": use_ap1,
    })
    has_means = _fill_sweep_means_from_ap_rows(metrics, ap_rows)
    if use_ap1 and has_means:
        metrics["metric_source"] = "ap1_fallback"

    mean_a = mean_p = None
    avg_detected = False
    plot_t10_a_rel = None
    plot_t10_p_rel = None
    if snips_p:
        arr_p = np.vstack(snips_p)
        arr_a = np.vstack(snips_a)
        mean_p = np.mean(arr_p, axis=0)
        mean_a = np.mean(arr_a, axis=0)
        base_avg = float(np.mean(mean_p[:n_pre]))
        seg_avg = mean_p[n_pre:]
        i_rel_p = _spikelet_local_peak_index(seg_avg, sr)
        i_rel_a = _spikelet_local_peak_index(mean_a[n_pre:], sr)
        amp_a_avg = (
            _round_or_none(statistics.mean(amps_a_for_avg), 4) if amps_a_for_avg else None
        )
        metrics["avg_amp_active_mV"] = amp_a_avg
        if i_rel_p is not None:
            amp_p_avg = float(seg_avg[i_rel_p]) - base_avg
            metrics["avg_amp_spikelet_mV"] = _round_or_none(amp_p_avg, 4)
            if amp_a_avg not in (None, 0):
                metrics["avg_amp_ratio"] = _round_or_none(amp_p_avg / amp_a_avg, 4)
            if i_rel_a is not None:
                metrics["avg_delay_ms"] = _round_or_none(
                    _samples_to_ms(i_rel_p - i_rel_a, sr), 4
                )
            i_peak_a_avg = n_pre + int(i_rel_a) if i_rel_a is not None else None
            i_peak_p_avg = n_pre + int(i_rel_p)
            v_base_a_avg = float(mean_a[n_pre]) if mean_a is not None else None
            amp_a_wave = None
            if i_peak_a_avg is not None and v_base_a_avg is not None:
                amp_a_wave = float(mean_a[i_peak_a_avg]) - v_base_a_avg
            t10_a = _frac_rise_time_ms(
                mean_a, n_pre, i_peak_a_avg, v_base_a_avg, amp_a_wave, sr
            ) if i_peak_a_avg is not None else None
            t10_p = _frac_rise_time_ms(
                mean_p, n_pre, i_peak_p_avg, base_avg, amp_p_avg, sr
            )
            if t10_a is not None and t10_p is not None:
                metrics["avg_delay_10_ms"] = _round_or_none(t10_p - t10_a, 4)
            plot_t10_a_rel = None if t10_a is None else _round_or_none(
                t10_a - _samples_to_ms(n_pre, sr), 4
            )
            plot_t10_p_rel = None if t10_p is None else _round_or_none(
                t10_p - _samples_to_ms(n_pre, sr), 4
            )
            rms_avg = rms_unified
            if SPIKELET_USE_NOISE_GATE:
                ok, _ = spikelet_amp_passes(amp_p_avg, rms_avg)
            else:
                ok = True
            avg_detected = ok
        metrics.update(
            _compute_meantrace10_metrics(mean_a, mean_p, n_pre, sr, rms_unified)
        )
        if not has_means:
            if avg_detected:
                metrics["metric_source"] = "average"
                metrics["skip_reason"] = None
                for r in ap_rows:
                    if r["used_in_average"]:
                        r["metric_source"] = "average"
                if metrics.get("mean_amp_ratio") is None and metrics.get("avg_amp_ratio") is not None:
                    if not _ratio_gt_one(metrics.get("avg_amp_ratio")):
                        metrics["mean_amp_ratio"] = metrics.get("avg_amp_ratio")
                    if metrics.get("mean_delay_ms") is None:
                        metrics["mean_delay_ms"] = metrics.get("avg_delay_ms")
                    if metrics.get("mean_delay_10_ms") is None:
                        metrics["mean_delay_10_ms"] = metrics.get("avg_delay_10_ms")
                    if metrics.get("mean_amp_active_mV") is None:
                        metrics["mean_amp_active_mV"] = metrics.get("avg_amp_active_mV")
                    if metrics.get("mean_amp_spikelet_mV") is None:
                        metrics["mean_amp_spikelet_mV"] = metrics.get("avg_amp_spikelet_mV")
            else:
                metrics["metric_source"] = None
                metrics["skip_reason"] = "no spikelet on AP2+ (individual or average)"
    else:
        if not has_means:
            metrics["skip_reason"] = metrics["skip_reason"] or "no AP2+ windows for average"
        mean_a = mean_p = None

    if metrics.get("mean_vm_begin_mV") is None:
        vm0 = [r.get("vm_begin_active_mV") for r in ap_rows if r.get("vm_begin_active_mV") is not None]
        metrics["mean_vm_begin_mV"] = _round_or_none(statistics.mean(vm0), 4) if vm0 else None
    if metrics.get("mean_baseline_passive_mV") is None:
        bsl = [r.get("baseline_passive_mV") for r in ap_rows if r.get("baseline_passive_mV") is not None]
        metrics["mean_baseline_passive_mV"] = (
            _round_or_none(statistics.mean(bsl), 4) if bsl else None
        )

    plot_meta = {
        "direction": direction,
        "sweep": sweep,
        "active_ch": active_ch,
        "passive_ch": passive_ch,
        "stim_start": stim_start,
        "stim_stop": stim_stop,
        "pre_start": pre_start,
        "sr": sr,
        "n_pre": n_pre,
        "n_post": n_post,
        "ind_peaks": ind_peaks,
        "ap_starts": ap_starts,
        "ap_rows": ap_rows,
        "mean_a": mean_a,
        "mean_p": mean_p,
        "snips_a": snips_a,
        "snips_p": snips_p,
        "snips_ok": snips_ok,
        "tier": tier,
        "n_ap": n_ap,
        "rms": rms_unified,
        "noise_thr_mV": thr_u,
        "metrics": metrics,
        "t10_a_rel_ms": plot_t10_a_rel,
        "t10_p_rel_ms": plot_t10_p_rel,
        "error": metrics.get("skip_reason"),
    }
    return ap_rows, metrics, plot_meta


def _aggregate_spikelet_sweep_metrics(sweep_metrics, primary_sn):
    """File-level spikelet fields: means of the primary QC sweep (for vs-time)."""
    out = _empty_spikelet_metrics(None)
    n_all = len(sweep_metrics)
    out["n_sweeps"] = n_all
    out["sweep"] = primary_sn
    if not sweep_metrics:
        out["skip_reason"] = "no sweep with APs in stim window"
        return out
    primary = None
    for m in sweep_metrics:
        if m.get("sweep") == primary_sn:
            primary = m
            break
    if primary is None:
        primary = sweep_metrics[0]
        out["sweep"] = primary.get("sweep")
    for key in (
        "sweep_tier",
        "n_AP_active",
        "rms_noise_mV",
        "mean_amp_active_mV",
        "mean_amp_spikelet_mV",
        "mean_amp_ratio",
        "mean_delay_ms",
        "mean_delay_10_ms",
        "mean_vm_begin_mV",
        "mean_baseline_passive_mV",
        "noise_thr_mV",
        "avg_amp_active_mV",
        "avg_amp_spikelet_mV",
        "avg_amp_ratio",
        "avg_delay_ms",
        "avg_delay_10_ms",
        "meantrace10_amp_active_mV",
        "meantrace10_amp_spikelet_mV",
        "meantrace10_amp_ratio",
        "meantrace10_delay_ms",
        "meantrace10_delay_10_ms",
        "meantrace10_detected",
        "meantrace10_skip_reason",
        "metric_source",
        "skip_reason",
        "n_with_amp",
        "n_AP_used",
        "n_delay_negative",
        "n_delay_10_negative",
        "n_ratio_gt_1",
        "ap1_fallback",
    ):
        out[key] = primary.get(key)
    out["n_spikelet_detected"] = int(
        sum(m.get("n_spikelet_detected") or 0 for m in sweep_metrics)
    )
    if out.get("mean_amp_ratio") is not None or (out.get("n_with_amp") or 0) > 0:
        out["metric_source"] = out.get("metric_source") or "primary_sweep"
        out["skip_reason"] = None
    else:
        out["metric_source"] = None
        reasons = [m.get("skip_reason") for m in sweep_metrics if m.get("skip_reason")]
        out["skip_reason"] = reasons[0] if reasons else "no spikelet amplitudes in file"
    return out


def _spikelet_row(row_dict):
    return {k: row_dict.get(k) for k in SPIKELET_AP_KEYS}


def _meantrace10_qc_panel(
    ax, mean_a, mean_p, n_pre, sr, metrics, rms=None, snips_p=None, snips_a=None,
):
    """QC subplot for parallel meantrace10 scheme (15 ms passive window after t=0)."""
    mt = metrics or {}
    title_base = (
        f"Meantrace10 ({SPIKELET_MEANTRACE_PEAK_MS:g} ms, "
        f"local peak, drop>={SPIKELET_MEANTRACE_MIN_PROMINENCE_MV:g} mV)"
    )
    ax.set_xlabel("Time from active AP start (ms)")
    ax.set_ylabel("Passive Vm (mV)", color="C1")
    ax.tick_params(axis="y", labelcolor="C1")

    if mean_a is None or mean_p is None or sr in (None, 0):
        reason = mt.get("meantrace10_skip_reason") or "no mean traces"
        ax.text(
            0.5, 0.5, reason,
            transform=ax.transAxes, ha="center", va="center",
            color="C3", fontsize=10,
        )
        ax.set_title(f"{title_base}: no aligned mean")
        return

    n_peak = _ms_to_samples(SPIKELET_MEANTRACE_PEAK_MS, sr)
    i1 = min(len(mean_p), n_pre + n_peak + 1)
    if i1 <= n_pre + 2 or len(mean_a) < i1:
        reason = mt.get("meantrace10_skip_reason") or "meantrace10 window too short"
        ax.text(
            0.5, 0.5, reason,
            transform=ax.transAxes, ha="center", va="center",
            color="C3", fontsize=10,
        )
        ax.set_title(title_base)
        return

    t_rel = (np.arange(i1) - n_pre) / float(sr) * 1000.0
    mp = np.asarray(mean_p[:i1], dtype=float)
    ma = np.asarray(mean_a[:i1], dtype=float)
    seg_p = mp[n_pre:i1]
    seg_a = ma[n_pre:i1]
    mp_smooth = _spikelet_smooth_for_peak(mp, sr)

    n_snips = 0
    labeled_snip = False
    for sn in snips_p or []:
        if sn is None or len(sn) < i1:
            continue
        ax.plot(
            t_rel, np.asarray(sn[:i1], dtype=float),
            color="0.55", lw=0.6, alpha=0.45,
            label="aligned APs" if not labeled_snip else None,
        )
        labeled_snip = True
        n_snips += 1

    ax.plot(t_rel, mp, color="C1", lw=2.4, zorder=3, label="mean passive")
    ax.plot(
        t_rel, mp_smooth, color="darkorange", lw=1.3, zorder=4,
        label=f"smoothed (σ={SPIKELET_PEAK_SMOOTH_MS:g} ms)",
    )
    base_p = float(np.mean(mp_smooth[:n_pre])) if n_pre > 0 else float(mp_smooth[0])
    ax.axhline(base_p, color="0.3", ls="-", lw=0.8, alpha=0.7, label="baseline")

    i_rel_p = _spikelet_local_peak_index(
        seg_p, sr, prominence=SPIKELET_MEANTRACE_MIN_PROMINENCE_MV,
    )
    i_rel_a = _spikelet_local_peak_index(seg_a, sr)
    if i_rel_a is None and len(seg_a):
        i_rel_a = int(np.argmax(seg_a))

    detected = bool(mt.get("meantrace10_detected"))
    if i_rel_p is not None:
        ip = n_pre + int(i_rel_p)
        ax.scatter(
            t_rel[ip], mp_smooth[ip],
            c="darkorange" if detected else "0.45",
            s=55, zorder=7,
            marker="o" if detected else "x",
            edgecolors="k", linewidths=0.4,
            label="spikelet peak (smoothed)" if detected else "peak search (no detection)",
        )
        amp_p = mt.get("meantrace10_amp_spikelet_mV")
        if amp_p is not None:
            ax.annotate(
                f"{amp_p:.2f} mV",
                (t_rel[ip], mp_smooth[ip]),
                textcoords="offset points", xytext=(4, 6),
                fontsize=8, color="0.15",
            )

    ax.axvline(0, color="limegreen", ls="--", lw=1.2)
    ax.axvspan(-SPIKELET_BASELINE_MS, 0, color="0.7", alpha=0.25)
    ax.axvspan(0, SPIKELET_MEANTRACE_PEAK_MS, color="C4", alpha=0.12)

    ax2 = ax.twinx()
    labeled_snip_a = False
    for sna in snips_a or []:
        if sna is None or len(sna) < i1:
            continue
        ax2.plot(
            t_rel, np.asarray(sna[:i1], dtype=float),
            color="C0", lw=0.5, alpha=0.25,
            label="aligned APs (active)" if not labeled_snip_a else None,
        )
        labeled_snip_a = True
    ax2.plot(t_rel, ma, color="C0", lw=1.8, alpha=0.9, zorder=3, label="mean active")
    if i_rel_a is not None:
        ia = n_pre + int(i_rel_a)
        ax2.scatter(
            t_rel[ia], ma[ia], c="C0", s=55, zorder=8, marker="v",
            edgecolors="k", linewidths=0.4,
            label="AP peak on mean",
        )
        ax.axvline(t_rel[ia], color="C0", ls=":", lw=1.0, alpha=0.7)
    ax2.set_ylabel("Active Vm (mV)", color="C0")
    ax2.tick_params(axis="y", labelcolor="C0")

    win_parts = [mp[n_pre:i1], mp_smooth[n_pre:i1], np.array([base_p])]
    for sn in snips_p or []:
        if sn is not None and len(sn) >= i1:
            win_parts.append(np.asarray(sn[n_pre:i1], dtype=float))
    win_p = np.concatenate(win_parts)
    y_lo = float(np.min(win_p))
    y_hi = float(np.max(win_p))
    span_p = max(y_hi - y_lo, 0.05)
    pad_p = max(0.03, 0.12 * span_p)
    ax.set_ylim(y_lo - pad_p, y_hi + pad_p)
    win_a = [ma[n_pre:i1]]
    for sna in snips_a or []:
        if sna is not None and len(sna) >= i1:
            win_a.append(np.asarray(sna[n_pre:i1], dtype=float))
    win_a = np.concatenate(win_a) if win_a else ma[n_pre:i1]
    if len(win_a):
        span_a = max(float(np.max(win_a) - np.min(win_a)), 1.0)
        pad_a = max(0.5, 0.12 * span_a)
        ax2.set_ylim(float(np.min(win_a)) - pad_a, float(np.max(win_a)) + pad_a)

    skip = mt.get("meantrace10_skip_reason")
    if detected:
        status = "DETECTED"
        detail = (
            f"  amp={mt.get('meantrace10_amp_spikelet_mV')} mV (smoothed), "
            f"delay={mt.get('meantrace10_delay_ms')} ms, "
            f"ratio={mt.get('meantrace10_amp_ratio')}"
        )
        status_color = "0.15"
    else:
        status = "NOT detected"
        detail = f" ({skip})" if skip else ""
        status_color = "C3"
    extra = f", n={n_snips}" if n_snips else ""
    ax.set_title(f"{title_base}: {status}{detail}{extra}", color=status_color)

    if not detected and skip:
        ax.text(
            0.01, 0.99, skip,
            transform=ax.transAxes, va="top", ha="left",
            fontsize=8, color="C3", wrap=True,
            bbox=dict(boxstyle="round,pad=0.25", facecolor="white", edgecolor="C3", alpha=0.9),
            zorder=10,
        )

    handles, labels = ax.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax.legend(handles + h2, labels + l2, loc="upper left", fontsize=7)
    ax.set_xlim(t_rel[0], t_rel[-1])
    ax.grid(True, alpha=0.3)


def save_spikelet_qc_plot(abf, plot_meta, plots_dir, stem, dir_tag=None):
    """Sweep overlay + aligned AP mean + meantrace10 QC (three subplots)."""
    if not plot_meta:
        return None
    import os

    os.makedirs(plots_dir, exist_ok=True)
    plt = _get_agg_plt()
    sweep = plot_meta["sweep"]
    active_ch = plot_meta["active_ch"]
    passive_ch = plot_meta["passive_ch"]
    sr = plot_meta["sr"]
    n_pre = plot_meta["n_pre"]
    n_post = plot_meta["n_post"]
    direction = plot_meta["direction"]

    abf.setSweep(sweepNumber=sweep, channel=active_ch)
    t = abf.sweepX
    y_a = np.asarray(abf.sweepY, dtype=float)
    abf.setSweep(sweepNumber=sweep, channel=passive_ch)
    y_p = np.asarray(abf.sweepY, dtype=float)

    fig, axes = plt.subplots(
        3, 1, figsize=(12, 10.5), sharex=False,
        gridspec_kw={"height_ratios": [1.35, 1.0, 0.85]},
    )
    ax_ov, ax_avg, ax_mt10 = axes

    ss = int(plot_meta["stim_start"])
    se = int(plot_meta["stim_stop"])
    i_left = max(int(plot_meta["pre_start"]), 0)
    i_right = min(se + int(0.05 * sr), len(t) - 1)

    ax_ov.plot(t, y_a, color="C0", lw=1.0, label=f"active ch{active_ch}")
    ax_ov.plot(t, y_p, color="C1", lw=1.0, label=f"passive ch{passive_ch}")
    if len(t) == 0:
        ax_ov.text(0.5, 0.5, "empty sweep", transform=ax_ov.transAxes, ha="center")
        err = plot_meta.get("error") or "empty sweep"
        ax_ov.set_title(f"{stem} — {direction}  {err}")
        ax_avg.axis("off")
        _meantrace10_qc_panel(
            ax_mt10,
            plot_meta.get("mean_a"),
            plot_meta.get("mean_p"),
            n_pre, sr,
            plot_meta.get("metrics"),
            plot_meta.get("rms"),
            snips_p=plot_meta.get("snips_p"),
            snips_a=plot_meta.get("snips_a"),
        )
        tag = dir_tag or SPIKELET_DIR_TAG.get(direction) or "na"
        fname = f"{stem}_{tag}_spikelets.png"
        path = os.path.join(plots_dir, fname)
        _savefig_white(fig, path)
        plt.close(fig)
        print(f"  saved spikelet QC (empty sweep): {path}")
        return path
    ax_ov.axvline(
        t[min(ss, len(t) - 1)], color="0.35", ls="--", lw=1.0, label="stim start / end",
    )
    ax_ov.axvline(t[min(se, len(t) - 1)], color="0.35", ls="--", lw=1.0)

    analyzed_ap = set()
    for r in plot_meta.get("ap_rows") or []:
        ai = r.get("ap_index")
        if ai is not None:
            analyzed_ap.add(int(ai) - 1)
    starts = plot_meta.get("ap_starts") if plot_meta.get("ap_starts") is not None else []
    peaks = plot_meta.get("ind_peaks") if plot_meta.get("ind_peaks") is not None else []
    for i, ip in enumerate(peaks):
        ip = int(ip)
        if 0 <= ip < len(t):
            ax_ov.scatter(
                t[ip], y_a[ip], c="C3", s=28, zorder=5, marker="o",
                label="AP peak" if i == 0 else None,
            )
        if i < len(starts) and np.isfinite(starts[i]):
            i0 = int(starts[i])
            if 0 <= i0 < len(t):
                ax_ov.scatter(
                    t[i0], y_a[i0], c="limegreen", s=28, zorder=5, marker="v",
                    label="AP start (t=0, 2nd inflection before peak)" if i == 0 else None,
                )
            if i in analyzed_ap and 0 <= i0 < len(t):
                t0 = t[i0]
                ax_ov.axvspan(t0 - SPIKELET_BASELINE_MS / 1000.0, t0, color="0.7", alpha=0.25)
                ax_ov.axvspan(t0, t0 + SPIKELET_PEAK_MS / 1000.0, color="C4", alpha=0.12)

    labeled_pass = False
    labeled_fail = False
    labeled_ratio = False
    labeled_10a = False
    labeled_10p = False
    thr_u = plot_meta.get("noise_thr_mV")
    if thr_u is None:
        thr_u = spikelet_amp_threshold(plot_meta.get("rms"))
    for r in plot_meta.get("ap_rows") or []:
        thr_loc = r.get("noise_thr_mV")
        if thr_loc is None:
            thr_loc = thr_u
        idx = r.get("i_peak_passive")
        if idx is None and r.get("t_peak_passive_ms") is not None:
            idx = int(round((r["t_peak_passive_ms"] / 1000.0) * sr))
        if idx is not None:
            idx = min(max(int(idx), 0), len(y_p) - 1)
            amp = r.get("amp_spikelet_mV")
            if r.get("ratio_gt_1"):
                ax_ov.scatter(
                    t[idx], y_p[idx], c="C3", s=70, zorder=7, marker="*",
                    edgecolors="k", linewidths=0.4,
                    label="ratio > 1 (omitted from folder plots)" if not labeled_ratio else None,
                )
                labeled_ratio = True
            elif r.get("detected"):
                ax_ov.scatter(
                    t[idx], y_p[idx], c="darkorange", s=48, zorder=6, marker="o",
                    edgecolors="k", linewidths=0.4,
                    label="spikelet pass" if not labeled_pass else None,
                )
                labeled_pass = True
            else:
                ax_ov.scatter(
                    t[idx], y_p[idx], c="0.45", s=42, zorder=6, marker="x",
                    label="spikelet below noise" if not labeled_fail else None,
                )
                labeled_fail = True
            if amp is not None:
                tag_txt = "local max" if r.get("detected") else "no peak"
                extra = ""
                if r.get("ratio_gt_1"):
                    extra += f"  ratio {r.get('amp_ratio')} > 1 (error, not plotted)"
                if r.get("delay_negative"):
                    extra += f"  delay {r.get('delay_ms')} ms < 0"
                ax_ov.annotate(
                    f"AP{r.get('ap_index')} {amp:.2f} mV {tag_txt}{extra}",
                    (t[idx], y_p[idx]),
                    textcoords="offset points", xytext=(4, 6),
                    fontsize=7,
                    color="C3" if (r.get("ratio_gt_1") or r.get("delay_negative")) else "0.15",
                    zorder=8,
                )
            elif r.get("skip_reason"):
                i0 = r.get("t_start_ms")
                if i0 is not None:
                    ax_ov.annotate(
                        f"AP{r.get('ap_index')} {r.get('skip_reason')}",
                        (float(i0) / 1000.0, ax_ov.get_ylim()[1]),
                        textcoords="offset points", xytext=(4, -12),
                        fontsize=7, color="C3", zorder=8,
                    )
        if r.get("t_10_active_ms") is not None:
            t10a = r["t_10_active_ms"] / 1000.0
            idx = min(max(int(round(t10a * sr)), 0), len(y_a) - 1)
            ax_ov.scatter(
                t[idx], y_a[idx], c="cyan", s=28, zorder=7, marker="D",
                label="AP 10%" if not labeled_10a else None,
            )
            labeled_10a = True
        if r.get("detected") and r.get("t_10_spikelet_ms") is not None:
            t10p = r["t_10_spikelet_ms"] / 1000.0
            idx = min(max(int(round(t10p * sr)), 0), len(y_p) - 1)
            ax_ov.scatter(
                t[idx], y_p[idx], c="magenta", s=28, zorder=7, marker="D",
                label="spikelet 10%" if not labeled_10p else None,
            )
            labeled_10p = True

    n_meas = sum(1 for r in (plot_meta.get("ap_rows") or []) if r.get("amp_spikelet_mV") is not None)
    n_pass = sum(1 for r in (plot_meta.get("ap_rows") or []) if r.get("detected"))
    n_bad_ratio = sum(1 for r in (plot_meta.get("ap_rows") or []) if r.get("ratio_gt_1"))
    src = (plot_meta.get("metrics") or {}).get("metric_source")
    err = plot_meta.get("error") or (plot_meta.get("metrics") or {}).get("skip_reason")
    ax_ov.set_ylabel("Vm (mV)")
    ratio_note = f", ratio>1={n_bad_ratio} omitted from folder plots" if n_bad_ratio else ""
    ax_ov.set_title(
        f"{stem} — {direction}  sweep {sweep} ({plot_meta.get('tier')}, "
        f"n_AP={plot_meta.get('n_ap')}, local_max={n_pass}/{n_meas}{ratio_note}  "
        f"{'noise gate OFF' if not SPIKELET_USE_NOISE_GATE else f'thr={_round_or_none(thr_u, 3)} mV'}  "
        f"source={src})"
    )
    if err:
        ax_ov.text(
            0.01, 0.99,
            f"note: {err}",
            transform=ax_ov.transAxes, va="top", ha="left",
            fontsize=9, color="C3", wrap=True,
            bbox=dict(boxstyle="round,pad=0.3", facecolor="white", edgecolor="C3", alpha=0.9),
            zorder=10,
        )
    ax_ov.legend(loc="upper right", fontsize=7)
    ax_ov.set_xlim(t[i_left], t[i_right])
    ax_ov.grid(True, axis="y", alpha=0.25)

    t_snip = (np.arange(-n_pre, n_post) / float(sr)) * 1000.0
    snips_p = plot_meta.get("snips_p") or []
    snips_a = plot_meta.get("snips_a") or []
    snips_ok = plot_meta.get("snips_ok")
    if snips_ok is None or len(snips_ok) != len(snips_p):
        snips_ok = [False] * len(snips_p)
    labeled_pass_tr = False
    labeled_fail_tr = False
    pass_p, pass_a = [], []
    for sn, sna, ok in zip(snips_p, snips_a, snips_ok):
        if len(sn) != len(t_snip):
            continue
        if ok:
            ax_avg.plot(
                t_snip, sn, color="C1", lw=0.8, alpha=0.45,
                label="pass traces" if not labeled_pass_tr else None,
            )
            labeled_pass_tr = True
            pass_p.append(sn)
            if len(sna) == len(t_snip):
                pass_a.append(sna)
        else:
            ax_avg.plot(
                t_snip, sn, color="0.55", lw=0.7, alpha=0.35,
                label="below-noise traces" if not labeled_fail_tr else None,
            )
            labeled_fail_tr = True
    mean_p_plot = None
    mean_p_label = None
    if pass_p:
        mean_p_plot = np.mean(np.vstack(pass_p), axis=0)
        mean_p_label = f"mean of pass (n={len(pass_p)})"
    elif plot_meta.get("mean_p") is not None:
        mean_p_plot = plot_meta.get("mean_p")
        mean_p_label = "mean of aligned (none passed gate)"
    if mean_p_plot is not None and len(mean_p_plot) == len(t_snip):
        ax_avg.plot(
            t_snip, mean_p_plot, color="C1", lw=2.2,
            label=mean_p_label,
        )
        base_avg = float(np.mean(mean_p_plot[:n_pre])) if n_pre > 0 else float(mean_p_plot[0])
        ax_avg.axhline(
            base_avg, color="0.3", ls="-", lw=0.8, alpha=0.7, label="baseline",
        )
    ax_avg.axvline(0, color="limegreen", ls="--", lw=1.2, label="t=0 AP start (inflection)")
    t10a = plot_meta.get("t10_a_rel_ms")
    t10p = plot_meta.get("t10_p_rel_ms")
    if t10a is not None:
        ax_avg.axvline(t10a, color="cyan", ls=":", lw=1.2, label="mean AP 10%")
    if t10p is not None:
        ax_avg.axvline(t10p, color="magenta", ls=":", lw=1.2, label="mean spikelet 10%")
    ax_avg.axvspan(
        -SPIKELET_BASELINE_MS, 0, color="0.7", alpha=0.25,
        label="passive baseline 1 ms (not t=0)",
    )
    ax_avg.axvspan(0, SPIKELET_PEAK_MS, color="C4", alpha=0.12)
    ax_avg.set_xlabel("Time from active AP start (ms)")
    ax_avg.set_ylabel("Passive Vm (mV)", color="C1")
    ax_avg.tick_params(axis="y", labelcolor="C1")

    ax_avg2 = ax_avg.twinx()
    for sna, ok in zip(snips_a, snips_ok):
        if len(sna) == len(t_snip) and ok:
            ax_avg2.plot(t_snip, sna, color="C0", lw=0.6, alpha=0.25)
    mean_a_plot = np.mean(np.vstack(pass_a), axis=0) if pass_a else plot_meta.get("mean_a")
    if mean_a_plot is not None and len(mean_a_plot) == len(t_snip):
        ax_avg2.plot(
            t_snip, mean_a_plot, color="C0", lw=1.6, alpha=0.9,
            label=f"mean AP pass (n={len(pass_a)})",
        )
    ax_avg2.set_ylabel("Active Vm (mV)", color="C0")
    ax_avg2.tick_params(axis="y", labelcolor="C0")

    handles, labels = ax_avg.get_legend_handles_labels()
    h2, l2 = ax_avg2.get_legend_handles_labels()
    ax_avg.legend(handles + h2, labels + l2, loc="upper left", fontsize=7)
    ap_align = (
        "Aligned AP1 (no multi-spike sweep)"
        if (plot_meta.get("n_ap") or 0) == 1 or plot_meta.get("tier") == "1"
        else "Aligned AP2+ (not AP1)"
    )
    ax_avg.set_title(
        f"{ap_align}: local_max={n_pass}/{n_meas}; "
        f"thick=mean of {len(pass_p)}; noise gate "
        f"{'OFF' if not SPIKELET_USE_NOISE_GATE else 'ON'}"
    )
    ax_avg.grid(True, alpha=0.3)

    _meantrace10_qc_panel(
        ax_mt10,
        plot_meta.get("mean_a"),
        plot_meta.get("mean_p"),
        n_pre, sr,
        plot_meta.get("metrics"),
        plot_meta.get("rms"),
        snips_p=plot_meta.get("snips_p"),
        snips_a=plot_meta.get("snips_a"),
    )

    try:
        fig.tight_layout()
    except Exception:
        pass
    tag = dir_tag or SPIKELET_DIR_TAG.get(direction) or "na"
    fname = f"{stem}_{tag}_spikelets.png"
    path = os.path.join(plots_dir, fname)
    _savefig_white(fig, path)
    plt.close(fig)
    print(f"  saved spikelet QC (primary sweep {sweep}): {path}")
    return path


def _save_spikelet_error_png(plots_dir, stem, dir_tag, direction, message):
    """Always write a 12/21 PNG so a failed direction is visible."""
    import os

    os.makedirs(plots_dir, exist_ok=True)
    plt = _get_agg_plt()
    fig, ax = plt.subplots(figsize=(10, 4))
    ax.axis("off")
    ax.set_title(f"{stem} — {direction}  (QC could not be drawn)")
    ax.text(
        0.5, 0.5, str(message),
        ha="center", va="center", color="C3", fontsize=11, wrap=True,
        transform=ax.transAxes,
    )
    tag = dir_tag or SPIKELET_DIR_TAG.get(direction) or "na"
    path = os.path.join(plots_dir, f"{stem}_{tag}_spikelets.png")
    _savefig_white(fig, path)
    plt.close(fig)
    print(f"  saved spikelet QC (error page): {path}")
    return path


def _spikelet_sweep_row(row_dict):
    return {k: row_dict.get(k) for k in SPIKELET_SWEEP_KEYS}


def _print_spikelet_pipeline_status(name, direction, metrics, sweep_metrics, ap_rows=None):
    """Console check: are sweep-mean ratio / delays / baseline actually filled?"""
    fname = os.path.basename(str(name or ""))
    print(
        f"  [spikelet] {fname} {direction}  "
        f"primary_sweep={metrics.get('sweep')}  n_sweeps={metrics.get('n_sweeps')}  "
        f"n_AP_used={metrics.get('n_AP_used')}  n_with_amp={metrics.get('n_with_amp')}  "
        f"n_pass={metrics.get('n_spikelet_detected')}  "
        f"noise_gate={'ON' if SPIKELET_USE_NOISE_GATE else 'OFF'}  "
        f"local_thr={metrics.get('noise_thr_mV')}"
    )
    print(
        f"    PRIMARY mean: spike={metrics.get('mean_amp_active_mV')}  "
        f"spikelet={metrics.get('mean_amp_spikelet_mV')}  "
        f"ratio={metrics.get('mean_amp_ratio')}  "
        f"n_ratio>1={metrics.get('n_ratio_gt_1')}  "
        f"delay_pk={metrics.get('mean_delay_ms')}  "
        f"delay_10={metrics.get('mean_delay_10_ms')}  "
        f"V_base={metrics.get('mean_baseline_passive_mV')}  "
        f"ap1_fallback={metrics.get('ap1_fallback')}  "
        f"skip={metrics.get('skip_reason')}"
    )
    print(
        f"    MEANTRACE10: spike={metrics.get('meantrace10_amp_active_mV')}  "
        f"spikelet={metrics.get('meantrace10_amp_spikelet_mV')}  "
        f"ratio={metrics.get('meantrace10_amp_ratio')}  "
        f"delay_pk={metrics.get('meantrace10_delay_ms')}  "
        f"delay_10={metrics.get('meantrace10_delay_10_ms')}  "
        f"det={metrics.get('meantrace10_detected')}  "
        f"skip={metrics.get('meantrace10_skip_reason')}"
    )
    by_sw = {}
    for r in ap_rows or []:
        by_sw.setdefault(r.get("sweep"), []).append(r)
    for m in sweep_metrics or []:
        flag = " PRIMARY" if m.get("sweep") == metrics.get("sweep") else ""
        print(
            f"    sweep {m.get('sweep')}{flag}: n_AP={m.get('n_AP_active')}  "
            f"n_amp={m.get('n_with_amp')}  n_det={m.get('n_spikelet_detected')}  "
            f"ratio={m.get('mean_amp_ratio')}  "
            f"n_ratio>1={m.get('n_ratio_gt_1')}  "
            f"d_pk={m.get('mean_delay_ms')}  d10={m.get('mean_delay_10_ms')}  "
            f"V_base={m.get('mean_baseline_passive_mV')}  "
            f"thr={m.get('noise_thr_mV')}"
        )
        show_aps = (m.get("n_with_amp") or 0) > 0 and (
            m.get("sweep") == metrics.get("sweep") or (m.get("n_spikelet_detected") or 0) == 0
        )
        if not show_aps:
            continue
        for r in by_sw.get(m.get("sweep"), []):
            if r.get("amp_spikelet_mV") is None:
                continue
            print(
                f"      AP{r.get('ap_index')}: amp={r.get('amp_spikelet_mV')}  "
                f"thr={r.get('noise_thr_mV')}  local_rms={r.get('rms_local_mV')}  "
                f"det={r.get('detected')}  {r.get('skip_reason') or ''}"
            )


def spikelets_for_file(abf, name, rec_dt, plots_dir=None, stem=None):
    """Both directions. Returns (ap_rows, summary_fields, plot_paths, sweep_rows)."""
    summary = empty_spikelet_summary_fields()
    all_rows = []
    sweep_rows = []
    plot_paths = []
    stem = stem or _abf_stem(name)
    if plots_dir and SAVE_SPIKELET_PLOTS:
        print(f"  Spikelet QC folder: {os.path.abspath(plots_dir)}")
    elif SAVE_SPIKELET_PLOTS:
        print("  Spikelet QC: plots_dir is empty, PNG will not be written")

    for direction, active, passive, tag in (
        ("ch0->ch2", 0, 2, "12"),
        ("ch2->ch0", 2, 0, "21"),
    ):
        try:
            result = analyze_spikelets_direction(
                abf, active, passive, direction
            )
            ap_rows, metrics, meta = result[0], result[1], result[2]
            sweep_metrics = result[3] if len(result) > 3 else []
        except Exception as exc:
            print(f"  Spikelet analysis error ({direction}): {exc}")
            traceback.print_exc()
            metrics = _empty_spikelet_metrics(str(exc))
            ap_rows, meta, sweep_metrics = [], None, []
        for r in ap_rows:
            all_rows.append(_spikelet_row({
                "file": name,
                "recording_datetime": rec_dt,
                "direction": direction,
                **r,
            }))
        for m in sweep_metrics:
            sweep_rows.append(_spikelet_sweep_row({
                "file": name,
                "recording_datetime": rec_dt,
                "direction": direction,
                "is_primary": m.get("sweep") == metrics.get("sweep"),
                **m,
            }))
        _print_spikelet_pipeline_status(
            name, direction, metrics, sweep_metrics, ap_rows=ap_rows,
        )
        for key, val in metrics.items():
            summary[f"spikelet_{key}_{tag}"] = val
        if plots_dir and SAVE_SPIKELET_PLOTS:
            if not meta:
                try:
                    meta = _fallback_spikelet_qc_meta(
                        abf, active, passive, direction,
                        metrics.get("skip_reason") or "no QC meta",
                    )
                except Exception as exc:
                    print(f"  Spikelet QC fallback failed ({direction}): {exc}")
                    traceback.print_exc()
                    meta = None
            if meta:
                try:
                    p = save_spikelet_qc_plot(
                        abf, meta, plots_dir, stem, dir_tag=tag,
                    )
                    if p:
                        plot_paths.append(p)
                except Exception as exc:
                    print(
                        f"  Spikelet QC error ({direction} sweep {meta.get('sweep')}): {exc}"
                    )
                    traceback.print_exc()
                    try:
                        p = _save_spikelet_error_png(
                            plots_dir, stem, tag, direction, str(exc),
                        )
                        if p:
                            plot_paths.append(p)
                    except Exception:
                        traceback.print_exc()
            else:
                try:
                    p = _save_spikelet_error_png(
                        plots_dir, stem, tag, direction,
                        metrics.get("skip_reason") or "could not build QC figure",
                    )
                    if p:
                        plot_paths.append(p)
                except Exception:
                    traceback.print_exc()

    return all_rows, summary, plot_paths, sweep_rows


# =============================================================================
# Tau / Cm (Tau Cm calculations.ipynb)
# =============================================================================


def cc_epochs_for_tau(borders, data_rate, voltage_ch, sweep_len=None):
    """Epochs for tau/Cm on double_cciv: Vm channels 0 and 2, CC fixed borders.

    Tau search starts at pre_end (end of pre-stim epoch), not at post_start.
    """
    sr = int(data_rate)
    if voltage_ch == 0:
        pre_start, pre_end = borders["start10"], borders["end10"]
        post_start, post_end = borders["start11"], borders["end11"]
    elif voltage_ch == 2:
        pre_start, pre_end = borders["start20"], borders["end20"]
        post_start, post_end = borders["start21"], borders["end21"]
    else:
        raise ValueError(f"unsupported voltage channel {voltage_ch}")
    # tau clock starts at end of pre epoch (onset of step), not at post mean window
    stim_start = pre_end
    post_stop = post_end + int(0.5 * sr)
    if sweep_len is not None:
        post_stop = min(post_stop, sweep_len)
    return {
        "pre_start": pre_start,
        "pre_end": pre_end,
        "post_start": post_start,
        "post_end": post_end,
        "stim_start": stim_start,
        "post_stop": post_stop,
        "voltage_ch": voltage_ch,
    }


def compute_tau_cm(abf, voltage_ch, rin_MOhm, borders=None):
    """
    Prefer sweep with ΔV=mean(post)-mean(pre) in [-35, -10] mV and post_min >= -90
    (closest |ΔV| to zero). If none: sweep with ΔV closest to -35 mV (post_min may be
    below -90; values still exported).

    v_63 = V_pre + 0.63 * (V_post - V_pre);
    tau [ms] from pre_end; Cm = tau_ms / Rin * 1000 (CC/Gj Rin).
    """
    ch = f"ch{voltage_ch}"
    empty = {f"{k}_{ch}": None for k in TAU_CM_FIELD_SUFFIXES}

    if borders is not None:
        abf.setSweep(sweepNumber=abf.sweepList[0], channel=voltage_ch)
        epochs = cc_epochs_for_tau(borders, abf.dataRate, voltage_ch, sweep_len=len(abf.sweepY))
    else:
        ce = channel_epochs(abf, voltage_ch)
        if ce is None:
            return {**empty, f"tau_skip_reason_{ch}": "no stimulus detected in sweepC"}
        sr = int(abf.dataRate)
        epochs = {
            "pre_start": ce["pre_start"],
            "pre_end": ce["pre_start"] + int(TCM_PRE_MS * sr),
            "post_start": ce["start"],
            "post_end": ce["stop"],
            "stim_start": ce["pre_start"] + int(TCM_PRE_MS * sr),  # pre_end
            "post_stop": ce["post_stop"],
        }

    sr = int(abf.dataRate)
    pre_start = epochs["pre_start"]
    pre_end = epochs["pre_end"]
    post_start = epochs["post_start"]
    post_end = epochs["post_end"]
    stim_start = epochs["stim_start"]  # tau clock origin = pre_end
    post_stop = epochs["post_stop"]
    epoch_label = (
        f"ch{voltage_ch} V_pre[{pre_start}:{pre_end}] "
        f"V_post[{post_start}:{post_end}] tau_from@{stim_start}(=pre_end) "
        f"tau→[{stim_start}:{post_stop}]"
    )

    # cand: (sweep, v_pre, v_post, delta_v, post_min)
    in_range = []
    all_sweeps = []

    for sweep_num in abf.sweepList:
        abf.setSweep(sweepNumber=sweep_num, channel=voltage_ch)
        y = abf.sweepY
        if pre_end > len(y) or post_end > len(y) or post_stop > len(y):
            continue
        if post_start >= post_end:
            continue

        sweep_v_pre = float(np.mean(y[pre_start:pre_end]))
        sweep_v_post = float(np.mean(y[post_start:post_end]))
        post_min = float(np.min(y[post_start:post_end]))
        sweep_dv = sweep_v_post - sweep_v_pre
        cand = (sweep_num, sweep_v_pre, sweep_v_post, sweep_dv, post_min)
        all_sweeps.append(cand)

        if (
            TCM_DV_SEARCH_MIN <= sweep_dv <= TCM_DV_SEARCH_MAX
            and post_min >= TCM_VMIN_LIMIT
        ):
            in_range.append(cand)

    selection_note = None
    if in_range:
        best_sweep, v_pre, v_post, delta_v, post_min = min(in_range, key=lambda c: abs(c[3]))
        selection_note = (
            f"in-range [{TCM_DV_SEARCH_MIN},{TCM_DV_SEARCH_MAX}] mV, "
            f"post_min>={TCM_VMIN_LIMIT}; "
            f"delta_V={delta_v:.3f} mV; V_post={v_post:.3f} mV; "
            f"V_post_min={post_min:.3f} mV"
        )
    elif all_sweeps:
        # fallback: closest to -35; post_min may be < -90 — still export
        target = float(TCM_DV_SEARCH_MIN)
        best_sweep, v_pre, v_post, delta_v, post_min = min(
            all_sweeps, key=lambda c: abs(c[3] - target)
        )
        selection_note = (
            f"fallback: closest to {target:g} mV "
            f"(no in-range with post_min>={TCM_VMIN_LIMIT}); "
            f"delta_V={delta_v:.3f} mV; V_post={v_post:.3f} mV; "
            f"V_post_min={post_min:.3f} mV"
        )
    else:
        reason = f"no sweeps with valid pre/post epochs ({post_start}:{post_end})"
        return {
            **empty,
            f"tau_skip_reason_{ch}": reason,
            f"_tau_plot_{ch}": {
                "sweep": None,
                "voltage_ch": voltage_ch,
                "pre_start": pre_start,
                "pre_end": pre_end,
                "post_start": post_start,
                "post_end": post_end,
                "stim_start": stim_start,
                "post_stop": post_stop,
                "epoch_label": epoch_label,
                "skip_reason": reason,
                "rin_MOhm": rin_MOhm,
            },
        }

    abf.setSweep(sweepNumber=best_sweep, channel=voltage_ch)
    v_63 = v_pre + 0.63 * (v_post - v_pre)
    tau_samples = None
    trace = abf.sweepY[stim_start:post_stop]
    for i, v in enumerate(trace):
        if v <= v_63:
            tau_samples = i
            break

    plot_base = {
        "sweep": best_sweep,
        "voltage_ch": voltage_ch,
        "pre_start": pre_start,
        "pre_end": pre_end,
        "post_start": post_start,
        "post_end": post_end,
        "stim_start": stim_start,
        "post_stop": post_stop,
        "epoch_label": epoch_label,
        "v_pre": v_pre,
        "v_post": v_post,
        "v_63": v_63,
        "delta_v": delta_v,
        "post_min": post_min,
        "selection_note": selection_note,
        "rin_MOhm": rin_MOhm,
    }

    common_fields = {
        f"V_pre_mV_{ch}": round(v_pre, 3),
        f"V_post_mV_{ch}": round(v_post, 3),
        f"V_post_min_mV_{ch}": round(post_min, 3),
        f"delta_V_mV_{ch}": round(delta_v, 3),
        f"tau_sweep_{ch}": best_sweep,
        f"tau_selection_note_{ch}": selection_note,
    }

    if tau_samples is None:
        reason = "trace never reached 63% of delta_V for tau"
        return {
            **empty,
            **common_fields,
            f"tau_skip_reason_{ch}": reason,
            f"_tau_plot_{ch}": {**plot_base, "tau_ms": None, "cm_pF": None, "skip_reason": reason},
        }

    tau_ms = round((tau_samples / sr) * 1000, 3)
    tau_index = stim_start + tau_samples
    cm_pF = None
    cm_skip = None
    if rin_MOhm is None:
        cm_skip = "missing Rin for Cm"
    elif rin_MOhm == 0:
        cm_skip = "Rin = 0"
    else:
        cm_pF = round((tau_ms / rin_MOhm) * 1000, 3)

    return {
        **common_fields,
        f"tau_ms_{ch}": tau_ms,
        f"Cm_pF_{ch}": cm_pF,
        f"tau_skip_reason_{ch}": None,
        f"Cm_skip_reason_{ch}": cm_skip,
        f"_tau_plot_{ch}": {
            **plot_base,
            "tau_index": tau_index,
            "tau_ms": tau_ms,
            "cm_pF": cm_pF,
            "cm_skip": cm_skip,
        },
    }


def tau_cm_for_file(abf, rin_ch0, rin_ch2, borders=None):
    """Tau and Cm for ch0 and ch2 using CC/Gj Rin."""
    out = empty_tau_cm_fields()
    plot_meta = {}
    for voltage_ch, rin in ((0, rin_ch0), (2, rin_ch2)):
        ch = f"ch{voltage_ch}"
        result = compute_tau_cm(abf, voltage_ch, rin, borders=borders)
        plot_key = f"_tau_plot_{ch}"
        if plot_key in result:
            plot_meta[ch] = result.pop(plot_key)
        out.update({k: v for k, v in result.items() if not k.startswith("_")})
    out["_tau_plot_meta"] = plot_meta
    return out


# =============================================================================
# QC plots (saved to disk)
# =============================================================================


def _get_agg_plt():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.facecolor": "white",
        "savefig.edgecolor": "none",
        "savefig.transparent": False,
        "axes.edgecolor": "black",
        "axes.labelcolor": "black",
        "xtick.color": "black",
        "ytick.color": "black",
        "text.color": "black",
    })
    return plt


def _savefig_white(fig, path, dpi=None):
    """Save PNG with opaque white figure/axes background (readable axes)."""
    if dpi is None:
        dpi = PLOT_DPI
    fig.patch.set_facecolor("white")
    for ax in fig.get_axes():
        ax.set_facecolor("white")
        ax.tick_params(colors="black")
        for attr in ("xaxis", "yaxis"):
            axis = getattr(ax, attr, None)
            label = getattr(axis, "label", None) if axis is not None else None
            if label is not None:
                label.set_color("black")
        title = getattr(ax, "title", None)
        if title is not None:
            title.set_color("black")
        for spine in ax.spines.values():
            spine.set_color("black")
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    fig.savefig(
        path, dpi=dpi, bbox_inches="tight",
        facecolor="white", edgecolor="none", transparent=False,
    )


def _mpl_cmap(name=None):
    """Colormap that works on matplotlib 3.7–3.11 (cm.get_cmap was removed)."""
    plt = _get_agg_plt()
    name = name or CC_VM_CMAP
    try:
        return plt.colormaps[name]
    except Exception:
        pass
    try:
        import matplotlib.cm as mplcm
        return mplcm.get_cmap(name)
    except Exception:
        return plt.cm.viridis


def _finish_folder_fig(fig, out_path, skip_tight=False):
    """tight_layout (optional) then save; never leave the caller without a PNG."""
    if not skip_tight:
        try:
            fig.tight_layout()
        except Exception:
            pass
    _savefig_white(fig, out_path)
    plt_mod = _get_agg_plt()
    plt_mod.close(fig)
    print(f"  saved {out_path}")
    return out_path


def cell_props_plots_dir(folder_path):
    """Return ``folder_path/Cell_properties_plots``, creating it if needed."""
    path = os.path.join(folder_path, CELL_PROPS_PLOTS_SUBDIR)
    os.makedirs(path, exist_ok=True)
    return path


def qc_plots_dir(folder_path):
    """Alias for ``cell_props_plots_dir`` (backward compatible)."""
    return cell_props_plots_dir(folder_path)


def cc_plots_dir(folder_path):
    """Return ``folder_path/CC_plots``, creating it if needed."""
    path = os.path.join(folder_path, CC_PLOTS_SUBDIR)
    os.makedirs(path, exist_ok=True)
    return path


def spikelet_plots_dir(folder_path):
    """Return ``folder_path/Spikelet_plots``, creating it if needed."""
    path = os.path.join(folder_path, SPIKELET_PLOTS_SUBDIR)
    os.makedirs(path, exist_ok=True)
    return path


def _abf_stem(filepath):
    return os.path.splitext(os.path.basename(filepath))[0]


def _save_ap_qc_plot(abf, voltage_ch, plots_dir, stem):
    """Vm trace on >=4 AP sweep: peaks, AP starts, stim lines, 1st-spike FWHM."""
    plt = _get_agg_plt()

    epochs = channel_epochs(abf, voltage_ch)
    if epochs is None:
        return None

    sr = int(abf.dataRate)
    start, stop = epochs["start"], epochs["stop"]
    pre_start, post_stop = epochs["pre_start"], epochs["post_stop"]

    sweep_4 = cp_sweep_4aps(abf, voltage_ch, start, stop)
    if sweep_4 is None:
        return None

    ind_infls = cp_infl_points(abf, sweep_4, voltage_ch, start, stop)
    ind_peaks = cp_peak_indices(abf, sweep_4, voltage_ch, start, stop)
    if len(ind_peaks) == 0:
        return None

    ap_start, ap_end = cp_spike_begin(ind_peaks, ind_infls)
    abf.setSweep(sweepNumber=sweep_4, channel=voltage_ch)
    t = abf.sweepX
    y = abf.sweepY

    fig, ax = plt.subplots(figsize=(14, 6))
    ax.plot(t[pre_start:post_stop], y[pre_start:post_stop], color="C0", lw=1, label="Vm")
    ax.scatter(t[ind_peaks], y[ind_peaks], color="red", s=40, zorder=5, label="peaks")
    starts = [int(s) for s in ap_start if np.isfinite(s)]
    if starts:
        ax.scatter(
            t[starts], y[starts], color="limegreen", s=40, zorder=5, label="AP start"
        )
    ax.axvline(start / sr, color="gray", ls=":", lw=1, alpha=0.8)
    ax.axvline(stop / sr, color="gray", ls=":", lw=1, alpha=0.8)

    if ap_start and np.isfinite(ap_start[0]) and ap_end and np.isfinite(ap_end[0]):
        fwhm_ms, fwhm_idx, fwhm_h = cp_fwhm_details(
            abf, sweep_4, voltage_ch, ap_start[0], ap_end[0], ind_peaks[0], sr
        )
        if fwhm_idx is not None and fwhm_h is not None:
            t0 = fwhm_idx[0] / sr
            t1 = fwhm_idx[1] / sr
            ax.hlines(
                fwhm_h, t0, t1,
                colors="darkorange",
                lw=5,
                zorder=6,
                label=f"FWHM 1st AP ({fwhm_ms} ms)" if fwhm_ms is not None else "FWHM 1st AP",
            )
            ax.plot([t0, t1], [fwhm_h, fwhm_h], "o", color="darkorange", ms=8, zorder=7)
            # mark AP baseline and peak used for 50% level
            i_ap0 = int(ap_start[0])
            i_pk = int(ind_peaks[0])
            ax.axhline(y[i_ap0], color="limegreen", ls="--", lw=1, alpha=0.7)
            ax.axhline(y[i_pk], color="red", ls=":", lw=1, alpha=0.5)
            ax.axhline(fwhm_h, color="darkorange", ls=":", lw=1, alpha=0.5)

    ax.set_xlabel("Time (s)")
    ax.set_ylabel("Vm (mV)")
    ax.set_title(f"{stem} — ch{voltage_ch} sweep {sweep_4} (>= {CP_MIN_APS} APs)")
    ax.legend(loc="upper right", fontsize=9)
    fig.tight_layout()
    path = os.path.join(plots_dir, f"{stem}_ch{voltage_ch}_AP.png")
    _savefig_white(fig, path)
    plt.close(fig)
    return path


def _save_rin_qc_plot(abf, voltage_ch, current_ch, plots_dir, stem, borders, rin_mohm, r2, vm_range_label, rin_note):
    """Single I–V QC plot for Rin used in CC, Gj, and Cm."""
    plt = _get_agg_plt()

    if voltage_ch == 0:
        v_start, v_end, i_index = borders["start11"], borders["end11"], borders["Rtime_ch0"]
    else:
        v_start, v_end, i_index = borders["start21"], borders["end21"], borders["Rtime_ch2"]

    if voltage_ch == 0:
        spike_start, spike_end = borders["start11"], borders["end11"]
    else:
        spike_start, spike_end = borders["start21"], borders["end21"]
    series = collect_cc_iv_series(
        abf, voltage_ch, current_ch, v_start, v_end, i_index,
        spike_start=spike_start, spike_end=spike_end,
    )
    currents = np.array(series["currents"])
    voltages = np.array(series["voltages"])
    if len(currents) == 0:
        return None

    primary = np.array(series["in_primary"])
    fallback = np.array(series["in_fallback"]) & ~primary
    subth = np.array(series.get("subthreshold", [True] * len(currents)))
    spiked = ~subth
    other = subth & ~primary & ~fallback

    fig, ax = plt.subplots(figsize=(8, 6))
    if spiked.any():
        ax.scatter(
            currents[spiked], voltages[spiked], color="0.75", marker="x", s=50,
            label="after first AP (excluded)",
        )
    if other.any():
        ax.scatter(currents[other], voltages[other], color="lightgray", s=50, label="other Vm (subth)")
    if fallback.any():
        ax.scatter(
            currents[fallback], voltages[fallback], color="C1", s=60,
            label=f"fallback {RIN_VMIN_FALLBACK}...{RIN_VMAX_FALLBACK} mV",
        )
    if primary.any():
        ax.scatter(
            currents[primary], voltages[primary], color="C0", s=60,
            label=f"primary {RIN_VMIN}...{RIN_VMAX} mV",
        )

    if rin_mohm is not None:
        fit_mask = primary
        if fit_mask.sum() < RIN_MIN_POINTS:
            fit_mask = primary | fallback
        if fit_mask.sum() >= RIN_MIN_POINTS and r2 is not None:
            i_fit = np.linspace(currents[fit_mask].min(), currents[fit_mask].max(), 100)
            slope_mV_pA = rin_mohm / 1000.0
            intercept = float(np.mean(voltages[fit_mask] - slope_mV_pA * currents[fit_mask]))
            ax.plot(
                i_fit, slope_mV_pA * i_fit + intercept, color="red", lw=2,
                label=f"Rin = {rin_mohm} MΩ, R² = {r2}",
            )
        else:
            ax.plot([], [], color="red", lw=2, label=f"Rin = {rin_mohm} MΩ ({rin_note or 'delta'})")

    ax.set_xlabel("Current at Rtime (pA)")
    ax.set_ylabel("Mean Vm in stim epoch (mV)")
    ax.set_title(f"{stem} — ch{voltage_ch} I–V (Rin for CC / Gj / Cm)")
    ax.legend(loc="best", fontsize=9)
    fig.tight_layout()
    path = os.path.join(plots_dir, f"{stem}_ch{voltage_ch}_IV_Rin.png")
    _savefig_white(fig, path)
    plt.close(fig)
    return path


def _save_tau_cm_qc_plot(abf, voltage_ch, plots_dir, stem, plot_meta):
    """Vm sweep for tau/Cm QC — saved even when tau/Cm failed (diagnostic)."""
    plt = _get_agg_plt()

    ch = f"ch{voltage_ch}"
    meta = plot_meta.get(ch)
    if not meta:
        fig, ax = plt.subplots(figsize=(10, 4))
        ax.text(0.5, 0.5, "no tau/Cm metadata", ha="center", va="center")
        ax.set_title(f"{stem} — ch{voltage_ch} tau/Cm (no data)")
        path = os.path.join(plots_dir, f"{stem}_ch{voltage_ch}_tau_Cm.png")
        _savefig_white(fig, path)
        plt.close(fig)
        return path

    if meta.get("sweep") is None:
        fig, ax = plt.subplots(figsize=(10, 4))
        msg = meta.get("skip_reason", "tau/Cm not computed")
        if meta.get("epoch_label"):
            msg = f"{msg}\n{meta['epoch_label']}"
        ax.text(0.5, 0.5, msg, ha="center", va="center", wrap=True, fontsize=10)
        ax.set_title(f"{stem} — ch{voltage_ch} tau/Cm SKIPPED")
        path = os.path.join(plots_dir, f"{stem}_ch{voltage_ch}_tau_Cm.png")
        _savefig_white(fig, path)
        plt.close(fig)
        return path

    abf.setSweep(sweepNumber=meta["sweep"], channel=voltage_ch)
    ps, pe = meta["pre_start"], meta["post_stop"]
    stim = meta["stim_start"]
    t = abf.sweepX[ps:pe]
    y = abf.sweepY[ps:pe]

    fig, ax = plt.subplots(figsize=(14, 6))
    ax.plot(t, y, color="C0", lw=1, label=f"Vm ch{voltage_ch}")
    if meta.get("pre_end") is not None:
        ax.axvspan(
            abf.sweepX[meta["pre_start"]], abf.sweepX[meta["pre_end"]],
            color="green", alpha=0.08, label="pre epoch",
        )
    if meta.get("post_end") is not None:
        ax.axvspan(
            abf.sweepX[meta["post_start"]], abf.sweepX[meta["post_end"]],
            color="blue", alpha=0.08, label="post epoch",
        )
    ax.axvline(abf.sweepX[stim], color="gray", ls=":", lw=1, alpha=0.8, label="tau start (pre_end)")
    if meta.get("v_pre") is not None:
        ax.axhline(meta["v_pre"], color="green", ls="--", lw=1.5, label=f"V_pre = {meta['v_pre']:.2f} mV")
    if meta.get("v_post") is not None:
        ax.axhline(meta["v_post"], color="blue", ls="--", lw=1.5, label=f"V_post = {meta['v_post']:.2f} mV")
    if meta.get("v_63") is not None:
        dv = meta.get("delta_v")
        dv_str = f" (ΔV={dv:.2f} mV)" if dv is not None else ""
        ax.axhline(
            meta["v_63"], color="darkorange", ls="--", lw=1.5,
            label=f"63% = {meta['v_63']:.2f} mV{dv_str}",
        )
    if meta.get("tau_index") is not None and meta.get("tau_ms") is not None:
        ax.scatter(
            abf.sweepX[meta["tau_index"]], abf.sweepY[meta["tau_index"]],
            color="purple", s=60, zorder=5, label=f"tau = {meta['tau_ms']:.3f} ms",
        )
        ax.axvline(abf.sweepX[meta["tau_index"]], color="purple", ls=":", alpha=0.5)

    rin = meta.get("rin_MOhm")
    tau_ms = meta.get("tau_ms")
    cm_pF = meta.get("cm_pF")
    cm_skip = meta.get("cm_skip")
    if tau_ms is not None and rin is not None and cm_pF is not None:
        cm_line = f"Cm = {cm_pF:.3f} pF  (tau/Rin×1000; Rin={rin} MΩ)"
    elif cm_skip:
        cm_line = f"Cm n/a ({cm_skip}; Rin={rin})"
    else:
        cm_line = meta.get("skip_reason", "")

    tau_str = f"{tau_ms:.3f} ms" if tau_ms is not None else "n/a"
    epoch_lbl = meta.get("epoch_label", "")
    sel_note = meta.get("selection_note") or ""
    title = (
        f"{stem} — ch{voltage_ch} sweep {meta['sweep']}  |  tau={tau_str}  |  {cm_line}\n"
        f"{epoch_lbl}"
    )
    if sel_note:
        title = f"{title}\n{sel_note}"
    ax.set_title(title, fontsize=8)
    ax.set_xlabel("Time (s)")
    ax.set_ylabel("Vm (mV)")
    ax.legend(loc="upper right", fontsize=8)
    fig.tight_layout()
    path = os.path.join(plots_dir, f"{stem}_ch{voltage_ch}_tau_Cm.png")
    _savefig_white(fig, path)
    plt.close(fig)
    return path


def save_qc_plots(
    abf, filepath, plots_dir, borders,
    rin_ch0, r2_ch0, rin_note_0,
    rin_ch2, r2_ch2, rin_note_2,
    tau_plot_meta=None,
    save_ap=True,
    save_rin=True,
    save_tau=True,
):
    """Save AP + Rin I–V + tau/Cm QC PNGs into plots_dir."""
    os.makedirs(plots_dir, exist_ok=True)
    stem = _abf_stem(filepath)
    saved = []

    for voltage_ch, current_ch, rin, r2, note in (
        (0, 1, rin_ch0, r2_ch0, rin_note_0),
        (2, 3, rin_ch2, r2_ch2, rin_note_2),
    ):
        jobs = []
        if save_ap:
            jobs.append((_save_ap_qc_plot, (abf, voltage_ch, plots_dir, stem)))
        if save_rin:
            jobs.append((
                _save_rin_qc_plot,
                (abf, voltage_ch, current_ch, plots_dir, stem, borders, rin, r2, None, note),
            ))
        if save_tau:
            jobs.append((_save_tau_cm_qc_plot, (abf, voltage_ch, plots_dir, stem, tau_plot_meta or {})))
        for saver, extra in jobs:
            try:
                path = saver(*extra)
                if path:
                    saved.append(path)
            except Exception as exc:
                print(f"  QC plot skip ({saver.__name__} ch{voltage_ch}): {exc}")

    return saved


def _cc_direction_epochs(borders, direction):
    """Return pre_start, pre_end, post_start, post_end, rtime, active_ch, passive_ch, label."""
    if direction == "ch0->ch2":
        return (
            borders["start10"], borders["end10"],
            borders["start11"], borders["end11"],
            borders["Rtime_ch0"],
            0, 2, "CC12 (ch0→ch2)",
        )
    if direction == "ch2->ch0":
        return (
            borders["start20"], borders["end20"],
            borders["start21"], borders["end21"],
            borders["Rtime_ch2"],
            2, 0, "CC21 (ch2→ch0)",
        )
    raise ValueError(direction)


def _cc_block_vpost(abf, block, active_ch, post_start, post_end):
    """Mean Vm in post epoch (active cell) for each row in coupling block."""
    vposts = []
    for r in block:
        sn = r["sweep"]
        abf.setSweep(sweepNumber=sn, channel=active_ch)
        vposts.append(float(statistics.mean(abf.sweepY[post_start:post_end])))
    return vposts


def _attach_gj_for_plot(block, rin_passive):
    """Copy coupling rows and fill Gj_sweep_nS from Rin_passive when missing."""
    out = []
    for r in block:
        rr = dict(r)
        gj = rr.get("Gj_sweep_nS")
        skip = rr.get("Gj_sweep_skip_reason")
        if gj is None and rin_passive is not None:
            try:
                gj, skip = gj_nS(rr.get("CC"), rin_passive)
            except Exception:
                gj, skip = None, "gj_plot_error"
        rr["Gj_sweep_nS"] = gj
        rr["Gj_sweep_skip_reason"] = skip
        out.append(rr)
    return out


def _save_cc_traces_plot(abf, borders, block_02, block_20, plots_dir, stem):
    """Two panels: selected CC sweeps (valid vs skipped). Spike-cut sweeps omitted."""
    plt = _get_agg_plt()
    fig, axes = plt.subplots(2, 1, figsize=(14, 10), sharex=False)

    for ax, block, direction in (
        (axes[0], block_02, "ch0->ch2"),
        (axes[1], block_20, "ch2->ch0"),
    ):
        pre_s, pre_e, post_s, post_e, rtime, active_ch, passive_ch, label = (
            _cc_direction_epochs(borders, direction)
        )
        t_pre0 = abf.sweepX[pre_s]
        t_pre1 = abf.sweepX[min(pre_e, len(abf.sweepX) - 1)]
        t_post0 = abf.sweepX[post_s]
        t_post1 = abf.sweepX[min(post_e, len(abf.sweepX) - 1)]
        t_rt = abf.sweepX[min(rtime, len(abf.sweepX) - 1)]

        ax.axvspan(t_pre0, t_pre1, color="green", alpha=0.12, label="pre epoch")
        ax.axvspan(t_post0, t_post1, color="blue", alpha=0.12, label="post epoch")
        ax.axvline(t_rt, color="purple", ls="--", lw=1.5, label=f"Rtime (I sample @{rtime})")

        plotted_valid = False
        plotted_skip = False
        view_start = pre_s
        view_end = min(post_e + int(0.1 * abf.dataRate), len(abf.sweepY))

        for r in block:
            sn = r["sweep"]
            cc = r["CC"]
            skipped = cc is None
            color = "0.55" if skipped else "C0"
            alpha = 0.55 if skipped else 0.85
            lw = 1.0 if skipped else 1.4

            abf.setSweep(sweepNumber=sn, channel=active_ch)
            t = abf.sweepX[view_start:view_end]
            y_a = abf.sweepY[view_start:view_end]
            lbl = None
            if skipped and not plotted_skip:
                lbl = "skipped (no CC)"
                plotted_skip = True
            elif (not skipped) and not plotted_valid:
                lbl = "valid CC"
                plotted_valid = True
            ax.plot(t, y_a, color=color, alpha=alpha, lw=lw, label=lbl)

            abf.setSweep(sweepNumber=sn, channel=passive_ch)
            y_p = abf.sweepY[view_start:view_end]
            ax.plot(t, y_p, color=color, alpha=alpha * 0.7, lw=0.9, ls="--")

            mid = (post_s + post_e) // 2
            abf.setSweep(sweepNumber=sn, channel=active_ch)
            y_ann = float(abf.sweepY[mid])
            if cc is None:
                cc_txt = "skip"
            else:
                cc_txt = f"{cc:.3f}"
                gj = r.get("Gj_sweep_nS")
                if gj is not None:
                    cc_txt += f" Gj={gj:.2f}"
            ax.annotate(
                f"s{sn}:{cc_txt}",
                xy=(abf.sweepX[mid], y_ann),
                fontsize=7,
                color=color,
                alpha=0.9,
            )

        ax.set_ylabel("Vm (mV)")
        ax.set_title(
            f"{stem} — {label}  |  solid=active ch{active_ch}, dashed=passive ch{passive_ch}",
            fontsize=10,
        )
        ax.legend(loc="upper right", fontsize=8)

    axes[1].set_xlabel("Time (s)")
    fig.tight_layout()
    os.makedirs(plots_dir, exist_ok=True)
    path = os.path.join(plots_dir, f"{stem}_CC_traces.png")
    _savefig_white(fig, path)
    plt.close(fig)
    print(f"  saved {path}")
    return path


def _save_cc_vs_vpost_plot(abf, borders, block_02, block_20, plots_dir, stem):
    """CC vs mean active Vm in post epoch; Gj on twin axis when present."""
    plt = _get_agg_plt()
    fig, axes = plt.subplots(1, 2, figsize=(12, 5.4))

    for ax, block, direction, title in (
        (axes[0], block_02, "ch0->ch2", "CC12 (ch0→ch2)"),
        (axes[1], block_20, "ch2->ch0", "CC21 (ch2→ch0)"),
    ):
        _pre_s, _pre_e, post_s, post_e, _rtime, active_ch, _passive_ch, _ = (
            _cc_direction_epochs(borders, direction)
        )
        vposts = _cc_block_vpost(abf, block, active_ch, post_s, post_e)

        xv, yv, xs, ys = [], [], [], []
        xg, yg = [], []
        for r, vp in zip(block, vposts):
            sw_label = f"s{r['sweep']}"
            if r["CC"] is not None:
                xv.append(vp)
                yv.append(r["CC"])
                ax.annotate(
                    sw_label, (vp, r["CC"]),
                    textcoords="offset points", xytext=(5, 5),
                    fontsize=8, color="0.2", zorder=4,
                )
                gj = r.get("Gj_sweep_nS")
                if gj is not None:
                    xg.append(vp)
                    yg.append(gj)
            else:
                xs.append(vp)
                ys.append(0.0)
                ax.annotate(
                    sw_label, (vp, 0.0),
                    textcoords="offset points", xytext=(5, 5),
                    fontsize=8, color="0.45", zorder=4,
                )

        if xv:
            colors = ["C3" if c < 0 else "C0" for c in yv]
            ax.scatter(xv, yv, c=colors, s=55, zorder=3, label="valid CC")
        if xs:
            ax.scatter(xs, ys, c="0.5", marker="x", s=55, zorder=3, label="skipped")
        ax.axhline(0, color="black", lw=0.6, alpha=0.4)
        ax.set_xlabel(f"Mean Vm active ch{active_ch} in post [{post_s}:{post_e}] (mV)")
        ax.set_ylabel("CC", color="C0")
        ax.tick_params(axis="y", labelcolor="C0")
        ax.set_title(f"{stem} — {title}")

        if xg:
            ax2 = ax.twinx()
            ax2.scatter(xg, yg, c="C3", marker="s", s=36, zorder=2, alpha=0.85, label="Gj")
            ax2.set_ylabel("Gj (nS)", color="C3")
            ax2.tick_params(axis="y", labelcolor="C3")
            handles, labels = ax.get_legend_handles_labels()
            h2, l2 = ax2.get_legend_handles_labels()
            ax.legend(handles + h2, labels + l2, loc="best", fontsize=8)
        else:
            ax.legend(loc="best", fontsize=8)

    fig.tight_layout()
    os.makedirs(plots_dir, exist_ok=True)
    path = os.path.join(plots_dir, f"{stem}_CC_vs_Vpost.png")
    _savefig_white(fig, path)
    plt.close(fig)
    print(f"  saved {path}")
    return path


def save_cc_qc_plots(
    abf, filepath, plots_dir, borders,
    block_02, block_20,
    rin_ch0=None, rin_ch2=None,
):
    """Save CC QC figures into CC_plots (traces + CC/Gj vs Vm)."""
    if not plots_dir:
        plots_dir = os.path.join(os.path.dirname(os.path.abspath(filepath)), CC_PLOTS_SUBDIR)
    os.makedirs(plots_dir, exist_ok=True)
    stem = _abf_stem(filepath)
    block_02 = _attach_gj_for_plot(block_02, rin_ch2)
    block_20 = _attach_gj_for_plot(block_20, rin_ch0)
    saved = []
    jobs = []
    if SAVE_CC_TRACE_PLOTS:
        jobs.append((_save_cc_traces_plot, (abf, borders, block_02, block_20, plots_dir, stem)))
    if SAVE_CC_VPOST_PLOTS:
        jobs.append((_save_cc_vs_vpost_plot, (abf, borders, block_02, block_20, plots_dir, stem)))
    for saver, args in jobs:
        try:
            path = saver(*args)
            if path:
                saved.append(path)
        except Exception as exc:
            print(f"  CC plot skip ({saver.__name__}): {exc}")
            traceback.print_exc()
    return saved


def gj_nS(cc, rin_passive_MOhm):
    """
    Gj [nS] = (CC/(1-CC)) * (1/Rin_passive [MOhm]) * 1000.

    Rin_passive: for Gj12 (ch0→ch2) use Rin of cell 2 (ch2);
                 for Gj21 (ch2→ch0) use Rin of cell 1 (ch0).
    Requires 0 < CC < 1 and Rin_passive > 0.
    """
    if cc is None or rin_passive_MOhm is None or rin_passive_MOhm == 0:
        return None, "missing CC or Rin"
    if cc <= 0:
        return None, f"CC <= 0 (got {cc})"
    if cc >= 1:
        return None, f"CC >= 1 (got {cc})"
    if rin_passive_MOhm <= 0:
        return None, f"Rin <= 0 (got {rin_passive_MOhm}); Gj requires positive passive Rin"
    try:
        return round((cc / (1.0 - cc)) * (1000.0 / rin_passive_MOhm), 4), None
    except (ZeroDivisionError, TypeError):
        return None, "Gj calculation error"



def mean_valid_cc(cc_list):
    vals = [c for c in cc_list if c is not None]
    return round(statistics.mean(vals), 4) if vals else None


def n_valid_cc(cc_list):
    return sum(c is not None for c in cc_list)


def n_negative_cc(cc_list):
    """Count of valid (non-None) CC values that are negative."""
    return sum(c is not None and c < 0 for c in cc_list)


def _parse_recording_datetime(value):
    """Parse ABF recording time from summary row to datetime, or None."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    s = str(value).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def folder_summary_timed_rows(summary_rows):
    """All File_summary rows with a parseable recording datetime, sorted by time."""
    out = []
    for row in summary_rows:
        dt = _parse_recording_datetime(row.get("recording_datetime"))
        if dt is None:
            continue
        out.append((dt, row))
    out.sort(key=lambda item: item[0])
    return out


def folder_summary_complete_rows(summary_rows):
    """
    Files with Rin1, Rin2, CC12, CC21, Gj12, Gj21 all computed, sorted by time.

    Returns list of (datetime, row_dict).
    """
    keys = ("Rin1_MOhm", "Rin2_MOhm", "CC12", "CC21", "Gj12_nS", "Gj21_nS")
    out = []
    for dt, row in folder_summary_timed_rows(summary_rows):
        if any(row.get(k) is None for k in keys):
            continue
        out.append((dt, row))
    return out


def _finite_xy(times, values):
    """Keep (time, value) pairs where value is a finite number."""
    xs, ys = [], []
    for t, v in zip(times, values):
        if v is None:
            continue
        try:
            fv = float(v)
        except (TypeError, ValueError):
            continue
        if np.isfinite(fv):
            xs.append(t)
            ys.append(fv)
    return xs, ys


def _plot_timed(ax, times, values, *args, **kwargs):
    xs, ys = _finite_xy(times, values)
    if xs:
        ax.plot(xs, ys, *args, **kwargs)
    return len(xs)


def _mark_negative_delay_points(ax, xs, ys, already_labeled=False):
    """Red X on delay < 0 (spikelet event before the AP event)."""
    nx, ny = [], []
    for x, y in zip(xs or [], ys or []):
        fy = _finite_number(y)
        if fy is not None and fy < 0:
            nx.append(x)
            ny.append(fy)
    if nx:
        ax.scatter(
            nx, ny, marker="x", c="C3", s=52, zorder=6,
            label="delay < 0" if not already_labeled else None,
        )
    return len(nx)


def _vm_for_channel(row, ch_tag):
    """V_rest from staged linear I–V intercept; hold_V if no fit passed checks."""
    v = row.get(f"V_rest_mV_{ch_tag}")
    if v is not None:
        return v
    return row.get(f"hold_V_mV_{ch_tag}")


def _spikelet_file_metric(row, tag, metric):
    """File-level spikelet value (primary-sweep means stored on File_summary)."""
    keys = (f"spikelet_mean_{metric}_{tag}",)
    if not str(metric).endswith("amp_ratio"):
        keys = (
            f"spikelet_mean_{metric}_{tag}",
            f"spikelet_avg_{metric}_{tag}",
        )
    for key in keys:
        val = row.get(key)
        if val is None:
            continue
        v = _finite_number(val)
        if v is None:
            continue
        if str(metric).endswith("amp_ratio"):
            v = _plottable_amp_ratio(v)
            if v is None:
                continue
        return v
    return None


def _mean_numeric(items, key):
    vals = []
    for item in items or []:
        v = _finite_number(item.get(key) if item else None)
        if v is not None:
            vals.append(v)
    return statistics.mean(vals) if vals else None


def _spikelet_file_names_match(row_file, fname):
    names = {fname, os.path.basename(str(fname or ""))}
    rf = row_file or ""
    return os.path.basename(str(rf)) in names or str(rf) in names


def _primary_sweep_row(sweep_rows, fname, direction):
    """The QC/primary sweep row for one file and direction."""
    matched = []
    for rec in sweep_rows or []:
        if rec.get("direction") != direction:
            continue
        if not _spikelet_file_names_match(rec.get("file"), fname):
            continue
        matched.append(rec)
    if not matched:
        return None
    for rec in matched:
        if rec.get("is_primary"):
            return rec
    return matched[0]


def _collect_spikelet_baselines(sweep_rows, ap_rows=None):
    """All sweep-mean spikelet baselines in the folder (both directions)."""
    vals = []
    for rec in sweep_rows or []:
        vm = _sweep_row_metric(rec, "baseline_mV")
        if vm is not None:
            vals.append(float(vm))
    if vals:
        return vals
    for rec in ap_rows or []:
        vm = _finite_number(rec.get("baseline_passive_mV"))
        if vm is not None:
            vals.append(float(vm))
    return vals


def spikelet_baseline_mode_vm(sweep_rows, ap_rows=None, n_grid=512):
    """
    Vm where the density of spikelet baselines in the folder is highest.

    Uses a Gaussian KDE of sweep-mean baselines (passive, 1 ms before t=0).
    """
    vals = _collect_spikelet_baselines(sweep_rows, ap_rows)
    if not vals:
        return None
    arr = np.asarray(vals, dtype=float)
    arr = arr[np.isfinite(arr)]
    if arr.size == 0:
        return None
    if arr.size == 1 or float(np.ptp(arr)) < 1e-6:
        return _round_or_none(float(np.median(arr)), 3)
    try:
        from scipy.stats import gaussian_kde
        kde = gaussian_kde(arr)
        lo, hi = float(np.min(arr)), float(np.max(arr))
        pad = 0.05 * (hi - lo)
        grid = np.linspace(lo - pad, hi + pad, int(n_grid))
        dens = kde(grid)
        return _round_or_none(float(grid[int(np.argmax(dens))]), 3)
    except Exception:
        n_bins = max(8, min(40, int(np.sqrt(arr.size) * 2)))
        counts, edges = np.histogram(arr, bins=n_bins)
        i = int(np.argmax(counts))
        return _round_or_none(0.5 * (float(edges[i]) + float(edges[i + 1])), 3)


def _sweep_closest_to_vm(sweep_rows, fname, direction, target_mv, ap_rows=None):
    """Sweep whose mean spikelet baseline is closest to ``target_mv``."""
    target = _finite_number(target_mv)
    if target is None:
        return None
    candidates = []

    def _add(rec):
        vm = _sweep_row_metric(rec, "baseline_mV")
        if vm is None:
            return
        sw = rec.get("sweep")
        candidates.append((abs(float(vm) - target), sw if sw is not None else 10**9, rec))

    for rec in sweep_rows or []:
        if rec.get("direction") != direction:
            continue
        if not _spikelet_file_names_match(rec.get("file"), fname):
            continue
        _add(rec)
    if not candidates and ap_rows:
        rebuilt = _spikelet_means_from_ap_rows(
            [r for r in ap_rows if r.get("direction") == direction]
        )
        for (fn, _dir, sw), rec in rebuilt.items():
            if not _spikelet_file_names_match(fn, fname):
                continue
            packed = dict(rec)
            packed.setdefault("file", fn)
            packed.setdefault("direction", direction)
            packed.setdefault("sweep", sw)
            _add(packed)
    if not candidates:
        return None
    candidates.sort(key=lambda t: (t[0], t[1]))
    return candidates[0][2]


def _sweep_row_metric(row, metric):
    if not row:
        return None
    key_map = {
        "amp_ratio": ("mean_amp_ratio",),
        "amp_active_mV": ("mean_amp_active_mV", "avg_amp_active_mV", "amp_active_mV"),
        "amp_spikelet_mV": (
            "mean_amp_spikelet_mV", "avg_amp_spikelet_mV", "amp_spikelet_mV",
        ),
        "delay_ms": ("mean_delay_ms", "avg_delay_ms", "delay_ms", "delay_peak_ms"),
        "delay_10_ms": ("mean_delay_10_ms", "avg_delay_10_ms", "delay_10_ms"),
        "baseline_mV": (
            "mean_baseline_passive_mV", "baseline_passive_mV", "mean_vm_begin_mV",
        ),
    }
    for key in key_map.get(metric, (metric,)):
        v = _finite_number(row.get(key))
        if v is None:
            continue
        if str(metric).endswith("amp_ratio"):
            v = _plottable_amp_ratio(v)
            if v is None:
                continue
        return v
    return None


def _spikelet_file_mean_from_aps(ap_rows, fname, direction, metric, primary_sweep=None):
    """Mean of AP2+ with a measured spikelet amp (optionally one sweep only)."""
    vals = []
    for r in ap_rows or []:
        if r.get("direction") != direction:
            continue
        if not _spikelet_file_names_match(r.get("file"), fname):
            continue
        if primary_sweep is not None and r.get("sweep") != primary_sweep:
            continue
        if _finite_number(r.get("amp_spikelet_mV")) is None:
            continue
        if metric == "amp_ratio" and (r.get("ratio_gt_1") or _ratio_gt_one(r.get("amp_ratio"))):
            continue
        v = r.get(metric)
        if v is None and metric == "delay_ms":
            v = r.get("delay_peak_ms")
        v = _finite_number(v)
        if v is not None:
            vals.append(v)
    return statistics.mean(vals) if vals else None


def _spikelet_means_from_ap_rows(spikelet_rows):
    """One mean per file, direction, and sweep from AP2+ with measured amp."""
    by = {}
    for r in spikelet_rows or []:
        if _finite_number(r.get("amp_spikelet_mV")) is None:
            continue
        fname = os.path.basename(str(r.get("file") or ""))
        direction = r.get("direction")
        if not fname or not direction:
            continue
        key = (fname, direction, r.get("sweep"))
        by.setdefault(key, []).append(r)

    out = {}
    for key, items in by.items():
        rec = {
            "sweep": key[2],
            "mean_amp_ratio": _mean_row_field(
                [r for r in items if not r.get("ratio_gt_1") and not _ratio_gt_one(r.get("amp_ratio"))],
                "amp_ratio",
            ),
            "mean_amp_active_mV": _mean_row_field(items, "amp_active_mV"),
            "mean_amp_spikelet_mV": _mean_row_field(items, "amp_spikelet_mV"),
            "mean_delay_ms": _mean_row_field(items, "delay_ms", "delay_peak_ms"),
            "mean_delay_10_ms": _mean_row_field(items, "delay_10_ms"),
            "mean_vm_begin_mV": _mean_row_field(items, "vm_begin_active_mV"),
            "mean_baseline_passive_mV": _mean_row_field(items, "baseline_passive_mV"),
            "amp_ratio": _mean_row_field(
                [r for r in items if not r.get("ratio_gt_1") and not _ratio_gt_one(r.get("amp_ratio"))],
                "amp_ratio",
            ),
            "amp_active_mV": _mean_row_field(items, "amp_active_mV"),
            "amp_spikelet_mV": _mean_row_field(items, "amp_spikelet_mV"),
            "delay_ms": _mean_row_field(items, "delay_ms", "delay_peak_ms"),
            "delay_10_ms": _mean_row_field(items, "delay_10_ms"),
        }
        out[key] = rec
    return out


def _format_time_axes(axes_flat):
    from matplotlib.dates import AutoDateLocator, ConciseDateFormatter

    for ax in np.atleast_1d(axes_flat).ravel():
        locator = AutoDateLocator()
        ax.xaxis.set_major_locator(locator)
        ax.xaxis.set_major_formatter(ConciseDateFormatter(locator))


def _cc_gj_twin_panel(ax, times, cc_vals, gj_vals, cc_label, gj_label, panel_title):
    n_cc = _plot_timed(ax, times, cc_vals, "o-", color="C0", label=cc_label)
    ax.set_ylabel("CC", color="C0")
    ax.tick_params(axis="y", labelcolor="C0")
    ax2 = ax.twinx()
    n_gj = _plot_timed(ax2, times, gj_vals, "s--", color="C3", label=gj_label)
    ax2.set_ylabel("Gj (nS)", color="C3")
    ax2.tick_params(axis="y", labelcolor="C3")
    ax.set_title(panel_title)
    ax.grid(True, alpha=0.3)
    n = n_cc + n_gj
    if n == 0:
        ax.text(0.5, 0.5, "no data", transform=ax.transAxes,
                ha="center", va="center", color="0.5", fontsize=9)
    else:
        lines = ax.get_lines() + ax2.get_lines()
        labels = [ln.get_label() for ln in lines]
        ax.legend(lines, labels, loc="best", fontsize=8)
    return n


def save_folder_summary_plot(summary_rows, out_path, title=None):
    """
    Folder time course, 4 stacked subplots (file-mean CC/Gj per direction):

    1) CC12 + Gj12 (twin axis)
    2) CC21 + Gj21 (twin axis)
    3) Rin1 (ch0) and Rin2 (ch2)
    4) Vm1 (ch0) and Vm2 (ch2) — V_rest (staged I–V intercept), else hold_V
    """
    pairs = folder_summary_timed_rows(summary_rows)
    if not pairs:
        return None

    times = [dt for dt, _ in pairs]
    rows = [r for _, r in pairs]
    plt = _get_agg_plt()
    fig, axes = plt.subplots(4, 1, figsize=(11, 12), sharex=True)
    if title:
        fig.suptitle(f"{title}  —  CC / Gj / Rin / Vm vs time", fontsize=12)

    n = 0
    n += _cc_gj_twin_panel(
        axes[0], times,
        [r.get("CC12") for r in rows],
        [r.get("Gj12_nS") for r in rows],
        "CC12", "Gj12", "CC12 / Gj12 (ch0→ch2)",
    )
    n += _cc_gj_twin_panel(
        axes[1], times,
        [r.get("CC21") for r in rows],
        [r.get("Gj21_nS") for r in rows],
        "CC21", "Gj21", "CC21 / Gj21 (ch2→ch0)",
    )

    ax_r = axes[2]
    n_r = 0
    n_r += _plot_timed(ax_r, times, [r.get("Rin1_MOhm") for r in rows],
                       "o-", color="C0", label="Rin1 (ch0)")
    n_r += _plot_timed(ax_r, times, [r.get("Rin2_MOhm") for r in rows],
                       "s-", color="C1", label="Rin2 (ch2)")
    n += n_r
    ax_r.set_ylabel("Rin (MΩ)")
    ax_r.set_title("Rin1 / Rin2")
    ax_r.grid(True, alpha=0.3)
    if n_r == 0:
        ax_r.text(0.5, 0.5, "no data", transform=ax_r.transAxes,
                  ha="center", va="center", color="0.5", fontsize=9)
    else:
        ax_r.legend(loc="best", fontsize=8)

    ax_v = axes[3]
    n_v = 0
    n_v += _plot_timed(
        ax_v, times, [r.get("V_rest_mV_ch0") for r in rows],
        "o-", color="C0", label="Vm1 V_rest",
    )
    n_v += _plot_timed(
        ax_v, times, [
            r.get("hold_V_mV_ch0") if r.get("V_rest_mV_ch0") is None else None
            for r in rows
        ],
        "o--", color="C0", alpha=0.55, label="Vm1 hold_V (no I–V fit)",
    )
    n_v += _plot_timed(
        ax_v, times, [r.get("V_rest_mV_ch2") for r in rows],
        "s-", color="C1", label="Vm2 V_rest",
    )
    n_v += _plot_timed(
        ax_v, times, [
            r.get("hold_V_mV_ch2") if r.get("V_rest_mV_ch2") is None else None
            for r in rows
        ],
        "s--", color="C1", alpha=0.55, label="Vm2 hold_V (no I–V fit)",
    )
    n += n_v
    ax_v.set_ylabel("Vm (mV)")
    ax_v.set_title(
        "Vm1 / Vm2  (V_rest = staged I–V intercept at I=0; dashed = hold_V)"
    )
    ax_v.set_xlabel("Recording time")
    ax_v.grid(True, alpha=0.3)
    if n_v == 0:
        ax_v.text(0.5, 0.5, "no data", transform=ax_v.transAxes,
                  ha="center", va="center", color="0.5", fontsize=9)
    else:
        ax_v.legend(loc="best", fontsize=8)

    if n == 0:
        plt.close(fig)
        return None

    _format_time_axes(axes[-1])
    fig.autofmt_xdate()
    try:
        fig.tight_layout()
    except Exception:
        pass
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    _savefig_white(fig, out_path)
    plt.close(fig)
    print(f"  saved {out_path}")
    return out_path


def save_folder_spikelet_over_time_plot(
    summary_rows, out_path, title=None, spikelet_rows=None, spikelet_sweep_rows=None,
    target_vm=None,
):
    """
    Spikelet/spike vs recording time: one point per file.

    If ``target_vm`` is None, Y = primary-sweep mean ratio and delays
    (AP2+, or AP1 if no multi-spike sweep). Primary = first sweep with
    >=4 APs, else 3, else 2, else 1.

    If ``target_vm`` is set, Y comes from the sweep whose mean spikelet
    baseline (passive, 1 ms before t=0) is closest to that Vm.

    X = recording datetime. Amplitude vs time is not plotted.
    Always writes the PNG (empty panels if no values).
    """
    pairs = folder_summary_timed_rows(summary_rows)
    use_dates = bool(pairs)
    if pairs:
        times = [dt for dt, _ in pairs]
        rows = [r for _, r in pairs]
    else:
        rows = list(summary_rows or [])
        rows.sort(key=lambda r: (_abf_file_number(r.get("file") or ""), str(r.get("file") or "")))
        times = list(range(len(rows)))
        print("  spikelet vs time: no recording datetime; using file order on X")
    if not rows:
        print("  spikelet vs time: skipped (no File_summary rows)")
        return None
    dir_12, dir_21 = "ch0->ch2", "ch2->ch0"
    near_vm = _finite_number(target_vm)

    def _pick_sweep(fname, direction):
        if near_vm is not None:
            return _sweep_closest_to_vm(
                spikelet_sweep_rows, fname, direction, near_vm,
                ap_rows=spikelet_rows,
            )
        picked = _primary_sweep_row(spikelet_sweep_rows, fname, direction)
        if picked is None:
            picked = _primary_sweep_row(
                spikelet_sweep_rows, str(fname or ""), direction
            )
        return picked

    def _val(row, tag, direction, metric):
        fname = os.path.basename(str(row.get("file") or ""))
        picked = _pick_sweep(fname, direction)
        if picked is None and fname != str(row.get("file") or ""):
            picked = _pick_sweep(str(row.get("file") or ""), direction)
        v = _sweep_row_metric(picked, metric)
        if v is not None:
            return v
        p_sw = None if picked is None else picked.get("sweep")
        v = _spikelet_file_mean_from_aps(
            spikelet_rows, fname, direction, metric, primary_sweep=p_sw,
        )
        if v is not None:
            return v
        if near_vm is not None:
            return None
        return _spikelet_file_metric(row, tag, metric)

    ratio12 = [_val(r, "12", dir_12, "amp_ratio") for r in rows]
    ratio21 = [_val(r, "21", dir_21, "amp_ratio") for r in rows]
    dpk12 = [_val(r, "12", dir_12, "delay_ms") for r in rows]
    dpk21 = [_val(r, "21", dir_21, "delay_ms") for r in rows]
    d1012 = [_val(r, "12", dir_12, "delay_10_ms") for r in rows]
    d1021 = [_val(r, "21", dir_21, "delay_10_ms") for r in rows]
    n_ratio = sum(v is not None for v in ratio12 + ratio21)
    n_del = sum(v is not None for v in dpk12 + dpk21 + d1012 + d1021)
    n_picked = 0
    for r in rows:
        fname = os.path.basename(str(r.get("file") or ""))
        if _pick_sweep(fname, dir_12) or _pick_sweep(fname, dir_21):
            n_picked += 1
    if near_vm is not None:
        print(
            f"  spikelet vs time (sweep nearest baseline {near_vm} mV vs recording time): "
            f"files={len(rows)}, files_with_sweep={n_picked}, "
            f"amp_ratio={n_ratio}, delays={n_del}"
        )
    else:
        print(
            f"  spikelet vs time (PRIMARY sweep mean vs recording time): "
            f"files={len(rows)}, files_with_primary_sweep={n_picked}, "
            f"amp_ratio={n_ratio}, delays={n_del}"
        )
    if n_ratio == 0 and n_del == 0:
        print(
            "  spikelet vs time: ratio/delay empty — check console [spikelet] lines "
            "and Excel Spikelet_sweeps (mean_amp_ratio, mean_delay_ms, "
            "mean_baseline_passive_mV, is_primary)"
        )

    plt = _get_agg_plt()
    fig, axes = plt.subplots(3, 1, figsize=(11, 10), sharex=True)
    if title:
        if near_vm is not None:
            fig.suptitle(
                f"{title}  —  ratio and delays vs recording time  "
                f"(sweep with spikelet baseline nearest {near_vm:g} mV)",
                fontsize=12,
            )
        else:
            fig.suptitle(
                f"{title}  —  primary-sweep ratio and delays vs recording time",
                fontsize=12,
            )

    panels = (
        (axes[0], "Amplitude ratio (spikelet / spike; ratio>1 omitted)", "ratio",
         ((ratio12, "o-", "C0", "12 (ch0→ch2)"),
          (ratio21, "s-", "C1", "21 (ch2→ch0)"))),
        (axes[1], "Delay peak (ms)", "delay peak (ms)",
         ((dpk12, "o-", "C0", "12 peak"),
          (dpk21, "s-", "C1", "21 peak"))),
        (axes[2], "Delay 10% (ms)", "delay 10% (ms)",
         ((d1012, "o-", "C0", "12 10%"),
          (d1021, "s-", "C1", "21 10%"))),
    )
    n_total = 0
    labeled_neg = False
    for ax, panel_title, ylabel, series in panels:
        n = 0
        is_delay = "delay" in ylabel.lower()
        for vals, style, color, label in series:
            n += _plot_timed(ax, times, vals, style, color=color, label=label)
            if is_delay:
                xs, ys = _finite_xy(times, vals)
                if _mark_negative_delay_points(ax, xs, ys, already_labeled=labeled_neg):
                    labeled_neg = True
        n_total += n
        ax.set_title(panel_title)
        ax.set_ylabel(ylabel)
        ax.grid(True, alpha=0.3)
        if n == 0:
            ax.text(0.5, 0.5, "no data", transform=ax.transAxes,
                    ha="center", va="center", color="0.5", fontsize=9)
        else:
            ax.legend(loc="best", fontsize=8)
    axes[-1].set_xlabel("Recording time" if use_dates else "File order")

    if use_dates:
        _format_time_axes(axes[-1])
        fig.autofmt_xdate()
    return _finish_folder_fig(fig, out_path)


def file_spikelet_vm_curves(sweep_rows, direction, y_key, ap_rows=None):
    """
    Per-file (spikelet baseline Vm, y) curves for one direction, sorted by Vm.

    One point per sweep. X = mean baseline of spikelets on that sweep
    (passive 1 ms before t=0). Returns list of
    (file_name, recording_datetime, vm_list, y_list).
    """
    y_alts = {
        "mean_amp_ratio": ("mean_amp_ratio",),
        "mean_delay_ms": ("mean_delay_ms", "avg_delay_ms", "delay_ms", "delay_peak_ms"),
        "mean_delay_10_ms": ("mean_delay_10_ms", "avg_delay_10_ms", "delay_10_ms"),
    }.get(y_key, (y_key,))

    def _y_of(row):
        for key in y_alts:
            v = _finite_number(row.get(key))
            if v is None:
                continue
            if str(y_key).endswith("amp_ratio"):
                v = _plottable_amp_ratio(v)
                if v is None:
                    continue
            return v
        return None

    def _vm_of(row):
        for key in (
            "mean_baseline_passive_mV",
            "baseline_passive_mV",
            "mean_vm_begin_mV",
            "vm_begin_active_mV",
        ):
            v = _finite_number(row.get(key))
            if v is not None:
                return v
        return None

    by_file = {}
    file_dt = {}
    for row in sweep_rows or []:
        if row.get("direction") != direction:
            continue
        y = _y_of(row)
        vm = _vm_of(row)
        if y is None or vm is None:
            continue
        fname = row.get("file") or "?"
        by_file.setdefault(fname, []).append((vm, y))
        if fname not in file_dt:
            file_dt[fname] = _parse_recording_datetime(row.get("recording_datetime"))

    if not by_file and ap_rows:
        rebuilt = _spikelet_means_from_ap_rows(
            [r for r in ap_rows if r.get("direction") == direction]
        )
        for (fname, _dir, _sw), rec in rebuilt.items():
            y = _y_of(rec)
            vm = _vm_of(rec)
            if y is None or vm is None:
                continue
            by_file.setdefault(fname, []).append((vm, y))
            if fname not in file_dt:
                file_dt[fname] = None

    curves = []
    for fname, pairs in by_file.items():
        pairs.sort(key=lambda p: p[0])
        curves.append((
            fname,
            file_dt.get(fname),
            [p[0] for p in pairs],
            [p[1] for p in pairs],
        ))
    curves.sort(key=lambda c: (c[1] is None, c[1] or datetime.min, c[0]))
    return curves


def save_folder_spikelet_vs_vm_plot(
    sweep_rows, out_path, title=None, summary_rows=None, spikelet_rows=None,
):
    """
    Folder overview like CC vs Vm: sweep-mean spikelet/spike ratio and delays vs
    mean spikelet baseline (passive, 1 ms before t=0). Color = first→last file.
    Colored lines = per-file linear fit; black = mean slope of those file lines.
    Always writes the PNG (empty panels if no detections).
    """
    from matplotlib.colors import Normalize

    plt = _get_agg_plt()
    fig, axes = plt.subplots(3, 2, figsize=(14, 11), sharex="col")
    if title:
        fig.suptitle(f"{title}  —  spikelet / spike vs mean spikelet baseline", fontsize=12)

    file_pos, first_lbl, last_lbl, n_files = _cc_vm_file_color_map(
        list(sweep_rows or []) + list(spikelet_rows or []),
        summary_rows=summary_rows,
    )
    print(f"  spikelet vs Vm color: first={first_lbl}  →  last={last_lbl}  ({n_files} files)")
    cmap = _mpl_cmap()

    panels = (
        ("mean_amp_ratio", "amp spikelet / amp spike (ratio>1 omitted)"),
        ("mean_delay_ms", "delay peak (ms)"),
        ("mean_delay_10_ms", "delay 10% (ms)"),
    )
    directions = (
        ("ch0->ch2", "ch0→ch2"),
        ("ch2->ch0", "ch2→ch0"),
    )
    n_curves = 0
    any_data = False
    for col, (direction, dir_title) in enumerate(directions):
        for row_i, (y_key, ylabel) in enumerate(panels):
            ax = axes[row_i, col]
            curves = file_spikelet_vm_curves(
                sweep_rows, direction, y_key, ap_rows=spikelet_rows,
            )
            n_curves += len(curves)
            if not curves:
                ax.set_title(f"{dir_title} — no data" if row_i == 0 else "")
                if col == 0:
                    ax.set_ylabel(ylabel)
                ax.grid(True, alpha=0.3)
                ax.text(0.5, 0.5, "no data", transform=ax.transAxes,
                        ha="center", va="center", color="0.5", fontsize=9)
                continue
            any_data = True
            file_xy = []
            labeled_neg = False
            for fname, _dt, vms, ys in curves:
                color = cmap(float(file_pos.get(
                    fname, file_pos.get(os.path.basename(str(fname)), 0.5)
                )))
                ax.scatter(vms, ys, color=[color], s=28, zorder=3, alpha=0.9)
                if "delay" in y_key:
                    if _mark_negative_delay_points(ax, vms, ys, already_labeled=labeled_neg):
                        labeled_neg = True
                if CC_VM_FIT_LINEAR:
                    x1, y1, _s, _b, _r2 = _linear_cc_vs_vm(vms, ys)
                    if x1 is not None:
                        ax.plot(x1, y1, "-", color=color, lw=1.3, alpha=0.85)
                file_xy.append((vms, ys))
            _plot_mean_of_file_lines(ax, file_xy)
            ax.grid(True, alpha=0.3)
            if row_i == 0:
                ax.set_title(f"{dir_title} — {len(curves)} file(s)")
            if col == 0:
                ax.set_ylabel(ylabel)
            if row_i == 2:
                ax.set_xlabel("Mean spikelet baseline (passive, mV)")

    print(f"  spikelet vs Vm curves: {n_curves} (need ratio/delay + mean spikelet baseline)")
    if not any_data:
        print("  spikelet vs Vm: no points yet; still saving empty figure")

    used_cbar = False
    try:
        sm = plt.cm.ScalarMappable(cmap=cmap, norm=Normalize(0.0, 1.0))
        sm.set_array([])
        cbar = fig.colorbar(sm, ax=list(axes.ravel()), fraction=0.046, pad=0.03)
        cbar.set_label("first file  →  last file")
        try:
            cbar.set_ticks([0.0, 1.0], labels=[f"first\n{first_lbl}", f"last\n{last_lbl}"])
        except TypeError:
            cbar.set_ticks([0.0, 1.0])
            cbar.ax.set_yticklabels([f"first  {first_lbl}", f"last  {last_lbl}"])
        cbar.ax.tick_params(labelsize=8)
        used_cbar = True
    except Exception as exc:
        print(f"  spikelet vs Vm colorbar skipped: {exc}")

    return _finish_folder_fig(fig, out_path, skip_tight=used_cbar)


def save_folder_spikelet_meantrace10_over_time_plot(
    summary_rows, out_path, title=None, spikelet_sweep_rows=None, target_vm=None,
):
    """Parallel mean-trace scheme: ratio/delays vs time from mean passive trace."""
    pairs = folder_summary_timed_rows(summary_rows)
    use_dates = bool(pairs)
    if pairs:
        times = [dt for dt, _ in pairs]
        rows = [r for _, r in pairs]
    else:
        rows = list(summary_rows or [])
        rows.sort(key=lambda r: (_abf_file_number(r.get("file") or ""), str(r.get("file") or "")))
        times = list(range(len(rows)))
        print("  spikelet meantrace10 vs time: no recording datetime; using file order on X")
    if not rows:
        print("  spikelet meantrace10 vs time: skipped (no File_summary rows)")
        return None
    dir_12, dir_21 = "ch0->ch2", "ch2->ch0"
    near_vm = _finite_number(target_vm)

    def _pick(fname, direction):
        if near_vm is not None:
            return _sweep_closest_to_vm(spikelet_sweep_rows, fname, direction, near_vm)
        picked = _primary_sweep_row(spikelet_sweep_rows, fname, direction)
        if picked is None:
            picked = _primary_sweep_row(spikelet_sweep_rows, str(fname or ""), direction)
        return picked

    def _val(row, direction, key):
        fname = os.path.basename(str(row.get("file") or ""))
        picked = _pick(fname, direction)
        if picked is None and fname != str(row.get("file") or ""):
            picked = _pick(str(row.get("file") or ""), direction)
        return _sweep_row_metric(picked, key)

    ratio12 = [_val(r, dir_12, "meantrace10_amp_ratio") for r in rows]
    ratio21 = [_val(r, dir_21, "meantrace10_amp_ratio") for r in rows]
    dpk12 = [_val(r, dir_12, "meantrace10_delay_ms") for r in rows]
    dpk21 = [_val(r, dir_21, "meantrace10_delay_ms") for r in rows]
    d1012 = [_val(r, dir_12, "meantrace10_delay_10_ms") for r in rows]
    d1021 = [_val(r, dir_21, "meantrace10_delay_10_ms") for r in rows]
    n_ratio = sum(v is not None for v in ratio12 + ratio21)
    n_del = sum(v is not None for v in dpk12 + dpk21 + d1012 + d1021)
    tag = "PRIMARY" if near_vm is None else f"nearest baseline {near_vm:g} mV"
    print(
        f"  spikelet meantrace10 vs time ({tag}): "
        f"files={len(rows)}, amp_ratio={n_ratio}, delays={n_del}"
    )

    plt = _get_agg_plt()
    fig, axes = plt.subplots(3, 1, figsize=(11, 10), sharex=True)
    if title:
        suffix = (
            "mean-trace 15 ms ratio and delays vs recording time"
            if near_vm is None
            else f"mean-trace 15 ms ratio and delays vs recording time (nearest {near_vm:g} mV)"
        )
        fig.suptitle(f"{title}  —  {suffix}", fontsize=12)

    panels = (
        (axes[0], "Amplitude ratio (meantrace10; ratio>1 omitted)",
         ((ratio12, "o-", "C0", "12 (ch0→ch2)"), (ratio21, "s-", "C1", "21 (ch2→ch0)"))),
        (axes[1], "Delay peak (ms)",
         ((dpk12, "o-", "C0", "12 peak"), (dpk21, "s-", "C1", "21 peak"))),
        (axes[2], "Delay 10% (ms)",
         ((d1012, "o-", "C0", "12 10%"), (d1021, "s-", "C1", "21 10%"))),
    )
    n_total = 0
    labeled_neg = False
    for ax, panel_title, series in panels:
        n = 0
        is_delay = "Delay" in panel_title
        for vals, style, color, label in series:
            n += _plot_timed(ax, times, vals, style, color=color, label=label)
            if is_delay:
                xs, ys = _finite_xy(times, vals)
                if _mark_negative_delay_points(ax, xs, ys, already_labeled=labeled_neg):
                    labeled_neg = True
        n_total += n
        ax.set_title(panel_title)
        ax.set_ylabel("ratio" if "ratio" in panel_title.lower() else panel_title.lower())
        ax.grid(True, alpha=0.3)
        if n == 0:
            ax.text(0.5, 0.5, "no data", transform=ax.transAxes,
                    ha="center", va="center", color="0.5", fontsize=9)
        else:
            ax.legend(loc="best", fontsize=8)
    axes[0].set_ylabel("ratio")
    axes[1].set_ylabel("delay peak (ms)")
    axes[2].set_ylabel("delay 10% (ms)")
    axes[-1].set_xlabel("Recording time" if use_dates else "File order")
    if n_total == 0:
        plt.close(fig)
        return None
    if use_dates:
        _format_time_axes(axes[-1])
        fig.autofmt_xdate()
    return _finish_folder_fig(fig, out_path)


def save_folder_spikelet_meantrace10_vs_vm_plot(
    sweep_rows, out_path, title=None, summary_rows=None,
):
    """Parallel mean-trace scheme: sweep-mean ratio/delays vs spikelet baseline."""
    from matplotlib.colors import Normalize

    plt = _get_agg_plt()
    fig, axes = plt.subplots(3, 2, figsize=(14, 11), sharex="col")
    if title:
        fig.suptitle(f"{title}  —  mean-trace 15 ms spikelet / spike vs mean spikelet baseline", fontsize=12)

    file_pos, first_lbl, last_lbl, n_files = _cc_vm_file_color_map(
        list(sweep_rows or []), summary_rows=summary_rows,
    )
    cmap = _mpl_cmap()
    panels = (
        ("meantrace10_amp_ratio", "meantrace10 amp spikelet / amp spike"),
        ("meantrace10_delay_ms", "meantrace10 delay peak (ms)"),
        ("meantrace10_delay_10_ms", "meantrace10 delay 10% (ms)"),
    )
    directions = (("ch0->ch2", "ch0→ch2"), ("ch2->ch0", "ch2→ch0"))
    any_data = False
    n_curves = 0
    for col, (direction, dir_title) in enumerate(directions):
        for row_i, (y_key, ylabel) in enumerate(panels):
            ax = axes[row_i, col]
            curves = file_spikelet_vm_curves(sweep_rows, direction, y_key, ap_rows=None)
            n_curves += len(curves)
            if not curves:
                ax.set_title(f"{dir_title} — no data" if row_i == 0 else "")
                if col == 0:
                    ax.set_ylabel(ylabel)
                ax.grid(True, alpha=0.3)
                ax.text(0.5, 0.5, "no data", transform=ax.transAxes,
                        ha="center", va="center", color="0.5", fontsize=9)
                continue
            any_data = True
            file_xy = []
            labeled_neg = False
            for fname, _dt, vms, ys in curves:
                color = cmap(float(file_pos.get(fname, file_pos.get(os.path.basename(str(fname)), 0.5))))
                ax.scatter(vms, ys, color=[color], s=28, zorder=3, alpha=0.9)
                if "delay" in y_key:
                    if _mark_negative_delay_points(ax, vms, ys, already_labeled=labeled_neg):
                        labeled_neg = True
                if CC_VM_FIT_LINEAR:
                    x1, y1, _s, _b, _r2 = _linear_cc_vs_vm(vms, ys)
                    if x1 is not None:
                        ax.plot(x1, y1, "-", color=color, lw=1.3, alpha=0.85)
                file_xy.append((vms, ys))
            _plot_mean_of_file_lines(ax, file_xy)
            ax.grid(True, alpha=0.3)
            if row_i == 0:
                ax.set_title(f"{dir_title} — {len(curves)} file(s)")
            if col == 0:
                ax.set_ylabel(ylabel)
            if row_i == 2:
                ax.set_xlabel("Mean spikelet baseline (passive, mV)")
    if not any_data:
        plt.close(fig)
        print("  spikelet meantrace10 vs Vm: no points yet")
        return None
    try:
        sm = plt.cm.ScalarMappable(cmap=cmap, norm=Normalize(0.0, 1.0))
        sm.set_array([])
        cbar = fig.colorbar(sm, ax=list(axes.ravel()), fraction=0.046, pad=0.03)
        cbar.set_label("first file  →  last file")
        try:
            cbar.set_ticks([0.0, 1.0], labels=[f"first\n{first_lbl}", f"last\n{last_lbl}"])
        except TypeError:
            cbar.set_ticks([0.0, 1.0])
            cbar.ax.set_yticklabels([f"first  {first_lbl}", f"last  {last_lbl}"])
        cbar.ax.tick_params(labelsize=8)
        used_cbar = True
    except Exception:
        used_cbar = False
    print(f"  spikelet meantrace10 vs Vm curves: {n_curves}")
    return _finish_folder_fig(fig, out_path, skip_tight=used_cbar)


def file_cc_norm_curves(all_rows, direction):
    """
    Per-file (Vm, CC_norm) curves for one direction, sorted by Vm within file.

    Returns list of (file_name, recording_datetime, vm_list, cc_norm_list).
    """
    by_file = {}
    file_dt = {}
    for row in all_rows:
        if row.get("direction") != direction:
            continue
        cc_n = row.get("CC_norm")
        vm = row.get("Vm_active_stim_mV")
        if cc_n is None or vm is None:
            continue
        fname = row.get("file") or "?"
        by_file.setdefault(fname, []).append((float(vm), float(cc_n)))
        if fname not in file_dt:
            file_dt[fname] = _parse_recording_datetime(row.get("recording_datetime"))

    curves = []
    for fname, pairs in by_file.items():
        pairs.sort(key=lambda p: p[0])
        vms = [p[0] for p in pairs]
        norms = [p[1] for p in pairs]
        curves.append((fname, file_dt.get(fname), vms, norms))
    curves.sort(key=lambda c: (c[1] is None, c[1] or datetime.min, c[0]))
    return curves


def _abf_file_number(fname):
    """Last integer in the ABF stem (Clampex run number), or -1."""
    stem = os.path.splitext(os.path.basename(str(fname)))[0]
    nums = re.findall(r"\d+", stem)
    return int(nums[-1]) if nums else -1


def _cc_vm_file_color_map(all_rows, summary_rows=None):
    """
    Color 0..1 by Clampex file number (same file → same color on both panels).

    Uses every file in All_data and File_summary, not only those with CC_norm.
    Returns (fname -> pos, first_label, last_label, n_files).
    """
    seen = {}
    for src in (all_rows, summary_rows or []):
        for row in src:
            fname = row.get("file")
            if not fname:
                continue
            key = os.path.basename(str(fname))
            if key not in seen:
                seen[key] = _parse_recording_datetime(row.get("recording_datetime"))
    items = sorted(seen.items(), key=lambda it: (_abf_file_number(it[0]), it[0]))
    n = max(len(items) - 1, 1)
    pos = {fname: (i / n) for i, (fname, _dt) in enumerate(items)}
    # also map full paths if some rows stored them
    for row_src in (all_rows, summary_rows or []):
        for row in row_src:
            fname = row.get("file")
            if fname:
                pos[str(fname)] = pos.get(os.path.basename(str(fname)), 0.5)

    def _lbl(fname):
        return os.path.splitext(os.path.basename(str(fname)))[0]

    first_lbl = _lbl(items[0][0]) if items else ""
    last_lbl = _lbl(items[-1][0]) if items else ""
    return pos, first_lbl, last_lbl, len(items)


def _poly_cc_vs_vm(vm, cc, degree, n_grid=80):
    """Return (x_fit, y_fit, coeffs, r2) or Nones if the poly is not defined."""
    vm = np.asarray(vm, dtype=float)
    cc = np.asarray(cc, dtype=float)
    ok = np.isfinite(vm) & np.isfinite(cc)
    vm, cc = vm[ok], cc[ok]
    if len(vm) < degree + 1 or float(np.ptp(vm)) < 1e-9:
        return None, None, None, None
    if len(np.unique(np.round(vm, 6))) < degree + 1:
        return None, None, None, None
    coeffs = np.polyfit(vm, cc, degree)
    pred = np.polyval(coeffs, vm)
    ss_res = float(np.sum((cc - pred) ** 2))
    ss_tot = float(np.sum((cc - np.mean(cc)) ** 2))
    r2 = 1.0 - ss_res / ss_tot if ss_tot > 0 else 1.0
    x = np.linspace(float(vm.min()), float(vm.max()), n_grid)
    y = np.polyval(coeffs, x)
    return x, y, coeffs, float(r2)


def _linear_cc_vs_vm(vm, cc, n_grid=80):
    """Return (x_fit, y_fit, slope, intercept, r2) or Nones if a line is not defined."""
    vm = np.asarray(vm, dtype=float)
    cc = np.asarray(cc, dtype=float)
    ok = np.isfinite(vm) & np.isfinite(cc)
    if int(np.sum(ok)) < CC_VM_LINEAR_MIN_POINTS:
        return None, None, None, None, None
    x, y, coeffs, r2 = _poly_cc_vs_vm(vm, cc, 1, n_grid=n_grid)
    if x is None:
        return None, None, None, None, None
    return x, y, float(coeffs[0]), float(coeffs[1]), r2


def _binned_mean_xy(vm, cc, n_bins=8):
    """Mean CC vs mean Vm in equal-width Vm bins (line through all files)."""
    vm = np.asarray(vm, dtype=float)
    cc = np.asarray(cc, dtype=float)
    ok = np.isfinite(vm) & np.isfinite(cc)
    vm, cc = vm[ok], cc[ok]
    if len(vm) < 2:
        return None, None
    vmin, vmax = float(np.min(vm)), float(np.max(vm))
    if vmax - vmin < 1e-9:
        return [vmin], [float(np.mean(cc))]
    n_bins = int(max(3, min(n_bins, max(3, len(vm) // 3))))
    edges = np.linspace(vmin, vmax, n_bins + 1)
    xs, ys = [], []
    for i in range(n_bins):
        left, right = edges[i], edges[i + 1]
        mask = (vm >= left) & (vm <= right) if i == n_bins - 1 else (vm >= left) & (vm < right)
        if np.any(mask):
            xs.append(float(np.mean(vm[mask])))
            ys.append(float(np.mean(cc[mask])))
    if len(xs) < 2:
        return None, None
    return xs, ys


def _mean_line_from_file_fits(file_curves, n_grid=80):
    """
    One line whose slope is the mean of per-file linear slopes.

    Each file with a valid fit counts once (not weighted by how many
    sweeps it has). The line is placed through the mean of those files'
    (Vm, y) centroids so it shows the typical angle of the colored lines.
    """
    slopes = []
    cx, cy = [], []
    xmin, xmax = np.inf, -np.inf
    for vms, ys in file_curves or []:
        _x, _y, slope, _b, _r2 = _linear_cc_vs_vm(vms, ys)
        if slope is None:
            continue
        vm = np.asarray(vms, dtype=float)
        yy = np.asarray(ys, dtype=float)
        ok = np.isfinite(vm) & np.isfinite(yy)
        if not np.any(ok):
            continue
        slopes.append(float(slope))
        cx.append(float(np.mean(vm[ok])))
        cy.append(float(np.mean(yy[ok])))
        xmin = min(xmin, float(np.min(vm[ok])))
        xmax = max(xmax, float(np.max(vm[ok])))
    if not slopes or not np.isfinite(xmin) or xmax <= xmin:
        return None, None, None, None
    slope = float(np.mean(slopes))
    x0 = float(np.mean(cx))
    y0 = float(np.mean(cy))
    intercept = y0 - slope * x0
    x = np.linspace(xmin, xmax, int(n_grid))
    y = slope * x + intercept
    return x, y, slope, intercept


def _plot_mean_of_file_lines(ax, file_curves):
    """Black line = mean slope of the per-file linear fits."""
    x1, y1, slope, intercept = _mean_line_from_file_fits(file_curves)
    if x1 is None:
        return None, None
    ax.plot(
        x1, y1, "-", color="black", lw=2.5, zorder=5, alpha=0.95,
        label="mean of file lines",
    )
    ax.legend(loc="best", fontsize=8)
    return slope, intercept


def _plot_all_files_cc_vm_mean(ax, vms, norms):
    """Black linear mean across every file's CC_norm vs Vm points."""
    x1, y1, slope, intercept, r2 = _linear_cc_vs_vm(vms, norms)
    if x1 is not None:
        ax.plot(
            x1, y1, "-", color="black", lw=2.5, zorder=5, alpha=0.95,
            label="mean (all files)",
        )
        ax.legend(loc="best", fontsize=8)
        return slope, intercept, r2
    xb, yb = _binned_mean_xy(vms, norms)
    if xb is not None:
        ax.plot(
            xb, yb, "D-", color="black", lw=2.0, ms=5, zorder=5,
            label="mean (all files)",
        )
        ax.legend(loc="best", fontsize=8)
    return None, None, None


def save_folder_cc_norm_vs_vm_plot(all_rows, out_path, title=None, summary_rows=None):
    """
    Folder overview: CC_norm vs Vm, color = first→last ABF file number.

    Linear fit only. Two panels: CC12 (ch0→ch2) and CC21 (ch2→ch0).
    """
    from matplotlib.colors import Normalize

    plt = _get_agg_plt()
    fig, axes = plt.subplots(1, 2, figsize=(14, 5.8), sharey=True)
    if title:
        fig.suptitle(title, fontsize=12)

    file_pos, first_lbl, last_lbl, n_files = _cc_vm_file_color_map(
        all_rows, summary_rows=summary_rows
    )
    print(f"  CC_norm vs Vm color: first={first_lbl}  →  last={last_lbl}  ({n_files} files)")
    cmap = _mpl_cmap()

    panels = (
        ("ch0->ch2", "CC12 (ch0→ch2)"),
        ("ch2->ch0", "CC21 (ch2→ch0)"),
    )
    any_data = False
    for ax, (direction, panel_title) in zip(axes, panels):
        curves = file_cc_norm_curves(all_rows, direction)
        if not curves:
            ax.set_title(f"{panel_title} — no data")
            ax.grid(True, alpha=0.3)
            continue
        any_data = True

        all_vm, all_cc = [], []
        for fname, _dt, vms, norms in curves:
            color = cmap(float(file_pos.get(fname, file_pos.get(os.path.basename(str(fname)), 0.5))))
            ax.scatter(vms, norms, color=[color], s=28, zorder=3, alpha=0.9)
            if CC_VM_FIT_LINEAR:
                x1, y1, _s, _b, _r2 = _linear_cc_vs_vm(vms, norms)
                if x1 is not None:
                    ax.plot(x1, y1, "-", color=color, lw=1.3, alpha=0.85)
            all_vm.extend(vms)
            all_cc.extend(norms)
        _plot_all_files_cc_vm_mean(ax, all_vm, all_cc)

        ax.axhline(1.0, color="0.45", ls=":", lw=0.8, alpha=0.7)
        ax.set_xlabel("Vm active during stim (mV)")
        ax.set_title(f"{panel_title} — {len(curves)} file(s) with CC_norm")
        ax.grid(True, alpha=0.3)

    if not any_data:
        plt.close(fig)
        print("  CC_norm vs Vm: no CC_norm + Vm_active_stim_mV points")
        return None

    axes[0].set_ylabel("CC_norm (CC / file mean)")
    used_cbar = False
    try:
        sm = plt.cm.ScalarMappable(cmap=cmap, norm=Normalize(0.0, 1.0))
        sm.set_array([])
        cbar = fig.colorbar(sm, ax=list(axes), fraction=0.046, pad=0.03)
        cbar.set_label("first file  →  last file")
        try:
            cbar.set_ticks([0.0, 1.0], labels=[f"first\n{first_lbl}", f"last\n{last_lbl}"])
        except TypeError:
            cbar.set_ticks([0.0, 1.0])
            cbar.ax.set_yticklabels([f"first  {first_lbl}", f"last  {last_lbl}"])
        cbar.ax.tick_params(labelsize=8)
        used_cbar = True
    except Exception as exc:
        print(f"  CC_norm vs Vm colorbar skipped: {exc}")

    return _finish_folder_fig(fig, out_path, skip_tight=used_cbar)


def cc_vm_fits_for_direction(all_rows, direction):
    """Per-file linear CC_norm vs Vm: slope (1/mV), intercept, R². Needs >=3 points."""
    out = []
    for fname, dt, vms, norms in file_cc_norm_curves(all_rows, direction):
        _x, _y, slope, intercept, r2 = _linear_cc_vs_vm(vms, norms)
        out.append({
            "file": fname,
            "recording_datetime": dt,
            "direction": direction,
            "n": len(vms),
            "slope": None if slope is None else round(float(slope), 6),
            "intercept": None if intercept is None else round(float(intercept), 6),
            "r2": None if r2 is None else round(float(r2), 4),
        })
    return out


def attach_cc_vm_slopes(summary_rows, all_rows):
    """Add CC12/CC21 vs Vm slope, intercept, R² to File_summary rows."""
    by_file = {}
    for direction, tag in (("ch0->ch2", "12"), ("ch2->ch0", "21")):
        for rec in cc_vm_fits_for_direction(all_rows, direction):
            key = os.path.basename(str(rec["file"]))
            by_file.setdefault(key, {})[tag] = rec
    for row in summary_rows:
        key = os.path.basename(str(row.get("file") or ""))
        fits = by_file.get(key, {})
        for tag in ("12", "21"):
            rec = fits.get(tag) or {}
            row[f"CC{tag}_vs_Vm_slope"] = rec.get("slope")
            row[f"CC{tag}_vs_Vm_intercept"] = rec.get("intercept")
            row[f"CC{tag}_vs_Vm_R2"] = rec.get("r2")
            row[f"n_CC_vs_Vm_{tag}"] = rec.get("n")
    return summary_rows


def save_folder_cc_vm_slope_over_time_plot(
    all_rows, out_path, title=None, summary_rows=None,
):
    """
    Slope of per-file CC_norm vs Vm linear fit, vs recording time.

    Same fit as the colored lines on CC_norm vs Vm (>=3 points).
    """
    dt_by_file = {}
    for row in summary_rows or []:
        fname = row.get("file")
        if not fname:
            continue
        dt = _parse_recording_datetime(row.get("recording_datetime"))
        dt_by_file[os.path.basename(str(fname))] = dt
        dt_by_file[str(fname)] = dt

    plt = _get_agg_plt()
    fig, axes = plt.subplots(2, 1, figsize=(11, 7.5), sharex=True)
    if title:
        fig.suptitle(f"{title}  —  CC vs Vm slope over files", fontsize=12)

    panels = (
        ("ch0->ch2", "CC12 vs Vm slope (ch0→ch2)"),
        ("ch2->ch0", "CC21 vs Vm slope (ch2→ch0)"),
    )
    n_total = 0
    for ax, (direction, panel_title) in zip(axes, panels):
        times, slopes, r2s = [], [], []
        for rec in cc_vm_fits_for_direction(all_rows, direction):
            if rec.get("slope") is None:
                continue
            dt = rec.get("recording_datetime")
            if dt is None:
                dt = dt_by_file.get(os.path.basename(str(rec["file"])))
                if dt is None:
                    dt = dt_by_file.get(str(rec["file"]))
            if dt is None:
                continue
            times.append(dt)
            slopes.append(rec["slope"])
            r2s.append(rec.get("r2"))
        n = _plot_timed(ax, times, slopes, "o-", color="C0", label="slope")
        n_total += n
        ax.axhline(0.0, color="0.45", ls=":", lw=0.8)
        ax.set_title(panel_title if n else f"{panel_title} — no data")
        ax.set_ylabel("slope (CC_norm / mV)")
        ax.grid(True, alpha=0.3)
        if n == 0:
            ax.text(0.5, 0.5, "no data", transform=ax.transAxes,
                    ha="center", va="center", color="0.5", fontsize=9)
        else:
            ax2 = ax.twinx()
            _plot_timed(ax2, times, r2s, "s--", color="0.4", label="R²")
            ax2.set_ylabel("R²", color="0.4")
            ax2.tick_params(axis="y", labelcolor="0.4")
            ax2.set_ylim(-0.05, 1.05)
            lines = ax.get_lines() + ax2.get_lines()
            ax.legend(lines, [ln.get_label() for ln in lines], loc="best", fontsize=8)
    axes[-1].set_xlabel("Recording time")

    if n_total == 0:
        plt.close(fig)
        print("  CC vs Vm slope over time: no file with >=3 CC_norm points")
        return None

    _format_time_axes(axes[-1])
    fig.autofmt_xdate()
    try:
        fig.tight_layout()
    except Exception:
        pass
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    _savefig_white(fig, out_path)
    plt.close(fig)
    print(f"  saved {out_path}  (slope points: {n_total})")
    return out_path


def format_excel_header_wrap(workbook, header_row=1, min_width=12, max_width=18):
    """Wrap header text to the column width and raise row 1 so names are fully visible."""
    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")
    fill_primary = PatternFill(fill_type="solid", fgColor="DDEBF7")
    fill_mode = PatternFill(fill_type="solid", fgColor="FCE4D6")
    for ws in workbook.worksheets:
        max_lines = 1
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(header_row, col)
            cell.alignment = wrap
            text = "" if cell.value is None else str(cell.value)
            letter = get_column_letter(col)
            width = min(max_width, max(min_width, len(text) + 1))
            ws.column_dimensions[letter].width = width
            if ws.title == "Summary_short":
                if text.startswith("primary_"):
                    for row_i in range(1, ws.max_row + 1):
                        ws.cell(row_i, col).fill = fill_primary
                elif text.startswith("near_mode_vm_"):
                    for row_i in range(1, ws.max_row + 1):
                        ws.cell(row_i, col).fill = fill_mode
            chars_per_line = max(8, int(width))
            n_lines = text.count("\n") + 1
            if "\n" not in text:
                n_lines = max(1, (len(text) + chars_per_line - 1) // chars_per_line)
            max_lines = max(max_lines, n_lines)
        ws.row_dimensions[header_row].height = max(30, 15 * max_lines + 8)


def _excel_cell(val):
    """openpyxl-safe cell: no numpy types, no inf, bool → Yes/No."""
    if val is None:
        return None
    if isinstance(val, (bool, np.bool_)):
        return "Yes" if val else "No"
    if isinstance(val, (np.integer,)):
        return int(val)
    if isinstance(val, (float, np.floating)):
        fv = float(val)
        return fv if np.isfinite(fv) else None
    return val


def _excel_dataframe(rows):
    import pandas as pd

    if not rows:
        return pd.DataFrame()
    clean = [{k: _excel_cell(v) for k, v in dict(row).items()} for row in rows]
    return pd.DataFrame(clean)


def _compact_excel_summary_rows(summary_rows, spikelet_sweep_rows=None, spikelet_rows=None):
    """Short user-facing cut from File_summary while keeping the full sheets."""

    def _add_spikelet_block(dst, prefix, picked, tag):
        dst[f"{prefix}_{tag}_sweep"] = picked.get("sweep") if picked else None
        dst[f"{prefix}_{tag}_baseline_mV"] = _sweep_row_metric(picked, "baseline_mV")
        dst[f"{prefix}_{tag}_amp_spikelet_mV"] = (
            _finite_number(picked.get("meantrace10_amp_spikelet_mV")) if picked else None
        )
        dst[f"{prefix}_{tag}_delay_peak_ms"] = (
            _finite_number(picked.get("meantrace10_delay_ms")) if picked else None
        )
        dst[f"{prefix}_{tag}_amp_ratio"] = (
            _finite_number(picked.get("meantrace10_amp_ratio")) if picked else None
        )
        dst[f"{prefix}_{tag}_detected"] = picked.get("meantrace10_detected") if picked else None
        dst[f"{prefix}_{tag}_skip_reason"] = picked.get("meantrace10_skip_reason") if picked else None

    rows = []
    mode_vm = spikelet_baseline_mode_vm(spikelet_sweep_rows, ap_rows=spikelet_rows)
    for src in summary_rows or []:
        row = {
            "file": src.get("file"),
            "recording_datetime": src.get("recording_datetime"),
            "file_skip_reason": src.get("file_skip_reason"),
            "analysis_blocks": src.get("analysis_blocks"),
            "CC_selection_mode": src.get("CC_selection_mode"),
            "CC_selection_note": src.get("CC_selection_note"),
            "CC12": src.get("CC12"),
            "CC21": src.get("CC21"),
            "Gj12_nS": src.get("Gj12_nS"),
            "Gj21_nS": src.get("Gj21_nS"),
            "Rin_ch0_MOhm": src.get("Rin_ch0_MOhm"),
            "Rin_ch2_MOhm": src.get("Rin_ch2_MOhm"),
            "V_rest_mV_ch0": src.get("V_rest_mV_ch0"),
            "V_rest_mV_ch2": src.get("V_rest_mV_ch2"),
            "AP21_ratio_ch0": src.get("AP21_ratio_ch0"),
            "AP21_ratio_ch2": src.get("AP21_ratio_ch2"),
            "FWHM_ms_ch0": src.get("FWHM_ms_ch0"),
            "FWHM_ms_ch2": src.get("FWHM_ms_ch2"),
            "props_sweep_ch0": src.get("props_sweep_ch0"),
            "props_sweep_ch2": src.get("props_sweep_ch2"),
            "inj_current_pA_ch0": src.get("inj_current_pA_ch0"),
            "inj_current_pA_ch2": src.get("inj_current_pA_ch2"),
            "tau_ms_ch0": src.get("tau_ms_ch0"),
            "tau_ms_ch2": src.get("tau_ms_ch2"),
            "Cm_pF_ch0": src.get("Cm_pF_ch0"),
            "Cm_pF_ch2": src.get("Cm_pF_ch2"),
            "tau_sweep_ch0": src.get("tau_sweep_ch0"),
            "tau_sweep_ch2": src.get("tau_sweep_ch2"),
            "delta_V_mV_ch0": src.get("delta_V_mV_ch0"),
            "delta_V_mV_ch2": src.get("delta_V_mV_ch2"),
            "V_post_min_mV_ch0": src.get("V_post_min_mV_ch0"),
            "V_post_min_mV_ch2": src.get("V_post_min_mV_ch2"),
            "mode_spikelet_baseline_mV": mode_vm,
        }
        fname = src.get("file")
        for direction, tag in (("ch0->ch2", "12"), ("ch2->ch0", "21")):
            primary = _primary_sweep_row(spikelet_sweep_rows, fname, direction)
            nearest = _sweep_closest_to_vm(
                spikelet_sweep_rows, fname, direction, mode_vm, ap_rows=spikelet_rows,
            )
            _add_spikelet_block(row, "primary", primary, tag)
            _add_spikelet_block(row, "near_mode_vm", nearest, tag)
        rows.append(row)
    return rows


def save_batch_excel(
    path,
    all_rows,
    summary_rows,
    spikelet_rows=None,
    spikelet_sweep_rows=None,
):
    """Write full sheets plus a compact summary cut for quick reading."""
    import pandas as pd

    path = os.path.abspath(path)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    compact_rows = _compact_excel_summary_rows(
        summary_rows or [],
        spikelet_sweep_rows=spikelet_sweep_rows or [],
        spikelet_rows=spikelet_rows or [],
    )
    sheets = (
        ("Summary_short", compact_rows),
        ("File_summary", summary_rows or []),
        ("All_data", all_rows or []),
        ("Spikelets", spikelet_rows or []),
        ("Spikelet_sweeps", spikelet_sweep_rows or []),
    )

    def _write(out_path):
        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            for name, rows in sheets:
                _excel_dataframe(rows).to_excel(writer, sheet_name=name, index=False)
            try:
                format_excel_header_wrap(writer.book)
            except Exception as exc:
                print(f"  Excel header wrap skipped: {exc}")

    try:
        _write(path)
        used = path
    except PermissionError:
        stem, ext = os.path.splitext(path)
        used = f"{stem}_{datetime.now().strftime('%H%M%S')}{ext or '.xlsx'}"
        print(f"  Excel file is open or locked; writing {used}")
        _write(used)

    size = os.path.getsize(used) if os.path.isfile(used) else 0
    print(f"Saved Excel: {used}")
    print(f"  size={size} bytes")
    for name, rows in sheets:
        print(f"  {name} rows: {len(rows)}")
    return used


def build_file_summary_row(
    name,
    rec_dt,
    rin_ch0,
    rin_ch2,
    rin_note_0,
    rin_note_2,
    rin_skip_0,
    rin_skip_2,
    cc_mean_02,
    cc_mean_20,
    n_avg_02,
    n_avg_20,
    gj_02,
    gj_20,
    gj_skip_02,
    gj_skip_20,
    n_sel_02,
    n_sel_20,
    n_neg_02=0,
    n_neg_20=0,
    cell_props=None,
    file_skip_reason=None,
):
    """One row per file for Excel sheet File_summary (CC12/Gj12 = ch0->ch2)."""
    row = {
        "file": name,
        "recording_datetime": rec_dt,
        "Rin1_MOhm": rin_ch0,
        "Rin2_MOhm": rin_ch2,
        "Rin1_note": rin_note_0,
        "Rin2_note": rin_note_2,
        "Rin1_skip_reason": rin_skip_0,
        "Rin2_skip_reason": rin_skip_2,
        "CC12": cc_mean_02,
        "CC21": cc_mean_20,
        "Gj12_nS": gj_02,
        "Gj21_nS": gj_20,
        "Gj12_skip_reason": gj_skip_02,
        "Gj21_skip_reason": gj_skip_20,
        "n_CC_avg_12": n_avg_02,
        "n_CC_avg_21": n_avg_20,
        "n_CC_neg_12": n_neg_02,
        "n_CC_neg_21": n_neg_20,
        "CC_used_neg_12": "Yes" if n_neg_02 > 0 else "No",
        "CC_used_neg_21": "Yes" if n_neg_20 > 0 else "No",
        "CC_neg_of_avg_12": f"{n_neg_02}/{n_avg_02}",
        "CC_neg_of_avg_21": f"{n_neg_20}/{n_avg_20}",
        "n_CC_selected_12": n_sel_02,
        "n_CC_selected_21": n_sel_20,
    }
    if cell_props:
        row.update(cell_props)
    if file_skip_reason:
        row["file_skip_reason"] = file_skip_reason
    return row


def _skipped_file_row(name, rec_dt, reason):
    return [sweep_only_row({
        "file": name,
        "recording_datetime": rec_dt,
        "sweep": None,
        "direction": None,
        "cur_step_pA": None,
        "delta_V_active_mV": None,
        "delta_V_passive_mV": None,
        "CC": None,
        "CC_skip_reason": reason,
        "Gj_sweep_nS": None,
        "Gj_sweep_skip_reason": reason,
    })]


def analyze_abf_file(
    filepath, plots_dir=None, cc_plots_dir_path=None, spikelet_plots_dir_path=None,
    blocks=None,
):
    """Return (per_sweep_rows, file_summary_row, qc_plot_paths, spikelet_ap_rows, spikelet_sweep_rows)."""
    bsel = resolve_analysis_blocks(
        blocks if blocks is not None else ensure_analysis_blocks()
    )
    name = os.path.basename(filepath)
    abf = pyabf.ABF(filepath)
    rec_dt = recording_datetime_str(abf)
    plot_paths = []

    if abf.channelCount != EXPECTED_CHANNELS:
        reason = f"expected {EXPECTED_CHANNELS} channels, got {abf.channelCount}"
        empty_props = {
            **empty_cell_props_fields(),
            **empty_tau_cm_fields(),
            **empty_spikelet_summary_fields(),
        }
        return (
            _skipped_file_row(name, rec_dt, reason),
            build_file_summary_row(
                name, rec_dt, None, None, None, None, reason, reason,
                None, None, 0, 0, None, None, reason, reason, 0, 0,
                cell_props=empty_props,
                file_skip_reason=reason,
            ),
            plot_paths,
            [],
            [],
        )

    cell_props = empty_cell_props_fields()
    if bsel["cell_props"]:
        try:
            cell_props = cell_properties_for_file(abf)
        except Exception as exc:
            msg = str(exc)
            cell_props["props_skip_reason_ch0"] = msg
            cell_props["props_skip_reason_ch2"] = msg

    b = time_period_borders(abf.dataRate)
    w0 = (b["start10"], b["end10"], b["start11"], b["end11"])
    w2 = (b["start20"], b["end20"], b["start21"], b["end21"])
    cc_mode = selected_cc_mode(bsel)
    cc_enabled = cc_mode is not None
    cc_selection_mode = CC_MODE_SELECTION_LABELS.get(cc_mode)
    cc_selection_note = CC_MODE_SELECTION_NOTES.get(cc_mode)

    sweeps_ch0, sweeps_ch2 = [], []
    block_02, block_20 = [], []
    cc_mean_02 = cc_mean_20 = None
    n_avg_02 = n_avg_20 = n_neg_02 = n_neg_20 = 0
    rin_ch0 = rin_ch2 = r2_ch0 = r2_ch2 = None
    n0 = n2 = 0
    rin_skip_0 = rin_skip_2 = None
    rin_range_0 = rin_range_2 = None
    rin_note_0 = rin_note_2 = None

    if cc_enabled:
        sweeps_ch0 = cc_select_sweeps_direction(
            abf, 0, 2, 1, b["start10"], b["end10"], b["start11"], b["end11"], mode_key=cc_mode,
        )
        sweeps_ch2 = cc_select_sweeps_direction(
            abf, 2, 0, 3, b["start20"], b["end20"], b["start21"], b["end21"], mode_key=cc_mode,
        )
        block_02 = coupling_block(abf, sweeps_ch0, 0, 2, 1, w0)
        block_20 = coupling_block(abf, sweeps_ch2, 2, 0, 3, w2)
        for r in block_02:
            r["CC_selection_mode"] = cc_selection_mode
        for r in block_20:
            r["CC_selection_mode"] = cc_selection_mode
        cc_list_02 = [r["CC"] for r in block_02]
        cc_list_20 = [r["CC"] for r in block_20]
        cc_mean_02 = mean_valid_cc(cc_list_02)
        cc_mean_20 = mean_valid_cc(cc_list_20)
        cc_normalize_block(block_02, cc_mean_02)
        cc_normalize_block(block_20, cc_mean_20)
        n_avg_02 = n_valid_cc(cc_list_02)
        n_avg_20 = n_valid_cc(cc_list_20)
        n_neg_02 = n_negative_cc(cc_list_02)
        n_neg_20 = n_negative_cc(cc_list_20)

    if cc_enabled or bsel["tau_cm"]:
        rin_ch0, r2_ch0, n0, rin_skip_0, rin_range_0, rin_note_0, _ = rin_for_channel(
            abf, 0, 1, b["start10"], b["end10"], b["start11"], b["end11"], b["Rtime_ch0"]
        )
        rin_ch2, r2_ch2, n2, rin_skip_2, rin_range_2, rin_note_2, _ = rin_for_channel(
            abf, 2, 3, b["start20"], b["end20"], b["start21"], b["end21"], b["Rtime_ch2"]
        )
    elif not cc_enabled:
        rin_skip_0 = rin_skip_2 = "CC / Rin block not selected"
        rin_note_0 = rin_note_2 = "CC / Rin block not selected"

    tau_cm = empty_tau_cm_fields()
    tau_plot_meta = {}
    if bsel["tau_cm"]:
        try:
            tau_cm_raw = tau_cm_for_file(abf, rin_ch0, rin_ch2, borders=b)
            tau_plot_meta = tau_cm_raw.pop("_tau_plot_meta", {})
            tau_cm = {k: v for k, v in tau_cm_raw.items() if not k.startswith("_")}
        except Exception as exc:
            msg = str(exc)
            tau_cm["tau_skip_reason_ch0"] = msg
            tau_cm["tau_skip_reason_ch2"] = msg
    else:
        tau_cm["tau_skip_reason_ch0"] = "Tau/Cm block not selected"
        tau_cm["tau_skip_reason_ch2"] = "Tau/Cm block not selected"

    spikelet_rows = []
    spikelet_sweep_rows = []
    spikelet_summary = empty_spikelet_summary_fields()
    if bsel["spikelets"]:
        try:
            sp_dir = spikelet_plots_dir_path
            if SAVE_SPIKELET_PLOTS and not sp_dir:
                sp_dir = os.path.join(os.path.dirname(os.path.abspath(filepath)), SPIKELET_PLOTS_SUBDIR)
            sp_result = spikelets_for_file(
                abf, name, rec_dt,
                plots_dir=sp_dir if SAVE_SPIKELET_PLOTS else None,
                stem=_abf_stem(filepath),
            )
            spikelet_rows, spikelet_summary, sp_paths = sp_result[0], sp_result[1], sp_result[2]
            spikelet_sweep_rows = sp_result[3] if len(sp_result) > 3 else []
            plot_paths.extend(sp_paths)
        except Exception as exc:
            print(f"  Spikelet analysis error: {exc}")
            traceback.print_exc()
            spikelet_summary = empty_spikelet_summary_fields()
            spikelet_summary["spikelet_skip_reason_12"] = str(exc)
            spikelet_summary["spikelet_skip_reason_21"] = str(exc)
    else:
        spikelet_summary["spikelet_skip_reason_12"] = "Spikelet block not selected"
        spikelet_summary["spikelet_skip_reason_21"] = "Spikelet block not selected"

    file_summary_extra = {
        **cell_props,
        **tau_cm,
        **spikelet_summary,
        "Rin_ch0_MOhm": rin_ch0,
        "R2_ch0": r2_ch0,
        "n_IV_ch0": n0,
        "Rin_ch0_skip_reason": rin_skip_0,
        "Rin_ch0_note": rin_note_0,
        "Rin_vm_range_ch0": rin_range_0,
        "Rin_ch2_MOhm": rin_ch2,
        "R2_ch2": r2_ch2,
        "n_IV_ch2": n2,
        "Rin_ch2_skip_reason": rin_skip_2,
        "Rin_ch2_note": rin_note_2,
        "Rin_vm_range_ch2": rin_range_2,
        "CC_selection_mode": cc_selection_mode if cc_enabled else None,
        "CC_selection_note": cc_selection_note if cc_enabled else None,
        "analysis_blocks": ",".join(k for k, on in bsel.items() if on),
    }

    gj_02, gj_skip_02 = gj_nS(cc_mean_02, rin_ch2) if cc_enabled else (None, "CC block not selected")
    gj_20, gj_skip_20 = gj_nS(cc_mean_20, rin_ch0) if cc_enabled else (None, "CC block not selected")

    summary_row = build_file_summary_row(
        name,
        rec_dt,
        rin_ch0,
        rin_ch2,
        rin_note_0,
        rin_note_2,
        rin_skip_0,
        rin_skip_2,
        cc_mean_02,
        cc_mean_20,
        n_avg_02,
        n_avg_20,
        gj_02,
        gj_20,
        gj_skip_02,
        gj_skip_20,
        len(sweeps_ch0),
        len(sweeps_ch2),
        n_neg_02=n_neg_02,
        n_neg_20=n_neg_20,
        cell_props={
            **file_summary_extra,
            "CC_mean_ch0to2": cc_mean_02,
            "CC_mean_ch2to0": cc_mean_20,
            "n_CC_avg_ch0to2": n_avg_02,
            "n_CC_avg_ch2to0": n_avg_20,
            "n_CC_neg_ch0to2": n_neg_02,
            "n_CC_neg_ch2to0": n_neg_20,
            "CC_used_neg_ch0to2": "Yes" if n_neg_02 > 0 else "No",
            "CC_used_neg_ch2to0": "Yes" if n_neg_20 > 0 else "No",
            "CC_neg_of_avg_ch0to2": f"{n_neg_02}/{n_avg_02}",
            "CC_neg_of_avg_ch2to0": f"{n_neg_20}/{n_avg_20}",
            "Gj_ch0to2_nS": gj_02,
            "Gj_ch0to2_skip_reason": gj_skip_02,
            "Gj_ch2to0_nS": gj_20,
            "Gj_ch2to0_skip_reason": gj_skip_20,
            "n_CC_sweeps_ch0to2": len(sweeps_ch0),
            "n_CC_sweeps_ch2to0": len(sweeps_ch2),
        },
    )

    rows = []
    if cc_enabled:
        for r in block_02:
            gj_sw, gj_sw_skip = gj_nS(r["CC"], rin_ch2)
            rows.append(sweep_only_row({
                "file": name,
                "recording_datetime": rec_dt,
                "direction": "ch0->ch2",
                "Gj_sweep_nS": gj_sw,
                "Gj_sweep_skip_reason": gj_sw_skip,
                **r,
            }))
        for r in block_20:
            gj_sw, gj_sw_skip = gj_nS(r["CC"], rin_ch0)
            rows.append(sweep_only_row({
                "file": name,
                "recording_datetime": rec_dt,
                "direction": "ch2->ch0",
                "Gj_sweep_nS": gj_sw,
                "Gj_sweep_skip_reason": gj_sw_skip,
                **r,
            }))
        if not rows:
            no_sw_reason = "no sweeps without spikes on active or passive in stimulus period"
            rows.append(sweep_only_row({
                "file": name,
                "recording_datetime": rec_dt,
                "CC_skip_reason": no_sw_reason,
            }))
            summary_row["file_skip_reason"] = no_sw_reason
    else:
        rows.append(sweep_only_row({
            "file": name,
            "recording_datetime": rec_dt,
            "CC_skip_reason": "CC block not selected",
        }))

    want_qc = bsel["cell_props"] or cc_enabled or bsel["tau_cm"]
    if SAVE_QC_PLOTS and want_qc:
        qc_dir = plots_dir
        if not qc_dir:
            qc_dir = cell_props_plots_dir(os.path.dirname(os.path.abspath(filepath)))
        try:
            qc_paths = save_qc_plots(
                abf,
                filepath,
                qc_dir,
                b,
                rin_ch0,
                r2_ch0,
                rin_note_0,
                rin_ch2,
                r2_ch2,
                rin_note_2,
                tau_plot_meta=tau_plot_meta,
                save_ap=bsel["cell_props"],
                save_rin=cc_enabled,
                save_tau=bsel["tau_cm"],
            )
            plot_paths.extend(qc_paths)
        except Exception as exc:
            print(f"  QC plots error: {exc}")

    if SAVE_CC_PLOTS and cc_enabled:
        cc_dir = cc_plots_dir_path
        if not cc_dir:
            cc_dir = os.path.join(os.path.dirname(os.path.abspath(filepath)), CC_PLOTS_SUBDIR)
        try:
            cc_paths = save_cc_qc_plots(
                abf,
                filepath,
                cc_dir,
                b,
                block_02,
                block_20,
                rin_ch0=rin_ch0,
                rin_ch2=rin_ch2,
            )
            plot_paths.extend(cc_paths)
        except Exception as exc:
            print(f"  CC plots error: {exc}")
            traceback.print_exc()

    return rows, summary_row, plot_paths, spikelet_rows, spikelet_sweep_rows
